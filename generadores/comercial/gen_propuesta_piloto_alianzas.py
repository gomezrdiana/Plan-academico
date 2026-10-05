# -*- coding: utf-8 -*-
"""Propuesta de trabajo del piloto de alianzas para el docente que lo lidera (por rol). Lenguaje llano; el otrosi es el respaldo.
Salida: recursos/comercial/B2B/PROPUESTA_PILOTO_ALIANZAS_DOCENTE.docx"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAR = RGBColor(0xE5, 0x4C, 0x09); GRIS = RGBColor(0x77, 0x77, 0x77)
d = Document()
for s in d.sections:
    s.top_margin = Cm(1.4); s.bottom_margin = Cm(1.2); s.left_margin = Cm(1.9); s.right_margin = Cm(1.9)
d.styles['Normal'].font.name = 'Calibri'; d.styles['Normal'].font.size = Pt(10.5)

def shd(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear'); el.set(qn('w:color'), 'auto'); el.set(qn('w:fill'), color); tcPr.append(el)

def h(txt, size=12.5, before=9):
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(3)
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = NAR

def t(txt, bold=False, after=3, size=10.5, color=None):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(txt); r.bold = bold; r.font.size = Pt(size)
    if color: r.font.color.rgb = color

def b(txt, after=2):
    p = d.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(txt); r.font.size = Pt(10.5)

def chk(txt):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(2); p.paragraph_format.left_indent = Cm(0.3)
    r = p.add_run('☐  ' + txt); r.font.size = Pt(10.5)

def tabla(filas, anchos):
    tb = d.add_table(rows=len(filas), cols=len(filas[0])); tb.style = 'Table Grid'
    for i, fila in enumerate(filas):
        for j, txt in enumerate(fila):
            c = tb.rows[i].cells[j]; c.width = Cm(anchos[j])
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            r = p.add_run(txt); r.font.size = Pt(9.5)
            if i == 0: r.bold = True; shd(c, 'FFF8E7')
    d.add_paragraph().paragraph_format.space_after = Pt(1)

h('PROPUESTA DE TRABAJO — PILOTO DE ALIANZAS EMPRESARIALES', 15, before=0)
t('Heiiu English Academy · Para el docente que lidera el piloto · 60 días · octubre a diciembre de 2026', color=GRIS, size=9.5, after=6)

h('1. Qué vamos a hacer')
t('Abrir alianzas con empresas de Bucaramanga. Llevas DOS ofertas y la frase es una sola: "si su empresa quiere premiar a su gente sin tocar presupuesto, firmamos la alianza hoy; si quiere formar a su equipo, le presento el plan empresa".')
t('OFERTA A — Alianza sin costo (la empresa como canal; la cierras tú con el convenio estándar): todos sus colaboradores, con carta de la empresa, entran a la columna Transformación del Fondo (45% de contado, 36% a cuotas), igual que un afiliado a Cajasan. Tres Becas [Empresa] por trimestre, para quien la empresa elija: 45% aunque paguen a cuotas, libros incluidos y prioridad de cupo, entregadas con certificado a nombre de la empresa. Si todos sus becados cumplen, el trimestre siguiente son cinco; si alguno falla, vuelven a tres; nunca más de cinco ni se acumulan. Para la empresa: reporte de avance de sus becados y diagnóstico de inglés para su equipo.', after=3)
t('OFERTA B — Plan empresa (la empresa como cliente; la cierra gerencia, tú la agendas): la empresa paga los cupos de sus colaboradores a tarifa corporativa, $4.782.800 por cupo del programa completo, mínimo 3, con factura y copago posible.', after=3)
t('Tu papel: abrir la puerta, presentar y cerrar el convenio de alianza con el formato estándar, sin cambiarle nada (el convenio no tiene costo ni condiciones negociables). Gerencia lo firma por la academia. Todo lo que tenga precio (formación corporativa pagada) o se salga del formato, lo presenta y cierra gerencia.', bold=True)

h('2. Qué se espera en 60 días')
tabla([
    ['Qué', 'Mínimo', 'Bueno'],
    ['Empresas contactadas (decisor real, con respuesta)', '10', '20'],
    ['Reuniones realizadas con el decisor', '5', '10'],
    ['Convenios firmados (los cierras tú con el formato; gerencia firma por la academia)', '3', '5'],
    ['Gerentes agendados con gerencia por formación pagada', '2', '4'],
    ['Personas que llegan por una alianza', '10', '20'],
    ['Matrículas por alianza (esto es lo que paga comisión)', '3', '6'],
], [10.0, 3.0, 3.0])
t('Ritmo semanal: 3 empresas contactadas y 1 o 2 reuniones. No hace falta más; hace falta que sea todas las semanas.', color=GRIS, size=9.5)

h('3. Lo tuyo')
b('Conseguir y contactar empresas nuevas (las que ya están en conversación con la academia te las decimos antes de empezar: esas no).')
b('Presentar la propuesta con el material aprobado, sin cambiarlo: carta de alianza, términos de las becas, tarifario cerrado.')
b('Cerrar solo con el mini-convenio estándar, tal cual. No negociar, no cotizar, no prometer condiciones ni descuentos. Si preguntan por precio o por formación pagada: "eso se lo presenta gerencia; le agendo 20 minutos".')
b('Terminar cada visita con algo: convenio firmado por la empresa, reunión con gerencia agendada (formación pagada), o un referido concreto.')
b('Llenar el tracker el mismo día de cada contacto y mandar el reporte diario. Lo que no está en el tracker no existe.')
b('Visitas fuera de tus horas de clase. Tus clases no se tocan.')
b('Lenguaje de la casa: "beca", nunca "descuento", "gratis" ni "patrocinio". Nada de promesas de empleo, visas o resultados.')

h('4. Lo de gerencia')
b('Entregarte el material listo: carta de alianza, términos, mini-convenio, tarifario, este documento y el tracker.')
b('Hacer contigo el ensayo antes de la primera visita (gerencia hace de gerente difícil).')
b('Responderte en menos de 24 horas lo que pidas en el reporte diario, y atender a cada gerente que agendes en menos de 48 horas.')
b('Firmar por la academia los convenios que cierres, y negociar, cotizar y cerrar la formación pagada con las empresas que abras.')
b('Liquidarte la comisión cada mes, con el reporte de ingresos a la vista.')

h('5. Lo que llenas y cuándo')
tabla([
    ['Qué', 'Cuándo', 'Cómo'],
    ['Tracker (una fila por empresa: contacto, fecha, canal, estado, próximo paso con fecha)', 'El mismo día de cada contacto', 'Archivo PILOTO_B2B_TRACKER.xlsx. La hoja RESUMEN calcula sola las métricas.'],
    ['Reporte diario', 'Días hábiles, 6 PM (antes si tienes clase en la noche)', 'Un mensaje de WhatsApp de 3 líneas: hoy contacté / reuniones / mañana. Y "necesito de ti", si hay algo.'],
    ['Revisión semanal', 'Viernes, 15 minutos por llamada', 'Tracker en pantalla: la semana contra la meta, y qué ajustar del discurso.'],
    ['Cortes', 'Día 30 y día 60', 'Reunión corta con gerencia: actividad al día 30, resultados al día 60.'],
], [6.0, 4.2, 6.6])

h('6. Cómo se te paga')
b('5% del recaudo de cada matrícula nueva que llegue por un convenio con una empresa que tú registraste y abriste, durante los 12 meses siguientes a la firma del convenio, cierre quien cierre.')
b('3% del recaudo del primer año si una empresa que tú abriste contrata formación corporativa pagada con gerencia.')
b('Se liquida cada mes, contra cuenta de cobro, cuando el pago del estudiante ya pasó el plazo legal de retracto. Transporte y comunicaciones por tu cuenta.')
b('Esto queda escrito en el otrosí a tu contrato docente, que se firma junto con esta propuesta.')

h('7. Qué pasa a los 30 y a los 60 días')
b('Día 30: miramos actividad (contactadas y reuniones) y que haya al menos 1 convenio firmado o 3 reuniones con decisor. Si va, seguimos y ajustamos el discurso. Si no va, una semana de ajuste; si tampoco, cerramos el piloto sin problema.')
b('Día 60: miramos resultados. Con 3 convenios o 3 matrículas por alianza, formalizamos el rol de docente con desarrollo comercial, con metas por trimestre. Si no, lo agradecemos, aprendemos y sigues siendo docente como siempre. Lo que ya se causó se sigue pagando.')

h('8. Checklist de arranque (todo marcado antes de la primera visita)')
for x in ['Contrato docente v3.1 y otrosí del piloto firmados.', 'Lista de empresas excluidas (las que ya conversan con la academia) entregada.', 'Material en la maleta: carta de alianza, términos de becas, mini-convenio, tarifario cerrado.', 'Tracker instalado y con las primeras 10 empresas objetivo escritas.', 'Ensayo con gerencia aprobado: pasó las preguntas de "¿qué cuesta?", "¿qué ganan?" y "deme el descuento en plata" sin regalar nada.', 'Horario de visitas definido, sin cruce con clases.', 'Hora del reporte diario acordada.']:
    chk(x)

d.add_paragraph()
t('Leído y aceptado:   Docente: ______________________________        Gerencia: ______________________________        Fecha: ____ / ____ / 2026', size=10)

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\B2B\PROPUESTA_PILOTO_ALIANZAS_DOCENTE.docx'
d.save(out); print('OK', out)
