# -*- coding: utf-8 -*-
"""TABLERO DE ALIANZAS — Ejecutivo(a) de Alianzas Empresariales.
Hojas: INICIO · TABLERO (semáforo del mes) · ACTIVIDAD_DIARIA (reporte diario) · PLAN_SEMANAL (jornada, para coordinación) · EMPRESAS · LIQUIDACION · METAS · PRIMERA_SEMANA.
Salida: recursos/comercial/INDUCCION_COMERCIAL/contratacion comercial/TABLERO_ALIANZAS.xlsx"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

NAR = 'E54C09'; CREMA = 'FFF8E7'; GRIS = '6B6B6B'; VERDE = 'C6EFCE'; AMAR = 'FFEB9C'; ROJO = 'FFC7CE'
H = Font(bold=True, color='FFFFFF'); B = Font(bold=True); T = Font(bold=True, size=14, color=NAR); S = Font(color=GRIS, italic=True, size=9)
FH = PatternFill('solid', fgColor=NAR); FC = PatternFill('solid', fgColor=CREMA)
thin = Side(style='thin', color='E6DCCB'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
PESOS = '"$"#,##0'; WRAP = Alignment(wrap_text=True, vertical='top')

wb = Workbook()

def header(ws, row, cols, widths=None):
    for j, c in enumerate(cols, 1):
        x = ws.cell(row=row, column=j, value=c); x.font = H; x.fill = FH; x.alignment = Alignment(wrap_text=True, vertical='center'); x.border = BOX
    if widths:
        for j, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[row].height = 32

# ---------------- INICIO ----------------
ws = wb.active; ws.title = 'INICIO'
ws['A1'] = 'TABLERO DE ALIANZAS — Heiiu English Academy'; ws['A1'].font = T
ws['A3'] = 'Fecha de inicio del contrato'; ws['A3'].font = B; ws['B3'] = None; ws['B3'].number_format = 'DD/MM/YYYY'; ws['B3'].fill = FC
ws['C3'] = '← la llena gerencia el primer día (dd/mm/aaaa)'; ws['C3'].font = S
ws['A4'] = 'Mes de contrato que se revisa (1, 2, 3…)'; ws['A4'].font = B; ws['B4'] = '=IF(B3="","",DATEDIF(B3,TODAY(),"m")+1)'; ws['B4'].fill = FC
ws['C4'] = '← por defecto el mes de contrato actual; se puede escribir otro número'; ws['C4'].font = S
ws['A5'] = 'Ese mes va del'; ws['A5'].font = B; ws['B5'] = '=IF(B3="","",EDATE(B3,B4-1))'; ws['B5'].number_format = 'DD/MM/YYYY'
ws['A6'] = 'al'; ws['A6'].font = B; ws['B6'] = '=IF(B3="","",EDATE(B3,B4)-1)'; ws['B6'].number_format = 'DD/MM/YYYY'
ws['A8'] = 'CÓMO SE USA'; ws['A8'].font = T
for i, t in enumerate([
    '1. EMPRESAS: cada empresa, fundación o colegio se registra el MISMO DÍA del primer contacto con quien decide. Quien registra primero, es dueño (atribución).',
    '2. ACTIVIDAD_DIARIA: al cierre de cada día, antes de las 6:00 p. m., una fila: contactos y reuniones con empresas, personas contactadas, citas, inscritos pagados, caja, lo más importante, mañana hago y bloqueos. Es el reporte diario; gerencia lo lee en Drive.',
    '2b. PLAN_SEMANAL: el lunes antes de las 8:30 a. m. se llena la semana (franjas, lugar, qué) y se comparte con coordinación. Es el registro de la jornada: 42 horas.',
    '3. TABLERO: se calcula solo. Muestra el mes contra los mínimos, el semáforo y las fechas de decisión. Gerencia lo revisa a diario.',
    '4. LIQUIDACION: contaduría o gerencia registra lo RECAUDADO en el mes por cada empresa y cada estudiante atribuido; la comisión se calcula sola. Se liquida el día 5 de cada mes y se paga con la nómina.',
    '5. METAS: tabla pactada en el anexo del contrato (no se cambia sin otrosí).',
    '6. PRIMERA_SEMANA: plan de arranque y benchmarking de la competencia.',
    'Definiciones: CONTACTO = habló con quien puede firmar (no recepción). REUNIÓN = presentación de 20+ min con quien decide. FIRMA = convenio firmado u orden de plan empresa firmada CON inicial pagada. ACTIVACIÓN = charla, postulación, feria o jornada en una alianza firmada. VENTA FIRMADA = valor total del contrato con inicial pagada (contado o cuotas).']):
    c = ws.cell(row=9 + i, column=1, value=t); c.alignment = WRAP
    ws.merge_cells(start_row=9 + i, start_column=1, end_row=9 + i, end_column=6); ws.row_dimensions[9 + i].height = 45
ws.column_dimensions['A'].width = 44; ws.column_dimensions['B'].width = 16; ws.column_dimensions['C'].width = 50

# ---------------- EMPRESAS ----------------
we = wb.create_sheet('EMPRESAS')
cols = ['Empresa / fundación / colegio', 'Tipo', 'Sector', 'Decisor (nombre)', 'Cargo', 'Teléfono / correo', 'Fecha 1er contacto con decisor', 'Canal (llamada, visita, feria, referido)',
        'Estado', 'Fecha reunión', 'Oferta', 'Cupos / estudiantes estimados', 'Valor estimado', 'Probabilidad', 'Fecha estimada de cierre', 'Valor firmado', 'Fecha de firma',
        'Inicial pagada (Sí/No)', 'Fecha 1a activación', 'Próximo paso', 'Fecha próximo paso', 'Valor ponderado']
header(we, 1, cols, [30, 13, 14, 20, 16, 20, 13, 16, 18, 12, 16, 11, 14, 11, 13, 14, 12, 10, 12, 30, 12, 14])
N = 400
for r in range(2, N + 2):
    for col in ('G', 'J', 'O', 'Q', 'S', 'U'): we[f'{col}{r}'].number_format = 'DD/MM/YYYY'
    for col in ('M', 'P', 'V'): we[f'{col}{r}'].number_format = PESOS
    we[f'N{r}'].number_format = '0%'
    we[f'V{r}'] = f'=IF(OR(M{r}="",N{r}=""),"",IF(OR(I{r}="Orden pagada",I{r}="Convenio firmado",I{r}="Perdida"),"",M{r}*N{r}))'
dv = DataValidation(type='list', formula1='"Empresa,Fundación,Colegio,Caja / gremio,Entidad pública"', allow_blank=True); we.add_data_validation(dv); dv.add(f'B2:B{N+1}')
dv = DataValidation(type='list', formula1='"Contactada,Reunión hecha,Propuesta enviada,Convenio firmado,Orden pagada,En pausa,Perdida"', allow_blank=True); we.add_data_validation(dv); dv.add(f'I2:I{N+1}')
dv = DataValidation(type='list', formula1='"A. Alianza sin costo,B. Plan empresa,Personalizado,Feria / base de datos"', allow_blank=True); we.add_data_validation(dv); dv.add(f'K2:K{N+1}')
dv = DataValidation(type='list', formula1='"Sí,No"', allow_blank=True); we.add_data_validation(dv); dv.add(f'R2:R{N+1}')
dv = DataValidation(type='list', formula1='"10%,25%,50%,75%,90%"', allow_blank=True); we.add_data_validation(dv); dv.add(f'N2:N{N+1}')
we.freeze_panes = 'B2'
# alerta: convenio firmado sin activación a 30 días
we.conditional_formatting.add(f'A2:A{N+1}', FormulaRule(formula=[f'AND($I2="Convenio firmado",$S2="",$Q2<>"",TODAY()-$Q2>30)'], fill=PatternFill('solid', fgColor=ROJO)))
we.conditional_formatting.add(f'U2:U{N+1}', FormulaRule(formula=[f'AND($U2<>"",$U2<TODAY(),$I2<>"Perdida",$I2<>"Orden pagada")'], fill=PatternFill('solid', fgColor=AMAR)))

# ---------------- ACTIVIDAD_DIARIA ----------------
wa = wb.create_sheet('ACTIVIDAD_DIARIA')
# Las columnas B-F las lee la hoja TABLERO (SUMIFS): no moverlas. Las de personas (G-J) se sumaron el 24/10/2026 porque la meta
# del gerente de alianzas cuenta personas inscritas y caja, no solo empresas.
header(wa, 1, ['Fecha', 'Contactos nuevos con decisor', 'Reuniones hechas', 'Activaciones', 'Firmas (convenio u orden pagada)', 'Ventas firmadas del día ($)',
               'Personas contactadas (referidos, conocidos, afiliados)', 'Citas con personas hechas', 'Inscritos pagados hoy (personas)', 'Caja recibida hoy ($)',
               'Lo más importante de hoy', 'Mañana hago (empresas / reuniones / citas)', 'Bloqueos / lo que necesito de gerencia'],
       [12, 12, 11, 11, 13, 16, 14, 11, 11, 15, 34, 34, 30])
for r in range(2, 302):
    wa[f'A{r}'].number_format = 'DD/MM/YYYY'; wa[f'F{r}'].number_format = PESOS; wa[f'J{r}'].number_format = PESOS
    for col in 'KLM': wa[f'{col}{r}'].alignment = WRAP
wa.freeze_panes = 'B2'
wa['O1'] = 'Se llena al cierre de cada día, antes de las 6:00 p. m., una fila por día (también los días sin resultados: se escribe 0). Gerencia la lee en Drive sin pedirla.'; wa['O1'].font = S
wa.column_dimensions['O'].width = 60

# ---------------- PLAN_SEMANAL ----------------
# Plan de organización del tiempo: el gerente de alianzas lo llena el lunes antes de las 8:30 a. m. y lo comparte con coordinación.
# Es el registro de la jornada (calidad y SST) cuando parte del trabajo se hace fuera de la sede. 42 horas a la semana.
wsem = wb.create_sheet('PLAN_SEMANAL')
wsem['A1'] = 'PLAN DE ORGANIZACIÓN DEL TIEMPO — semana a semana'; wsem['A1'].font = T
wsem['A2'] = ('Se llena el lunes antes de las 8:30 a. m. y se comparte con coordinación: es el registro de la jornada (calidad y seguridad en el trabajo). '
              '42 horas a la semana, en franjas concretas con lugar. Si la semana cambia, se corrige aquí el mismo día. '
              'Contacto con clientes solo de lunes a viernes de 7:00 a. m. a 7:00 p. m. y sábados de 8:00 a. m. a 3:00 p. m.; nunca domingos ni festivos.'); wsem['A2'].alignment = WRAP
wsem.merge_cells('A2:K2'); wsem.row_dimensions[2].height = 48
for j, w in enumerate([11, 9, 9, 15, 30, 9, 9, 15, 30, 10, 26], 1): wsem.column_dimensions[get_column_letter(j)].width = w
lugar = DataValidation(type='list', formula1='"Sede,Casa,Calle / empresa,Virtual"', allow_blank=True); wsem.add_data_validation(lugar)
row = 4
for sem in range(1, 9):
    wsem.cell(row=row, column=1, value='Semana %d · del' % sem).font = B
    c = wsem.cell(row=row, column=2); c.number_format = 'DD/MM/YYYY'; c.fill = FC
    wsem.cell(row=row, column=3, value='Enviado a coordinación:').font = S
    c = wsem.cell(row=row, column=4); c.number_format = 'DD/MM/YYYY HH:MM'; c.fill = FC
    header(wsem, row + 1, ['Día', 'Mañana de', 'a', 'Dónde', 'Qué (empresas, reuniones, citas, oficina)', 'Tarde de', 'a', 'Dónde', 'Qué', 'Horas', 'Observaciones / cambios'])
    first = row + 2
    for i, dia in enumerate(['Lunes', 'Martes', 'Miércoles', 'Jueves', 'Viernes', 'Sábado']):
        r = first + i
        wsem.cell(row=r, column=1, value=dia).font = B
        for col in 'BCFG': wsem[f'{col}{r}'].number_format = 'HH:MM'
        for col in 'EIK': wsem[f'{col}{r}'].alignment = WRAP
        wsem[f'J{r}'] = f'=(IF(AND(B{r}<>"",C{r}<>""),C{r}-B{r},0)+IF(AND(F{r}<>"",G{r}<>""),G{r}-F{r},0))*24'; wsem[f'J{r}'].number_format = '0.0'; wsem[f'J{r}'].fill = FC
        for col in range(1, 12): wsem.cell(row=r, column=col).border = BOX
        lugar.add(f'D{r}'); lugar.add(f'H{r}')
    last = first + 5
    wsem.cell(row=last + 1, column=9, value='Total de la semana').font = B
    wsem[f'J{last + 1}'] = f'=SUM(J{first}:J{last})'; wsem[f'J{last + 1}'].number_format = '0.0'; wsem[f'J{last + 1}'].font = B
    wsem[f'K{last + 1}'] = f'=IF(J{last + 1}=0,"",IF(J{last + 1}<42,"Faltan "&TEXT(42-J{last + 1},"0.0")&" horas para las 42",IF(J{last + 1}>42,"Pasa de 42: ajustar","Cumple 42 horas")))'; wsem[f'K{last + 1}'].font = S
    row = last + 3

# ---------------- METAS ----------------
wm = wb.create_sheet('METAS')
header(wm, 1, ['Trimestre', 'Meses de contrato desde', 'hasta', 'Meta del trimestre (ventas firmadas)', 'Meta al mes (referencia)', 'Programas completos al mes (≈)', 'Básico'], [10, 12, 10, 20, 18, 16, 34])
metas = [('T1', 1, 3, 90e6, 30e6, '6 + 1 Arranque', 'SMMLV + auxilio'), ('T2', 4, 6, 144e6, 48e6, 10, 'SMMLV + auxilio'), ('T3', 7, 9, 198e6, 66e6, 14, 'SMMLV + $550.000 + auxilio (si cumple)'),
         ('T4', 10, 12, 234e6, 78e6, 16, 'SMMLV + $550.000 + auxilio (si cumple)'), ('T5', 13, 15, 252e6, 84e6, 18, 'SMMLV + $850.000 + auxilio (si cumple)'), ('T6', 16, 18, 276e6, 92e6, 19, 'SMMLV + $850.000 + auxilio (si cumple)'),
         ('T7', 19, 21, 300e6, 100e6, 21, 'SMMLV + $850.000 + auxilio (si cumple)'), ('T8', 22, 24, 324e6, 108e6, 23, 'SMMLV + $850.000 + auxilio (si cumple)')]
for i, m in enumerate(metas, 2):
    for j, v in enumerate(m, 1):
        c = wm.cell(row=i, column=j, value=v)
        if j in (4, 5): c.number_format = PESOS
wm['A11'] = 'Desde T9: la meta del trimestre anterior + 2,5% (salvo otrosí). Acelerador: 10% hasta la meta, 15% del 100 al 130%, 20% sobre 130% (trimestres con 2+ empresas nuevas firmadas).'; wm['A11'].font = S
wm['A13'] = 'Mínimos de gestión (cada mes)'; wm['A13'].font = B
header(wm, 14, ['Mes de contrato', 'Contactos con decisor', 'Reuniones', 'Activaciones', 'Firmas acumuladas', 'Meta de ventas del mes'], None)
for i, row in enumerate([(1, 40, 12, 0, 2, 30e6), (2, 40, 12, 2, 4, 30e6), (3, 40, 12, 3, 6, 30e6), ('4+', 40, 12, '1 por alianza por trimestre', '—', 'según trimestre')], 15):
    for j, v in enumerate(row, 1):
        c = wm.cell(row=i, column=j, value=v)
        if j == 6 and isinstance(v, float): c.number_format = PESOS

# ---------------- TABLERO ----------------
wt = wb.create_sheet('TABLERO', 1)
wt['A1'] = 'TABLERO DEL MES'; wt['A1'].font = T
wt['A2'] = 'Periodo revisado (desde)'; wt['B2'] = '=INICIO!B5'; wt['B2'].number_format = 'DD/MM/YYYY'; wt['C2'] = '=INICIO!B6'; wt['C2'].number_format = 'DD/MM/YYYY'
wt['A3'] = 'Mes de contrato'; wt['B3'] = '=INICIO!B4'
wt['A4'] = 'Días hábiles transcurridos / del mes'; wt['B4'] = '=IF(B2="","",NETWORKDAYS(B2,MIN(TODAY(),C2)))'; wt['C4'] = '=IF(B2="","",NETWORKDAYS(B2,C2))'
wt['A5'] = 'Avance del mes (%)'; wt['B5'] = '=IF(B2="","",B4/C4)'; wt['B5'].number_format = '0%'
for c in ('A2', 'A3', 'A4', 'A5'): wt[c].font = B
header(wt, 7, ['Indicador', 'Mínimo / meta del mes', 'Llevamos', 'Debería llevar hoy', 'Semáforo'], [36, 18, 16, 18, 14])
rng = 'ACTIVIDAD_DIARIA!$A:$A,">="&$B$2,ACTIVIDAD_DIARIA!$A:$A,"<="&$C$2'
filas = [
 ('Contactos nuevos con decisor', '40', f'=SUMIFS(ACTIVIDAD_DIARIA!$B:$B,{rng})'),
 ('Reuniones', '12', f'=SUMIFS(ACTIVIDAD_DIARIA!$C:$C,{rng})'),
 ('Activaciones', '=IF($B$3=1,0,IF($B$3=2,2,3))', f'=SUMIFS(ACTIVIDAD_DIARIA!$D:$D,{rng})'),
 ('Firmas del mes', '2', f'=SUMIFS(ACTIVIDAD_DIARIA!$E:$E,{rng})'),
 ('Ventas firmadas del mes', '=IF($B$3="","",IF($B$3<=3,30000000,INDEX(METAS!$E$2:$E$9,MIN(8,ROUNDUP($B$3/3,0)))))', f'=SUMIFS(ACTIVIDAD_DIARIA!$F:$F,{rng})'),
]
for i, (n, meta, llev) in enumerate(filas, 8):
    wt.cell(row=i, column=1, value=n).border = BOX
    wt.cell(row=i, column=2, value=(int(meta) if meta.isdigit() else meta)).border = BOX
    wt.cell(row=i, column=3, value=llev).border = BOX
    wt.cell(row=i, column=4, value=f'=IF($B$5="","",ROUND(B{i}*$B$5,0))').border = BOX
    wt.cell(row=i, column=5, value=f'=IF(OR(B{i}="",$B$5=""),"",IF(C{i}>=D{i},"VERDE",IF(C{i}>=D{i}*0.7,"AMARILLO","ROJO")))').border = BOX
    if i == 12:
        for col in 'BCD': wt[f'{col}{i}'].number_format = PESOS
wt.conditional_formatting.add('E8:E12', CellIsRule(operator='equal', formula=['"VERDE"'], fill=PatternFill('solid', fgColor=VERDE)))
wt.conditional_formatting.add('E8:E12', CellIsRule(operator='equal', formula=['"AMARILLO"'], fill=PatternFill('solid', fgColor=AMAR)))
wt.conditional_formatting.add('E8:E12', CellIsRule(operator='equal', formula=['"ROJO"'], fill=PatternFill('solid', fgColor=ROJO)))
wt['A13'] = 'Firmas acumuladas desde el inicio'; wt['A13'].font = B
wt['C13'] = '=COUNTIFS(EMPRESAS!$I:$I,"Convenio firmado")+COUNTIFS(EMPRESAS!$I:$I,"Orden pagada")'
wt['B13'] = '=IF($B$3=1,2,IF($B$3=2,4,IF($B$3=3,6,"—")))'

wt['A15'] = 'EMBUDO DEL MES (de cada 10 contactos, ¿cuántos llegan?)'; wt['A15'].font = B
wt['A16'] = 'Reuniones / contactos'; wt['B16'] = '=IFERROR(C9/C8,"")'; wt['B16'].number_format = '0%'; wt['C16'] = 'referencia: 30% (12 de 40)'
wt['A17'] = 'Firmas / reuniones'; wt['B17'] = '=IFERROR(C11/C9,"")'; wt['B17'].number_format = '0%'; wt['C17'] = 'referencia: 17% (2 de 12)'
for c in ('C16', 'C17'): wt[c].font = S

wt['A19'] = 'PIPELINE (lo que viene)'; wt['A19'].font = B
wt['A20'] = 'Valor ponderado en negociación'; wt['B20'] = '=SUM(EMPRESAS!$V:$V)'; wt['B20'].number_format = PESOS
wt['A21'] = 'Con cierre estimado en los próximos 30 días'; wt['B21'] = '=SUMPRODUCT((EMPRESAS!$O$2:$O$401>=TODAY())*(EMPRESAS!$O$2:$O$401<=TODAY()+30)*(EMPRESAS!$V$2:$V$401<>"")*N(+EMPRESAS!$V$2:$V$401))'; wt['B21'].number_format = PESOS
wt['A22'] = 'Convenios firmados sin activar (más de 30 días)'; wt['B22'] = '=SUMPRODUCT((EMPRESAS!$I$2:$I$401="Convenio firmado")*(EMPRESAS!$S$2:$S$401="")*(EMPRESAS!$Q$2:$Q$401<>"")*(TODAY()-EMPRESAS!$Q$2:$Q$401>30))'
wt['A23'] = 'Próximos pasos vencidos'; wt['B23'] = '=SUMPRODUCT((EMPRESAS!$U$2:$U$401<>"")*(EMPRESAS!$U$2:$U$401<TODAY())*(EMPRESAS!$I$2:$I$401<>"Perdida")*(EMPRESAS!$I$2:$I$401<>"Orden pagada"))'

wt['A25'] = 'FECHAS DE DECISIÓN (primer trimestre)'; wt['A25'].font = B
header(wt, 26, ['Fecha', 'Qué se revisa', 'Si no se cumple'], None)
dec = [('=IF(INICIO!$B$3="","",WORKDAY(INICIO!$B$3-1,10))', 'Alerta temprana (≈ día 15): 20 contactos, 6 reuniones, 1 firma', 'Conversación de ajuste esa misma semana, con plan escrito'),
       ('=IF(INICIO!$B$3="","",EDATE(INICIO!$B$3,1)-1)', 'Cierre del mes 1: gestión + 2 firmas (gracia de 5 días para lo firmado)', 'Sin ingreso mínimo garantizado ese mes; plan de ajuste de una semana'),
       ('=IF(INICIO!$B$3="","",INICIO!$B$3+49)', 'DÍA 50: decisión de continuidad ANTES de vencer el período de prueba (2 meses)', 'Terminar dentro del período de prueba'),
       ('=IF(INICIO!$B$3="","",EDATE(INICIO!$B$3,3)-1)', 'Cierre del T1: $90M de meta y 6 firmas acumuladas', 'Revisión de continuidad y del plan del T2')]
for i, (f, q, s) in enumerate(dec, 27):
    wt.cell(row=i, column=1, value=f).number_format = 'DD/MM/YYYY'; wt.cell(row=i, column=2, value=q).alignment = WRAP; wt.cell(row=i, column=3, value=s).alignment = WRAP
    wt.row_dimensions[i].height = 30
wt.column_dimensions['B'].width = 46; wt.column_dimensions['C'].width = 40

# ---------------- LIQUIDACION ----------------
wl = wb.create_sheet('LIQUIDACION')
wl['A1'] = 'LIQUIDACIÓN DE COMISIONES — sobre lo RECAUDADO (no sobre lo firmado)'; wl['A1'].font = T
wl['A2'] = 'Mes liquidado (primer día)'; wl['B2'] = None; wl['B2'].number_format = 'MMMM YYYY'; wl['B2'].fill = FC; wl['A2'].font = B
wl['A3'] = 'Básico del mes (SMMLV + auxilio vigente)'; wl['B3'] = 2000000; wl['B3'].number_format = PESOS; wl['B3'].fill = FC; wl['A3'].font = B
wl['A4'] = '¿Cumplió gestión y firmas? (meses 1-3)'; wl['B4'] = 'Sí'; wl['B4'].fill = FC; wl['A4'].font = B
dv = DataValidation(type='list', formula1='"Sí,No,No aplica (mes 4+)"', allow_blank=False); wl.add_data_validation(dv); dv.add('B4')
header(wl, 6, ['Empresa o estudiante', 'Tipo de comisión', 'Fecha de firma de la matrícula / orden', 'Recaudado en el mes ($)', 'Meses desde la firma', '% que aplica', 'Comisión ($)', 'Soporte (recibo / factura)'],
       [32, 26, 16, 16, 12, 11, 14, 24])
dv = DataValidation(type='list', formula1='"Plan empresa / personalizado (10%),Cartera de alianza o base propia,Cierre de apoyo (comisión de asesora)"', allow_blank=True); wl.add_data_validation(dv); dv.add('B7:B306')
for r in range(7, 307):
    wl[f'C{r}'].number_format = 'DD/MM/YYYY'; wl[f'D{r}'].number_format = PESOS; wl[f'G{r}'].number_format = PESOS; wl[f'F{r}'].number_format = '0%'
    wl[f'E{r}'] = f'=IF(OR(C{r}="",$B$2=""),"",DATEDIF(C{r},EOMONTH($B$2,0),"m")+1)'
    wl[f'F{r}'] = f'=IF(B{r}="","",IF(LEFT(B{r},4)="Plan",IF(E{r}<=12,0.10,0),IF(LEFT(B{r},7)="Cartera",IF(E{r}<=12,0.06,IF(E{r}<=24,0.03,0)),"tarifa asesora")))'
    wl[f'G{r}'] = f'=IF(OR(D{r}="",F{r}=""),"",IF(ISNUMBER(F{r}),D{r}*F{r},""))'
wl['I6'] = 'Nota'; wl['I6'].font = H; wl['I6'].fill = FH
wl['I7'] = 'Cierre de apoyo: escribir la comisión de asesora a mano en G. El acelerador (15% / 20%) se liquida al cierre del trimestre aparte.'; wl['I7'].alignment = WRAP; wl.column_dimensions['I'].width = 40
wl['A309'] = 'Total comisiones del mes'; wl['A309'].font = B; wl['G309'] = '=SUM(G7:G306)'; wl['G309'].number_format = PESOS
wl['A310'] = 'Básico'; wl['G310'] = '=B3'; wl['G310'].number_format = PESOS
wl['A311'] = 'Complemento por ingreso mínimo garantizado ($3.500.000, meses 1-3 si cumplió)'; wl['G311'] = '=IF(B4="Sí",MAX(0,3500000-(G309+G310)),0)'; wl['G311'].number_format = PESOS
wl['A312'] = 'TOTAL A PAGAR (antes de deducciones de ley)'; wl['A312'].font = B; wl['G312'] = '=G309+G310+G311'; wl['G312'].number_format = PESOS; wl['G312'].font = B
wl.freeze_panes = 'A7'

# ---------------- PRIMERA_SEMANA ----------------
wp = wb.create_sheet('PRIMERA_SEMANA')
wp['A1'] = 'PRIMERA SEMANA — arranque'; wp['A1'].font = T
header(wp, 3, ['Día', 'Qué', 'Entregable', 'Hecho (Sí/No)'], [8, 60, 40, 12])
for i, (d, q, e) in enumerate([
    ('1', 'Inducción: producto, Fondo de Becas, convenio, orden de plan empresa, tarifa por volumen, módulo de graduación, garantía.', 'Kit leído y preguntas resueltas'),
    ('1-2', 'Benchmarking de 7 academias (abajo): llamada o visita como cliente, web y biblioteca de anuncios de Meta.', 'Tabla de benchmarking completa'),
    ('2-3', 'Lista de 100 empresas, fundaciones y colegios objetivo, con decisor identificado, cargadas en EMPRESAS.', '100 filas en EMPRESAS'),
    ('3', 'Simulación de presentación con gerencia (gerente de una empresa de 40 empleados).', 'Presentación aprobada'),
    ('4-5', 'Primeros contactos con decisor y primeras reuniones agendadas.', '≥ 10 contactos y ≥ 3 reuniones agendadas'),
    ('5', 'Formato de captura de datos para ferias con autorización de tratamiento de datos (Ley 1581).', 'Formato listo')], 4):
    wp.cell(row=i, column=1, value=d); wp.cell(row=i, column=2, value=q).alignment = WRAP; wp.cell(row=i, column=3, value=e).alignment = WRAP
    wp.row_dimensions[i].height = 32
wp['A12'] = 'BENCHMARKING DE LA COMPETENCIA'; wp['A12'].font = T
header(wp, 13, ['Academia', '¿Garantía de aprendizaje por escrito?', 'Oferta a empresas y precios', 'Convenio con cajas de compensación', 'Certificaciones', 'Qué prometen en sus anuncios', 'Fuente y fecha'], None)
for j, w in enumerate([18, 22, 30, 20, 20, 30, 18], 1): wp.column_dimensions[get_column_letter(j)].width = max(wp.column_dimensions[get_column_letter(j)].width or 0, w)
for i, a in enumerate(['Smart', 'Colombo Americano', 'Best English', 'Praxis', 'Lingui', 'Berlitz', 'Celai'], 14):
    wp.cell(row=i, column=1, value=a).font = B
    for j in range(2, 8): wp.cell(row=i, column=j).alignment = WRAP
    wp.row_dimensions[i].height = 30
wp['A22'] = 'Regla: hasta completar esta tabla, con empresas NO se dice "la única academia que la firma"; se dice "garantía de aprendizaje por escrito, sin letra pequeña".'; wp['A22'].font = S

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\INDUCCION_COMERCIAL\contratacion comercial\TABLERO_ALIANZAS.xlsx'
try: wb.save(out); print('OK', out)
except PermissionError: wb.save(out.replace('.xlsx', '_v2.xlsx')); print('ABIERTO -> _v2')
