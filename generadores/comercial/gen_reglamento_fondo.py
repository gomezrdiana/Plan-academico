# -*- coding: utf-8 -*-
"""Reglamento del Fondo + Carta de Aprobacion — entregable imprimible en formato de marca."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NARANJA = RGBColor(0xE5, 0x4C, 0x09)
GRIS = RGBColor(0x77, 0x77, 0x77)
CREMA = 'FFF8E7'
NARANJA_HEX = 'E54C09'

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.6); s.bottom_margin = Cm(1.4); s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)

def shd(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement('w:shd'); el.set(qn('w:val'), 'clear'); el.set(qn('w:fill'), color)
    tcPr.append(el)

def titulo(txt, size=17):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = NARANJA
    p.paragraph_format.space_after = Pt(4)

def texto(txt, size=11, bold=False):
    p = doc.add_paragraph(); r = p.add_run(txt); r.font.size = Pt(size); r.bold = bold
    p.paragraph_format.space_after = Pt(4)
    return p

def art(num, nombre, cuerpo):
    p = doc.add_paragraph()
    r = p.add_run('Artículo ' + num + ' — ' + nombre + '. '); r.bold = True; r.font.size = Pt(11)
    r2 = p.add_run(cuerpo); r2.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(5)

def tabla(filas, fs=10.5):
    tab = doc.add_table(rows=len(filas), cols=len(filas[0])); tab.style = 'Table Grid'
    for i, fila in enumerate(filas):
        for j, val in enumerate(fila):
            c = tab.rows[i].cells[j]
            p = c.paragraphs[0]; r = p.add_run(str(val)); r.font.size = Pt(fs)
            if i == 0:
                r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shd(c, NARANJA_HEX)
    esp = doc.add_paragraph(); esp_r = esp.add_run(''); esp_r.font.size = Pt(4)
    esp.paragraph_format.space_after = Pt(0)

# ================= PAGINA 1: REGLAMENTO =================
titulo('FONDO DE BECAS HEIIU — REGLAMENTO')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Heiiu English Academy · Global Teacher S.A.S. · v3 — Septiembre 2026'); r.font.size = Pt(10); r.font.color.rgb = GRIS

art('1', 'Naturaleza', 'El Fondo de Becas Heiiu es un programa institucional de Global Teacher S.A.S. que subsidia parcialmente el valor de los programas a estudiantes seleccionados. Los recursos del Fondo provienen de la propia institución.')
art('2', 'Administración', 'El Fondo es administrado y auditado por la contaduría externa de la institución. Ningún funcionario de Heiiu — incluida la gerencia — tiene facultad individual para otorgar, modificar o restituir becas por fuera de este reglamento.')
art('3', 'Niveles de beca', 'El PERFIL del aspirante define el techo de su beca: (a) Transformación — techo 45%, exclusiva para estrato 1-2 / SISBÉN con soporte, cupos limitados por trimestre; (b) General — techo 32%; (c) Ejecutivo — techo 15%. El COMPROMISO define cuánto del techo se otorga, según el paquete contratado y la forma de pago. Beca efectiva sobre la tarifa pública vigente:')
tabla([
 ['Paquete (contado / cuotas)', 'TRANSFORMACIÓN', 'GENERAL', 'EJECUTIVO'],
 ['A1→B2 completo', '45% / 36%', '32% / 25,6%', '15% / 12%'],
 ['A2→B2', '36% / 27%', '25,6% / 19,2%', '12% / 9%'],
 ['B1→B2', '24,8% / 15,8%', '17,6% / 11,2%', '8,3% / 5,3%'],
 ['Solo B2', '20,3% / 11,3%', '14,4% / 8%', '6,8% / 3,8%'],
 ['Nivel suelto / sabatino', 'Sin beca', 'Sin beca', 'Sin beca'],
])
texto('3.4 — Convenio Cajasan: el aspirante que acredite afiliación vigente a Cajasan (carné o certificado) y pague a cuotas recibe el porcentaje de contado de su perfil. Este beneficio no modifica el techo del perfil ni se acumula con otros.', 10.5)
art('4', 'Condiciones de permanencia', 'La beca se mantiene mientras el estudiante cumpla TODAS: (1) asistencia mínima del 80%; (2) puntualidad en los pagos acordados (ninguna cuota con más de 5 días hábiles de mora); (3) cumplimiento de tareas y entregables del programa, incluido el portafolio de práctica; (4) comportamiento conforme al contrato de matrícula.')
art('5', 'Pérdida y recuperación', 'El incumplimiento de cualquier condición suspende la beca: las cuotas siguientes se liquidan a tarifa plena. El estudiante puede solicitar POR UNA ÚNICA VEZ la recuperación de la beca acreditando 30 días continuos de cumplimiento total. La solicitud se presenta por escrito a coordinación y la resuelve la administración del Fondo.')
art('6', 'Vigencia de la aprobación', 'Toda beca aprobada se comunica por escrito con fecha de vencimiento. La aprobación caduca a las 72 horas si no se formaliza la matrícula; con abono de separación, el cupo y las condiciones se mantienen por 7 días calendario. Vencido el plazo, la solicitud puede presentarse de nuevo sin garantía del mismo porcentaje ni cupo.')
art('7', 'Sin excepciones', 'Este reglamento no admite excepciones individuales. Una excepción a un estudiante compromete la sostenibilidad del Fondo para todas las familias.')

doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ================= PAGINA 2: CARTA =================
titulo('FONDO DE BECAS HEIIU — CARTA DE APROBACIÓN')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Heiiu English Academy · Global Teacher S.A.S.'); r.font.size = Pt(10); r.font.color.rgb = GRIS
doc.add_paragraph()
texto('Estudiante: ______________________________________  Documento: ____________________', 12)
texto('Programa: ____________________________  Perfil: ____________________  Fecha: ____ / ____ / ______', 12)
doc.add_paragraph()
texto('Nos complace informarle que su solicitud ante el Fondo de Becas Heiiu fue APROBADA con una beca del ________ % sobre la tarifa pública vigente.', 12, bold=True)
texto('Valor cubierto por el Fondo: $ ____________________  ·  Inversión final del estudiante: $ ____________________', 12)
doc.add_paragraph()
texto('Esta aprobación vence el ____ / ____ / ______ a las 6:00 PM (Artículo 6 del Reglamento). Con abono de separación, sus condiciones se mantienen por 7 días calendario.', 11.5)
texto('El beneficio está sujeto al Reglamento del Fondo (entregado junto con esta carta), en particular a las condiciones de permanencia del Artículo 4. El estudiante y su acudiente declaran conocerlas y aceptarlas.', 11.5)
for _ in range(3): doc.add_paragraph()
texto('_______________________________                    _______________________________', 12)
texto('Administración del Fondo                                        Estudiante / Acudiente', 11)
doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Los artefactos hacen la realidad: esta carta se entrega SIEMPRE impresa y firmada, junto con el reglamento.'); r.font.size = Pt(9.5); r.font.color.rgb = GRIS

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\documentos_operativos\politicas\REGLAMENTO_FONDO_Y_CARTA_IMPRIMIBLE.docx'
doc.save(out); print('OK', out)
