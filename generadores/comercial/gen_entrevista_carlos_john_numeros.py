# -*- coding: utf-8 -*-
"""Entrevista de NÚMEROS que hace el socio financiero (Carlos) al candidato de alianzas (John), 30 min virtual, 2 págs.
Reemplaza la guía larga de cross-reference. Salida en Candidatos/ (fuera de git):
  - ENTREVISTA_CARLOS_JOHN_NUMEROS.docx (para Carlos)
  - LLAMADA_REFERENCIA_WELLI_DIANA.docx (para Diana: 6 preguntas al jefe directo)"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAR = RGBColor(0xE5, 0x4C, 0x09); GRIS = RGBColor(0x66, 0x66, 0x66)
BASE = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\INDUCCION_COMERCIAL\contratacion comercial\LINKEDIN DIANA GOMEZ\Candidatos'

def nuevo_doc():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.2); s.bottom_margin = Cm(1.0); s.left_margin = Cm(1.8); s.right_margin = Cm(1.8)
    doc.styles['Normal'].font.name = 'Calibri'; doc.styles['Normal'].font.size = Pt(10)
    return doc
def shd(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear'); el.set(qn('w:color'), 'auto'); el.set(qn('w:fill'), color); tcPr.append(el)
def titulo(doc, t1, t2):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
    r = p.add_run(t1); r.bold = True; r.font.size = Pt(13.5); r.font.color.rgb = NAR
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(6)
    r = p.add_run(t2); r.font.size = Pt(9); r.font.color.rgb = GRIS
def sec(doc, txt):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(7); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(11.5); r.font.color.rgb = NAR
def t(doc, txt, bold=False, size=10, color=None, after=3):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(txt); r.font.size = Pt(size); r.bold = bold
    if color: r.font.color.rgb = color
def preg(doc, txt):
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(1)
    r = p.add_run(txt); r.font.size = Pt(10)
def busca(doc, txt):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4); p.paragraph_format.left_indent = Cm(0.6)
    r = p.add_run('Qué buscas: '); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = GRIS
    r = p.add_run(txt); r.font.size = Pt(9.5); r.font.color.rgb = GRIS
def caja(doc, lines):
    tb = doc.add_table(rows=1, cols=1); tb.style = 'Table Grid'; c = tb.cell(0, 0); shd(c, 'FFF8E7')
    c.paragraphs[0].paragraph_format.space_after = Pt(0)
    r = c.paragraphs[0].add_run('\n'.join(lines)); r.font.size = Pt(9.5); r.italic = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

# ---------- 1. ENTREVISTA PARA CARLOS ----------
doc = nuevo_doc()
titulo(doc, 'ENTREVISTA DE NÚMEROS — JOHN HINESTROZA', 'Cargo: Gerente Comercial de Alianzas · La hace Carlos, socio financiero · Virtual · 30 minutos · Uso interno, no se entrega al candidato')
t(doc, 'Lo que John ya contó (su historia en Welli, Nexa y Dentix, por qué salió) lo verifica Diana con la ficha que él llenó y con su jefe directo de Welli. Esta entrevista no repite esa historia: lo pone a construir números en vivo. Un comercial de verdad se sabe su embudo de memoria; el que no, da cifras redondas y vagas. Tú anotas cada número exacto que diga.', bold=True)
t(doc, 'Lo que ya tienes para leer antes: la ficha de cifras que llenó el 14/10, su tarea práctica (5 empresas con decisor, mensaje de apertura, plan de 2 semanas) y la propuesta del cargo. En su plan él mismo propuso 5 indicadores y una meta de $30.000.000.', size=9.5, color=GRIS)

sec(doc, 'ARRANQUE (1 min)')
caja(doc, ['"Soy Carlos, socio financiero de Heiiu. Reviso los números de todo cargo comercial. Diana ya habló contigo de la causa; vamos a hablar solo de cifras."'])

sec(doc, '1. SU EMBUDO EN WELLI, DE MEMORIA (8 min)')
preg(doc, '"En una semana normal en Welli: ¿cuántas clínicas visitabas, cuántas te recibían, cuántas firmaban, cuántas de esas llegaban a desembolsar y cuánto tardaba el primer desembolso?"')
preg(doc, '"¿Cuál era tu tasa de visita a firma y de firma a aliado activo? Dime el número."')
preg(doc, '"¿Cuántos de los 180 aliados estaban activos el último mes? ¿Qué porcentaje de los desembolsos lo hacían los 20 primeros?"')
preg(doc, '"¿Cuánto le costaba a Welli cada aliado activo, contando tu salario y tu comisión? ¿Lo calculaste alguna vez?"')
busca(doc, 'números específicos y rápidos (ej. "15 visitas, me recibían 10, firmaban 3, activaban 2, primer desembolso a las 3 semanas"). Si dice "muchas", "bastante", "depende", o se demora, no manejaba sus indicadores: se los manejaban.')

sec(doc, '2. CÓMO REPORTABA (4 min)')
preg(doc, '"¿Qué le reportabas a tu director, cada cuánto y en qué formato? Descríbeme el tablero."')
preg(doc, '"¿Cuál era el indicador que te medían primero? ¿Y cuál era el que tú mirabas primero?"')
preg(doc, '"Cuéntame un mes en que ibas por debajo de meta: qué viste en los números, en qué semana, y qué cambiaste."')
busca(doc, 'que mencione indicadores de actividad (visitas, contactos), de conversión y de plata, y que sepa cuál fue el que le avisó. Si solo habla del resultado final, no leía el proceso.')

sec(doc, '3. QUE ARME EL EMBUDO DE HEIIU EN VIVO (10 min)')
preg(doc, '"Los mínimos del cargo son 40 empresas con decisor y 12 reuniones al mes, y 2, 4 y 6 firmas acumuladas en los meses 1 a 3. Constrúyeme el embudo: de 40 contactos, cuántas reuniones, cuántas propuestas, cuántas firmas, cuántas matrículas por convenio, cuánto recaudo y en qué mes entra."')
preg(doc, '"¿Qué cuenta para ti como decisor contactado?"')
busca(doc, 'la respuesta correcta es "hablé con la persona que decide y tengo nombre, cargo y celular". Si cuenta correos enviados o mensajes de LinkedIn, sus 40 son 8.')
preg(doc, '"Tus 5 empresas de la tarea son grandes, con directores nacionales; eso firma en 2 o 3 meses. ¿Qué firmas en el mes 1 mientras eso madura?"')
preg(doc, '"Dime los 5 indicadores que me vas a mostrar cada martes y, para cada uno, el número al que te sacas tú solo."')
busca(doc, 'que las cifras del embudo cuadren entre sí (si dice 40 contactos y 20 reuniones, no sabe lo que es un decisor). Y que el último punto lo diga sin pensar: un gerente comercial sabe cuándo está fallando antes de que se lo digan.')

sec(doc, '4. LA PLATA Y LAS CONDICIONES (5 min)')
preg(doc, '"Explícame con tus palabras cómo se te paga: el básico, el piso de los 3 meses y la comisión. ¿Sobre qué se paga la comisión?"')
busca(doc, 'sobre lo cobrado, no sobre lo firmado. Si no lo sabe, se le explica una vez y se anota.')
preg(doc, '"Tu plan dice $30 millones. ¿En qué mes entra esa plata a la cuenta?"')
busca(doc, 'que distinga firmado, inicial pagada y recaudo real.')
preg(doc, '"Diana te dijo que no tenías que ir a la oficina. Lo que sí hay: martes y jueves en la sede, con tu planilla abierta, con Diana y conmigo. ¿Te sirve?"')
busca(doc, 'Diana se lo dejó pasar en la reunión. Lo corriges tú, sin drama, como un hecho del cargo.')

sec(doc, 'CIERRE (2 min)')
caja(doc, ['"¿Hay algo de todo esto que te haría decir que no?"  → y te callas. Lo que diga es su objeción real; se anota, no se responde.',
           '"Gracias. Diana te escribe esta semana."  → sin oferta, sin fecha, sin hablar del título de gerente.'])

sec(doc, 'DESPUÉS: 3 LÍNEAS A DIANA, EL MISMO DÍA')
preg(doc, '¿Se sabe sus números o se los sabían por él?')
preg(doc, '¿El embudo de Heiiu que armó cuadra o es discurso?')
preg(doc, 'Veredicto: contratar, no contratar, o una condición para contratar.')
t(doc, 'Lo que no haces: venderle Heiiu, discutir el título, ofrecer plata o fecha. Tú verificas; Diana ofrece después. No preguntar por embarazo, hijos, estado civil, religión, política ni salud.', size=9.5, color=GRIS)

out1 = BASE + r'\ENTREVISTA_CARLOS_JOHN_NUMEROS.docx'
try: doc.save(out1); print('OK', out1)
except PermissionError: doc.save(out1.replace('.docx', '_v2.docx')); print('ABIERTO -> _v2')

# ---------- 2. LLAMADA A LA REFERENCIA (DIANA) ----------
doc = nuevo_doc()
titulo(doc, 'LLAMADA A LA REFERENCIA DE WELLI — LA HACE DIANA', 'Elkin Suárez, director comercial nacional de Welli · 318 716 9143 · 10 minutos · Antes de la entrevista con Carlos')
t(doc, 'No es una confrontación: es escuchar y anotar hechos. No le menciones ninguna cifra de John; que las diga él. Si una respuesta no cuadra con la ficha, no lo discutes con Elkin: lo anotas y se lo pasas a Carlos para la entrevista.', bold=True)
sec(doc, 'Apertura (tal cual)')
caja(doc, ['"Buenas, Elkin, le habla Diana Gómez, gerente de Heiiu English Academy en Bucaramanga. John Hinestroza lo dio como referencia y autorizó por escrito que lo llamáramos. ¿Tiene 10 minutos?"'])
sec(doc, 'Las 6 preguntas, en este orden')
for q, nota in [
 ('"¿Qué cargo tenía John con usted y de qué fechas a qué fechas?"', 'Su HV: gerente comercial Santanderes, abril 2025 a "actualmente".'),
 ('"¿Cuántos aliados tenía la zona cuando él llegó y cuántos cuando salió?"', 'Él escribió de 7 a 180. No se lo digas.'),
 ('"¿Esos aliados los abrió él solo o con un equipo? ¿Cuántos estaban activos al final?"', ''),
 ('"¿Cómo fue su cumplimiento de metas, mes a mes, el último semestre?"', 'Él dice 85 a 110%.'),
 ('"¿Por qué salió?"', 'Él dijo: la zona se volvió logística y cambiaron las condiciones comerciales. Si Elkin dice otra cosa, esa es la respuesta.'),
 ('"¿Lo volvería a contratar para abrir una zona desde cero?"', 'Y te callas. El silencio o el "pues..." ya responde.'),
]:
    p = doc.add_paragraph(style='List Number'); p.paragraph_format.space_after = Pt(1)
    r = p.add_run(q); r.font.size = Pt(10)
    if nota:
        r = p.add_run('  ' + nota); r.font.size = Pt(9.5); r.font.color.rgb = GRIS
    t(doc, 'Respuesta: ________________________________________________________________________', size=9.5, after=4)
sec(doc, 'Cierre')
caja(doc, ['"Le agradezco mucho su tiempo. Que esté bien."  → no cuentas si lo vas a contratar ni qué te pareció.'])
t(doc, 'Después: una línea a Carlos con lo que no cuadró (si nada, "Welli confirma la ficha").', size=9.5, color=GRIS)

out2 = BASE + r'\LLAMADA_REFERENCIA_WELLI_DIANA.docx'
try: doc.save(out2); print('OK', out2)
except PermissionError: doc.save(out2.replace('.docx', '_v2.docx')); print('ABIERTO -> _v2')
