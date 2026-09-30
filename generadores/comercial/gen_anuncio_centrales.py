# -*- coding: utf-8 -*-
"""Anuncio de prueba + mensaje de entrada por WhatsApp para la linea de personas con inmueble reportadas (negocio de Carlos).
Salida: recursos/comercial/CFO/ANUNCIO_PRUEBA_CENTRALES_v1.docx"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x16, 0x33, 0x5A); GRIS = RGBColor(0x66, 0x66, 0x66)
d = Document()
for s in d.sections:
    s.top_margin = Cm(1.5); s.bottom_margin = Cm(1.3); s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)
d.styles['Normal'].font.name = 'Calibri'; d.styles['Normal'].font.size = Pt(10.5)

def shd(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear'); el.set(qn('w:color'), 'auto'); el.set(qn('w:fill'), color); tcPr.append(el)

def h(txt, size=12.5, before=8):
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = NAVY

def t(txt, bold=False, after=4, size=10.5, color=None, italic=False):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(txt); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if color: r.font.color.rgb = color

def caja(txt):
    tb = d.add_table(rows=1, cols=1); tb.style = 'Table Grid'
    c = tb.rows[0].cells[0]; shd(c, 'F3F1EA')
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    r = p.add_run(txt); r.font.size = Pt(11)
    d.add_paragraph().paragraph_format.space_after = Pt(2)

h('ANUNCIO DE PRUEBA v1 — SALIR DE CENTRALES CON INMUEBLE', 15, before=0)
t('Para Carlos. Prueba pequeña en Meta (hasta $500.000). El anuncio no dice precio ni plazo ni promete resultado; el filtro (tener inmueble) va adentro para que solo escriba quien califica. Los textos van en los cuadros grises, listos para copiar.', color=GRIS, size=9.5, after=8)

h('1. Texto del anuncio (imagen fija o video corto de Carlos hablando a cámara, 30 segundos)')
t('Gancho:', bold=True, after=2)
caja('¿Tienes casa o apartamento a tu nombre, pero estás reportado y ningún banco te presta?')
t('Cuerpo:', bold=True, after=2)
caja('Muchos reportes en centrales de riesgo tienen errores: sin aviso previo, deudas que ya prescribieron, datos mal registrados. La ley permite reclamarlos. Estudiamos tu caso, te decimos por escrito qué se puede hacer y en cuánto tiempo, y si tu reporte sale, te conseguimos el crédito que necesitas con tu inmueble como respaldo. Más de 20 años en finanzas, en Colombia y Estados Unidos.')
t('Cierre:', bold=True, after=2)
caja('Escríbenos por WhatsApp y cuéntanos dos cosas: si el inmueble está a tu nombre, y con quién está tu reporte.')
t('Versión corta para el video (Carlos a cámara, 30 s): "Si tienes un inmueble a tu nombre y estás reportado, hay dos cosas que casi nadie te dice. Uno: muchos reportes tienen errores y la ley permite reclamarlos. Dos: con tu inmueble como respaldo sí hay quien te preste. Estudiamos tu caso y te decimos por escrito qué se puede hacer. Escríbeme."', size=9.5, color=GRIS, italic=True, after=8)

h('2. Primera respuesta por WhatsApp (la escribe el equipo de Carlos)')
caja('Hola, gracias por escribir. Soy [nombre], del equipo de Carlos Remolina. Para saber si tu caso se puede estudiar, tres preguntas rápidas:\n1. ¿El inmueble está a tu nombre y sin hipoteca vigente?\n2. ¿Con qué entidad está el reporte y desde cuándo?\n3. ¿La deuda sigue existiendo o ya la pagaste?\nCon eso te digo si tiene sentido estudiarlo.')

h('3. Segunda respuesta (si califica: tiene inmueble)')
caja('Con lo que me cuentas sí vale la pena estudiarlo. Cómo funciona: revisamos tu reporte y la deuda con nuestro equipo jurídico y te entregamos por escrito qué se puede hacer y el tiempo estimado. El valor depende del monto reportado y te lo cotizo después de revisar tu caso; se paga la mitad al empezar y la otra mitad cuando el reclamo queda radicado y tienes la constancia en la mano; si tomas el crédito, ese saldo se descuenta del desembolso. Si sale, te presentamos a un fondeador para el crédito con tu inmueble como garantía. ¿Te agendo una llamada de 20 minutos con Carlos para revisar tu caso? ¿Te queda mejor mañana en la mañana o en la tarde?')

h('4. Si NO califica (sin inmueble, o con hipoteca vigente)')
caja('Gracias por contarme. Sin un inmueble libre como respaldo no podemos ofrecerte el crédito, y no te voy a hacer perder tiempo. Lo que sí podemos es estudiar tu reporte: si tiene errores se puede reclamar. Si te interesa, me dices y te cuento el valor.')

h('5. Lo que NO se dice nunca (ni en el anuncio ni en el chat)')
for x in ['"Te borramos el reporte", "te sacamos de Datacrédito" ni "sales en X días".', 'Tasas ni montos de crédito por escrito: se hablan en la llamada.', 'Nombres de casos ni de clientes.', 'La palabra "gratis" y cualquier promesa de resultado.']:
    t('• ' + x, after=2)

h('6. Cómo se mide la prueba')
tb = d.add_table(rows=4, cols=2); tb.style = 'Table Grid'
for i, (a, b) in enumerate([('Qué se mide', 'Señal de que funciona'), ('Personas que escriben y contestan las tres preguntas', '10 o más con inmueble en el primer mes'), ('Personas que pagan la primera mitad', '2 o más'), ('Costo por persona con inmueble', 'Menos de $50.000')]):
    c0, c1 = tb.rows[i].cells; c0.width = Cm(8.5); c1.width = Cm(8.5)
    r0 = c0.paragraphs[0].add_run(a); r1 = c1.paragraphs[0].add_run(b); r0.font.size = r1.font.size = Pt(9.5)
    if i == 0: r0.bold = r1.bold = True; shd(c0, 'F3F1EA'); shd(c1, 'F3F1EA')
d.add_paragraph()
t('Regla de la prueba: se corre 3 semanas con el mismo anuncio. Si a la semana 2 no hay nadie con inmueble, se cambia el gancho, no el presupuesto.', size=9.5, color=GRIS, italic=True)


h('7. VARIANTE PARA EMPRESAS: FONDEO CON GARANTÍA INMOBILIARIA (misma cuenta, segundo anuncio)')
t('Público: dueños de empresa o negocio en Colombia. Mismo principio: filtro adentro, sin tasas ni promesas.', color=GRIS, size=9.5, after=4)
t('Gancho:', bold=True, after=2)
caja('¿Tu empresa necesita capital y el banco ya te dijo que no?')
t('Cuerpo:', bold=True, after=2)
caja('Cuando el banco cierra la puerta, no se acabaron las opciones. Con tus estados financieros y un inmueble como respaldo, hay fondeadores que sí prestan. Revisamos tus números, te decimos con claridad si tu empresa califica y cuánto, y te presentamos al fondeador correcto. Más de 20 años en finanzas, en Colombia y Estados Unidos.')
t('Cierre:', bold=True, after=2)
caja('Escríbenos por WhatsApp y cuéntanos dos cosas: cuánto necesitas y si tienes un inmueble para respaldar.')
t('Primera respuesta por WhatsApp:', bold=True, after=2)
caja('Hola, gracias por escribir. Soy [nombre], del equipo de Carlos Remolina. Para saber si podemos ayudarte, tres preguntas rápidas:\n1. ¿Cuánto necesita la empresa y para qué?\n2. ¿Tienes un inmueble libre (sin hipoteca vigente) para respaldar el crédito?\n3. ¿Tienes estados financieros del último año?\nCon eso te digo si tiene sentido revisar tu caso.')
t('Si califica:', bold=True, after=2)
caja('Con eso sí vale la pena revisarlo. Cómo funciona: nos envías los estados financieros y los documentos del inmueble, revisamos tu capacidad y te decimos cuánto y en qué condiciones se puede fondear; luego te presentamos al fondeador. Nuestra comisión es del 6% sobre lo desembolsado y se paga solo al desembolso. ¿Te agendo una llamada de 20 minutos con Carlos? ¿Mañana en la mañana o en la tarde?')
t('Si los números no alcanzan para presentarlo a un fondeador, se le ofrece la Radiografía Financiera: es el mismo cliente, en otro momento.', size=9.5, color=GRIS, italic=True, after=6)

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\CFO\ANUNCIO_PRUEBA_CENTRALES_v1.docx'
try:
    d.save(out); print('OK', out)
except PermissionError:
    alt = out.replace('_v1.docx', '_v1_NUEVO.docx'); d.save(alt); print('ABIERTO -> guardado como', alt)
