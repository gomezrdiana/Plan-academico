# -*- coding: utf-8 -*-
"""(1) Mensaje a los grupos de estudiantes activos sobre video diario y garantia (Experiencia al Cliente).
(2) Hoja de compromiso del estudiante activo: recibe el anexo de garantia vigente, video diario como tarea, puntualidad.
Salidas en documentos_operativos/plantillas/."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

NAR = RGBColor(0xE5, 0x4C, 0x09); GRIS = RGBColor(0x77, 0x77, 0x77)
BASE = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\documentos_operativos\plantillas'

def nuevo():
    d = Document()
    for s in d.sections:
        s.top_margin = Cm(1.6); s.bottom_margin = Cm(1.4); s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)
    d.styles['Normal'].font.name = 'Calibri'; d.styles['Normal'].font.size = Pt(10.5)
    return d

def t(d, txt, bold=False, after=5, size=10.5, color=None, align=None, italic=False):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    if align: p.alignment = align
    r = p.add_run(txt); r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = color
    return p

# ---------- 1. MENSAJE ----------
d = nuevo()
t(d, 'MENSAJE PARA LOS GRUPOS DE WHATSAPP DE LOS COHORTES ACTIVOS', bold=True, size=13, color=NAR)
t(d, 'Lo envía Experiencia al Cliente (no el profesor). Se copia y pega tal cual en cada grupo. Donde dice [minutos], se pone el mínimo del nivel del grupo: A1 un minuto · A2 dos minutos · B1 tres minutos · B2 cinco minutos.', color=GRIS, size=9.5, after=10)
for par in [
    'Hola a todos.',
    'Les recordamos cómo funciona su garantía de aprendizaje: pide asistir al menos al 80% de las clases, entregar al menos el 90% de las tareas y presentar todas las evaluaciones del nivel.',
    'Desde esta semana, la tarea de cada clase incluye un video corto en inglés: [minutos] hablando de corrido, sin leer, enviado a este grupo el mismo día. Cada video cuenta como tarea entregada. Su profesor lo ve y al empezar la clase le da a cada uno una recomendación concreta para mejorar.',
    'Es la práctica que más rápido los hace hablar, y es la prueba de su avance.',
    'Más adelante les entregaremos por escrito su garantía y estos compromisos, para que cada uno tenga claro qué le corresponde a la academia y qué le corresponde a usted.',
    'Experiencia al Cliente — Heiiu English Academy',
]:
    t(d, par, after=8)
d.save(BASE + r'\MENSAJE_GRUPOS_VIDEO_GARANTIA.docx'); print('OK mensaje')

# ---------- 2. HOJA DE COMPROMISO ----------
d = nuevo()
t(d, 'HOJA DE COMPROMISO DEL ESTUDIANTE', bold=True, size=14, color=NAR, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
t(d, 'Garantía de aprendizaje: asistencia, puntualidad, trabajo en clase y tareas · Heiiu English Academy — Global Teacher S.A.S., NIT 900.422.478-2', size=9, color=GRIS, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)

t(d, 'Yo, ______________________________________________, identificado(a) con documento No. ____________________, estudiante del nivel ______ en el grupo / jornada ____________________, declaro lo siguiente:', after=8)

items = [
    ('1. Recibo mi garantía por escrito. ',
     'Recibo en este acto el ANEXO DE GARANTÍA HEIIU DE APRENDIZAJE vigente, lo leí y lo acepto. Entiendo que la garantía aplica cuando cumplo sus tres condiciones: asistir al menos al 80% de las clases del nivel, entregar al menos el 90% de las tareas asignadas y presentar todas las evaluaciones del nivel.'),
    ('2. El video diario es parte de mis tareas. ',
     'Me comprometo a grabar y enviar al grupo de WhatsApp de mi cohorte, el mismo día de cada clase, un video en inglés hablado de corrido y sin leer, con la duración mínima de mi nivel (A1: 1 minuto · A2: 2 minutos · B1: 3 minutos · B2: 5 minutos). Entiendo que cada video es una tarea asignada y cuenta dentro del 90% de tareas de mi garantía. Puedo enviar audio en lugar de video solo de forma excepcional, máximo dos veces por semana.'),
    ('3. Trabajo en clase. ',
     'La clase es hablada y en inglés: participo en las simulaciones y en los ejercicios, entrego mi ticket de salida al final de cada clase y mantengo el celular guardado durante la clase. Lo que se practica en clase es lo que la garantía mide después.'),
    ('4. Recibo feedback, no notas, de mi profesor. ',
     'Mi profesor ve mis videos y me da una recomendación concreta en clase. Las notas, los resultados de las evaluaciones y todo lo relacionado con la garantía me los comunica coordinación académica.'),
    ('5. Puntualidad. ',
     'La clase empieza a la hora exacta de mi horario. Llegar más de quince (15) minutos después cuenta como llegada tarde, y cada tres (3) llegadas tarde se contabilizan como una (1) inasistencia. Entiendo que la puntualidad es parte de la formación del programa.'),
    ('6. Esta hoja no cambia mi contrato. ',
     'Precisa cómo se cumplen las condiciones de la garantía y hace parte de mi expediente académico. Todas las demás cláusulas de mi contrato de matrícula siguen iguales. La academia conserva mis videos y grabaciones únicamente como evidencia de mi avance académico, conforme a la Ley 1581 de 2012.'),
]
for tit, txt in items:
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(6); p.paragraph_format.left_indent = Cm(0.3)
    r = p.add_run(tit); r.bold = True; r.font.size = Pt(10.5)
    r2 = p.add_run(txt); r2.font.size = Pt(10.5)

d.add_paragraph()
t(d, 'Firmada en Bucaramanga el ____ / ____ / ______', after=14)
t(d, 'ESTUDIANTE                                                                    GLOBAL TEACHER S.A.S.', after=16)
t(d, '_____________________________                                _____________________________', after=2)
t(d, 'Nombre: ________________________                              Coordinación académica', after=2)
t(d, 'Documento: _____________________', after=12)
t(d, 'Si el estudiante es menor de edad, firma también su acudiente:', size=9.5, color=GRIS, after=10)
t(d, '_____________________________     Nombre del acudiente: ____________________     Documento: ______________', size=9.5, after=12)
t(d, 'Se entrega con esta hoja una copia del Anexo de Garantía vigente. Original al expediente del estudiante; copia al estudiante.', size=8.5, color=GRIS, italic=True)
d.save(BASE + r'\HOJA_COMPROMISO_ESTUDIANTE.docx'); print('OK hoja')
