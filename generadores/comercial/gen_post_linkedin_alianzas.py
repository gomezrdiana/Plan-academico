# -*- coding: utf-8 -*-
"""Post y kit de publicación de la vacante Ejecutivo(a) Senior de Alianzas Empresariales (LinkedIn, redes, portales).
Base: PROPUESTA_EJECUTIVA_ALIANZAS.pdf. Mismo formato que POST_LINKEDIN_VACANTE*.xlsx.
Salida: recursos/comercial/INDUCCION_COMERCIAL/contratacion comercial/POST_LINKEDIN_VACANTE_ALIANZAS.xlsx"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

EMAIL = 'citasheiiu@gmail.com'; ASUNTO = 'Alianzas Heiiu'
wb = Workbook()

def hoja(nombre, lineas, ancho=120, primera=False):
    ws = wb.active if primera else wb.create_sheet(nombre)
    ws.title = nombre; ws.column_dimensions['A'].width = ancho
    for i, t in enumerate(lineas, 1):
        c = ws.cell(row=i, column=1, value=t); c.alignment = Alignment(wrap_text=True, vertical='top')
        if t and (t.isupper() or t.startswith('TÍTULO') or t.startswith('QUÉ') or t.startswith('CÓMO')): c.font = Font(bold=True)
    return ws

hoja('POST LINKEDIN', [
 'Busco a la persona que abra las puertas de las empresas, los colegios y las fundaciones de Bucaramanga al inglés.',
 'En Heiiu English Academy abrimos UNA vacante de CONTRATACIÓN INMEDIATA: Ejecutivo(a) Senior de Alianzas Empresariales.',
 'Es un cargo de campo y de relacionamiento con la alta dirección: gerencias, talento humano, bienestar, rectorías y fundaciones. La misión es liderar el canal corporativo e institucional de Heiiu, construirlo y hacerlo crecer, con el respaldo de un equipo comercial y académico.',
 'QUÉ VAS A LLEVAR A LAS EMPRESAS',
 'No vendemos clases de inglés. Formamos personas: inglés con garantía de aprendizaje por escrito, formación en virtudes y hábitos en cada clase, y resultados que se pueden mostrar (la página web propia del estudiante, su pitch de negocio en inglés o su video para una vacante). Con el respaldo de una academia presencial con licencia desde 2012, triple certificación ICONTEC (NTC 5555, NTC 5580 e ISO 9001) y convenio con Cajasan. Las ofertas, el convenio y los precios ya están escritos: alianza sin costo para la empresa y plan empresa por volumen.',
 'QUÉ OFRECEMOS',
 '✔ Contrato laboral desde el primer día, con todas las prestaciones de ley.',
 '✔ Básico (salario mínimo vigente + auxilio) + comisiones sin techo: 10% de lo que paga cada empresa y cartera mensual por cada estudiante de tus alianzas.',
 '✔ Ingreso mínimo garantizado de $3.500.000 en tus 3 primeros meses, cumpliendo los indicadores de gestión.',
 '✔ Metas trimestrales claras, acelerador por sobrecumplimiento y aumentos de básico por resultados.',
 '✔ Un equipo que te apoya en las activaciones desde el primer día, y la posibilidad de crecer hacia liderar un equipo comercial o nuevas sedes.',
 'QUÉ BUSCAMOS',
 '• Ventas como tu fuerte: mínimo 5 años vendiendo y cerrando negocios corporativos (B2B), alianzas o convenios, con resultados que puedas demostrar con cifras. Idealmente en educación, salud, seguros o servicios.',
 '• Cierre con quien decide: gerencias, talento humano, bienestar, rectorías y fundaciones.',
 '• Manejo de números: metas, presupuestos comerciales, indicadores y embudo de ventas. Que leas tu propio tablero y anticipes tus resultados.',
 '• Mercadeo y activación: eventos, ferias, charlas corporativas, BTL.',
 '• Deseable: contratación pública (SECOP), manejo de cartera y recaudo, o liderazgo de equipos comerciales.',
 '• Formación profesional en áreas comerciales, de mercadeo, administración, economía, contaduría o afines.',
 '• Disciplina de registro, curiosidad para proponer nuevas oportunidades y trabajo en Bucaramanga y su área metropolitana.',
 'CÓMO APLICAR',
 f'Envía tu hoja de vida a {EMAIL} con el asunto "{ASUNTO}" y, esto es lo importante, un audio o video de 1 minuto contándonos la alianza o el convenio más importante que hayas cerrado y cómo lo lograste. Sin el audio no revisamos la hoja de vida. En la entrevista nos presentarás la alianza como si fuéramos el gerente de una empresa de 40 empleados.',
 'Si conoces a alguien así, etiquétalo(a). La persona correcta llega a un cargo donde su esfuerzo se nota en su bolsillo.',
 '#Bucaramanga #EmpleoBucaramanga #AlianzasEstratégicas #VentasB2B #DesarrolloDeNegocios #Educación',
], primera=True)

hoja('PIEZA GRAFICA', [
 'TEXTO PARA LA PIEZA GRÁFICA (imagen del post)',
 'Titular: BUSCAMOS EJECUTIVO(A) SENIOR DE ALIANZAS EMPRESARIALES',
 'Subtítulo: Mínimo 5 años de experiencia B2B. CONTRATACIÓN INMEDIATA.',
 '· Contrato laboral + prestaciones',
 '· $3.500.000 garantizados los 3 primeros meses',
 '· 10% de lo que paga cada empresa + cartera mensual',
 '· Comisiones sin techo · Bucaramanga',
 'Cierre: Aplica con tu hoja de vida + un audio de 1 minuto. Detalles en el post.',
 'Marca: logo Heiiu, naranja #E54C09 sobre crema #FFF8E7. El amarillo #FBC900 nunca como texto sobre blanco.',
])

hoja('PARA BRIAN', [
 'PARA BRIAN — SOLO LINKEDIN Y REDES',
 '1. PIEZA PRINCIPAL: crear la vacante gratuita en LinkedIn Empleos desde la página de empresa de Heiiu, con el texto del post (título: Ejecutivo(a) Senior de Alianzas Empresariales — Contratación inmediata · ubicación: Bucaramanga · presencial · tiempo completo). Aparece por búsqueda, no depende de seguidores.',
 '2. Publicar el post en la página de Heiiu con la pieza gráfica, y compartirlo desde su perfil como apoyo, siempre enlazando a la vacante oficial.',
 '3. Replicar la pieza gráfica en Instagram y Facebook de la academia.',
 '4. Si se le pone pauta: segmentar Bucaramanga y área metropolitana; cargos actuales tipo ejecutivo de cuenta corporativa, gerente comercial, director comercial, ejecutivo de alianzas, desarrollo de negocios, key account manager; sectores educación, salud prepagada y asistencia médica, seguros, cajas de compensación, servicios empresariales.',
 f'5. Las hojas de vida llegan a {EMAIL} con el asunto "{ASUNTO}".',
])

hoja('PORTALES (ACADEMIA)', [
 'PORTALES DE EMPLEO GRATUITOS — LOS PUBLICA LA ACADEMIA, NO BRIAN',
 'Dónde: Computrabajo · Magneto · elempleo · Servicio Público de Empleo (Agencia de Cajasan o Comfenalco) · grupos de empleo de Bucaramanga en Facebook · grupos de egresados de administración y mercadeo (UNAB, UIS, UPB, Santo Tomás).',
 'Datos para llenar los formularios:',
 '· Cargo: Ejecutivo(a) Senior de Alianzas Empresariales',
 '· Empresa: Heiiu English Academy (Global Teacher S.A.S.)',
 '· Ciudad: Bucaramanga y área metropolitana · tiempo completo · trabajo de campo',
 '· Contrato: laboral, con período de prueba de ley · CONTRATACIÓN INMEDIATA',
 '· Salario: básico (SMMLV + auxilio) + comisiones sin techo; ingreso mínimo garantizado de $3.500.000 los 3 primeros meses según indicadores',
 '· Experiencia: mínimo 5 años en VENTAS corporativas B2B, alianzas o convenios, con resultados demostrables; manejo de indicadores y presupuestos comerciales; mercadeo y activación (deseable SECOP, cartera o liderazgo de equipos)',
 '· Formación: profesional en áreas comerciales, mercadeo, administración, economía, contaduría o afines',
 '· Descripción: copiar el texto de la hoja POST LINKEDIN, sin los hashtags.',
 f'· Contacto: {EMAIL}, asunto "{ASUNTO}".',
 'Filtro al recibir: quien no mande el audio o video de 1 minuto queda descartado sin revisar. A los que sí lo manden se les envía la PROPUESTA_EJECUTIVA_ALIANZAS.pdf y se agenda la entrevista con la simulación (presentar la alianza a "un gerente de una empresa de 40 empleados").',
])

hoja('MENSAJE DIRECTO', [
 'MENSAJE DIRECTO PARA PERFILES QUE YA ESTÁN TRABAJANDO (búsqueda activa en LinkedIn)',
 'Hola [NOMBRE], vi tu perfil y tu experiencia en [alianzas / venta corporativa] en [EMPRESA]. En Heiiu English Academy, en Bucaramanga, estamos buscando, para contratación inmediata, UNA persona para construir nuestro canal de alianzas con empresas, colegios y fundaciones: contrato laboral, ingreso mínimo garantizado de $3.500.000 los tres primeros meses, 10% de lo que pague cada empresa y cartera mensual por cada estudiante de tus alianzas, sin techo.',
 'Si te interesa conocer la propuesta completa, o conoces a alguien con ese perfil, cuéntame y te la comparto.',
 'A quién buscar: personas en Bucaramanga con cargos de ejecutivo de cuenta corporativa, director o gerente comercial, ejecutivo de alianzas, desarrollo de negocios o key account en asistencia médica, salud prepagada, seguros, cajas de compensación, educación o servicios empresariales. Meta: 20 mensajes directos en la primera semana.',
])

hoja('NOTA INTERNA (no publicar)', [
 'NOTA INTERNA — no se publica ni se envía a Brian',
 'Documentos del proceso: PROPUESTA_EJECUTIVA_ALIANZAS.pdf (se envía a quien pase el filtro del audio), ANEXO_COMPENSACION_EJECUTIVO_ALIANZAS_BORRADOR.docx (abogado), TABLERO_ALIANZAS.xlsx (seguimiento desde el día 1).',
 'En la publicación NO se dice "la única academia que firma la garantía" hasta completar el benchmarking de la primera semana (Smart, Colombo, Best English, Praxis, Lingui, Berlitz, Celai).',
 'Ya hay una candidata en conversación (perfil ex directora comercial de AME); la publicación abre el abanico y sirve de referencia de mercado.',
 'Correo de recepción: citasheiiu@gmail.com, asunto "Alianzas Heiiu" (filtrar en Gmail por el asunto).',
])

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\INDUCCION_COMERCIAL\contratacion comercial\POST_LINKEDIN_VACANTE_ALIANZAS.xlsx'
try: wb.save(out); print('OK', out)
except PermissionError: wb.save(out.replace('.xlsx', '_v2.xlsx')); print('ABIERTO -> _v2')
