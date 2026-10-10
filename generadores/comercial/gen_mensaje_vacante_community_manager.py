# -*- coding: utf-8 -*-
"""Mensaje para candidatos a Community Manager (Heiiu + constructora) y criterios de descarte.
Salida: recursos/comercial/contratacion/COMMUNITY_MANAGER/MENSAJE_VACANTE_COMMUNITY_MANAGER.docx"""
import os, re
from docx import Document
from docx.shared import Pt, RGBColor

NAR = RGBColor(0xE5, 0x4C, 0x09); GRIS = RGBColor(0x66, 0x66, 0x66)
d = Document(); d.styles['Normal'].font.name = 'Calibri'; d.styles['Normal'].font.size = Pt(11)
for s in d.sections: s.left_margin = s.right_margin = Pt(60); s.top_margin = s.bottom_margin = Pt(50)

def p(t, size=11, color=None, after=6, bold_all=False):
    x = d.add_paragraph(); x.paragraph_format.space_after = Pt(after)
    for part in re.split(r'(\*\*[^*]+\*\*)', t):
        if not part: continue
        r = x.add_run(part[2:-2] if part.startswith('**') else part); r.bold = bold_all or part.startswith('**'); r.font.size = Pt(size)
        if color: r.font.color.rgb = color
def bullet(t):
    x = d.add_paragraph(style='List Bullet'); x.paragraph_format.space_after = Pt(2)
    for part in re.split(r'(\*\*[^*]+\*\*)', t):
        if not part: continue
        r = x.add_run(part[2:-2] if part.startswith('**') else part); r.bold = part.startswith('**')
def num(t):
    x = d.add_paragraph(style='List Number'); x.paragraph_format.space_after = Pt(3)
    for part in re.split(r'(\*\*[^*]+\*\*)', t):
        if not part: continue
        r = x.add_run(part[2:-2] if part.startswith('**') else part); r.bold = part.startswith('**')

p('MENSAJE PARA CANDIDATOS — COMMUNITY MANAGER', 14, NAR, 2, True)
p('Copiar desde la línea "¡Hola!" hasta "¡Gracias!". Antes de enviar, reemplazar [día] y [hora].', 9, GRIS, 12)

p('¡Hola! Gracias por tu interés en la vacante de **Community Manager**.')
p('Buscamos a una persona creativa y curiosa, que sepa de marketing y quiera entender qué contenido funciona y por qué, para manejar **dos marcas por igual**: una **academia de inglés** y una **constructora**.')
p('**El servicio incluye, para CADA una de las dos marcas, como MÍNIMO:**', after=2)
for t in ['Proponer la **parrilla de contenido** del mes.',
          '**Mínimo 5 historias diarias** (videos, encuestas, preguntas, reposts, fotos, detrás de cámaras).',
          '**Mínimo 20 videos editados al mes** (Reels/TikTok), es decir **mínimo 40 en total** entre las dos marcas.',
          '**Mínimo 8 publicaciones de foto, carrusel o pieza gráfica al mes**, mínimo 16 en total.',
          '**Testimonios grabados** de estudiantes y de clientes.',
          '**Publicar y programar** todo el contenido.',
          'Un **informe mensual** con lo que funcionó, lo que no y lo que propone.']:
    bullet(t)
p('Estas cantidades son el piso, no el techo: el material que se grabe y sirva se edita y se publica.', after=6)
p('**En la academia**, el contenido principal son las **clases reales y los testimonios de estudiantes**: la mayoría de sus videos los muestran, con **mínimo 2 testimonios por semana**.')
p('**Modalidad:** prestación de servicios, **semipresencial en Bucaramanga**: mínimo **3 jornadas de grabación por semana en la academia** (clases, estudiantes, testimonios) y **mínimo 1 por semana en la constructora**, agendadas cada semana. **La edición se hace de forma remota.** Disponibilidad prioritaria de lunes a sábado. Honorarios: **$1.800.000 mensuales**. Se requiere **equipo propio**: celular o cámara, micrófono y programas de edición de video e imagen.')
p('Si te interesa, envíanos **un solo correo a citasheiiu@gmail.com con el asunto "Vacante CM"**, **antes del [día] a las [hora]**, con:', after=2)
for t in ['**Tu hoja de vida** y el enlace a tu portafolio o a cuentas que hayas manejado.',
          '**Tu disponibilidad** de lunes a sábado (mañanas, tardes, noches) y en qué zona de Bucaramanga vives.',
          '**El equipo y los programas** con los que trabajas.',
          '**Un video corto grabado y editado por ti**, vertical, de 30 a 45 segundos, como un Reel para Instagram. Tema libre: un negocio, un lugar o una persona de tu entorno. Debe tener **subtítulos**, un **gancho en los primeros 3 segundos** y un **cierre que invite a hacer algo**. Adjúntalo al correo o envía el enlace de Google Drive (con acceso abierto), para que no pierda calidad.',
          '**En 3 líneas: ¿cómo conseguirías que un estudiante acepte grabar su testimonio y que se vea natural, no actuado?**']:
    num(t)
p('Si el correo no tiene el asunto "Vacante CM" o no incluye los 5 puntos, no podremos continuar con tu proceso. ¡Gracias!', after=16)

p('USO INTERNO — CÓMO DESCARTAR (no enviar)', 12, NAR, 4, True)
for t in ['1. No mandó los 5 puntos o no puso el asunto "Vacante CM" → fuera (si no sigue una instrucción simple, no cumple una parrilla).',
          '2. No puede venir a grabar a Bucaramanga o vive muy lejos para 4 jornadas semanales → fuera.',
          '3. El video no tiene gancho, subtítulos ni cierre → fuera. Es lo que más pesa.',
          '4. El punto 5 es genérico ("contenido de valor", "tips") → baja prioridad.',
          '5. No tiene equipo propio → fuera.',
          '6. Sin experiencia en cuentas reales → baja prioridad; se compensa con un video muy bueno.',
          'Contrato: prestación de servicios por entregables (mínimos, "piso, no techo"), con jornadas agendadas semanalmente — no horario fijo impuesto, para no configurar contrato realidad. Revisar con el abogado.']:
    p(t, 10, GRIS, 3)

out_dir = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\contratacion\COMMUNITY_MANAGER'
os.makedirs(out_dir, exist_ok=True)
out = os.path.join(out_dir, 'MENSAJE_VACANTE_COMMUNITY_MANAGER.docx')
try: d.save(out); print('OK', out)
except PermissionError: d.save(out.replace('.docx', '_v2.docx')); print('ABIERTO -> _v2')
