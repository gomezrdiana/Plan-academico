# -*- coding: utf-8 -*-
"""Hoja rápida de cierre a cuotas para el Gerente de Alianzas (por rol): qué beca aplica, reglas del plan de pagos,
tabla de cuotas tipo (programa completo, 1,9% mensual sobre saldo) y la frase de cierre. Números del motor de cuotas
(MOTOR_CUOTAS_FONDO_2026.xlsx) y de ACLARACION_BECAS_CUOTAS. Salida (git):
recursos/comercial/INDUCCION_COMERCIAL/contratacion comercial/HOJA_CIERRE_CUOTAS_ALIANZAS.docx (+PDF)."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAR = RGBColor(0xE5, 0x4C, 0x09); GRIS = RGBColor(0x66, 0x66, 0x66); BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
OUT = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\INDUCCION_COMERCIAL\contratacion comercial\HOJA_CIERRE_CUOTAS_ALIANZAS.docx'
r = 0.019
def cuota(s, n): return round(s * r / (1 - (1 + r) ** -n))
def m(v): return '$' + format(v, ',').replace(',', '.')

d = Document()
for s in d.sections:
    s.top_margin = Cm(0.9); s.bottom_margin = Cm(0.7); s.left_margin = Cm(1.3); s.right_margin = Cm(1.3)
d.styles['Normal'].font.name = 'Calibri'; d.styles['Normal'].font.size = Pt(9)
def shd(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear'); el.set(qn('w:color'), 'auto'); el.set(qn('w:fill'), color); tcPr.append(el)
def sec(t):
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(1)
    x = p.add_run(t); x.bold = True; x.font.size = Pt(10.5); x.font.color.rgb = NAR
def t(txt, size=9, color=None, after=1, bold=False):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    x = p.add_run(txt); x.font.size = Pt(size); x.bold = bold
    if color: x.font.color.rgb = color
def bullets(items, size=9):
    for it in items:
        p = d.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(0); x = p.add_run(it); x.font.size = Pt(size)
def caja(txt, size=9):
    tb = d.add_table(rows=1, cols=1); tb.style = 'Table Grid'; c = tb.cell(0, 0); shd(c, 'FFF8E7')
    c.paragraphs[0].paragraph_format.space_after = Pt(0); x = c.paragraphs[0].add_run(txt); x.font.size = Pt(size); x.italic = True
    d.add_paragraph().paragraph_format.space_after = Pt(0)
def tabla(hdr, filas, anchos, size=8.5):
    tb = d.add_table(rows=1, cols=len(hdr)); tb.style = 'Table Grid'
    for i, h in enumerate(hdr):
        c = tb.rows[0].cells[i]; shd(c, 'E54C09'); c.paragraphs[0].paragraph_format.space_after = Pt(0)
        x = c.paragraphs[0].add_run(h); x.bold = True; x.font.size = Pt(size); x.font.color.rgb = BLANCO
    for f in filas:
        row = tb.add_row().cells
        for i, txt in enumerate(f):
            row[i].paragraphs[0].paragraph_format.space_after = Pt(0)
            x = row[i].paragraphs[0].add_run(txt); x.font.size = Pt(size); x.bold = (i == 0)
    for row in tb.rows:
        for i, w in enumerate(anchos): row.cells[i].width = Cm(w)
    d.add_paragraph().paragraph_format.space_after = Pt(0)

p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
x = p.add_run('HOJA RÁPIDA DE CIERRE A CUOTAS — ALIANZAS EMPRESARIALES'); x.bold = True; x.font.size = Pt(13); x.font.color.rgb = NAR
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(4)
x = p.add_run('Uso interno · Programa completo A1 a B2 (lista $8.696.000) · Se da el plan completo: inicial + número de cuotas + valor de cada cuota, nunca solo la cuota'); x.font.size = Pt(8.5); x.font.color.rgb = GRIS

sec('1. Qué beca aplica (nunca se suman dos: se aplica la mejor)')
tabla(['Quién es', 'Contado', 'A cuotas (base)', 'Inicial mínima 20%'], [
 ('Con carta de empresa aliada (beca de alianza)', m(4782800), m(5565440), m(1113088)),
 ('Sin carta, pero estrato 1 o 2 con soporte, o afiliado a Cajasan con carné vigente', m(4782800), m(5565440), m(1113088)),
 ('Sin carta, todos los demás (beca General del Fondo)', m(5913280), m(6469824), m(1293965)),
 ('Arranque A1+A2 (sin beca, 200 h)', m(2990000), 'Dos pagos: mitad al inscribirse, mitad antes de la primera clase. Sin interés, sin más cuotas.', ''),
], [7.4, 2.6, 4.6, 3.4])
t('Contado = transferencia a la cuenta de Heiiu, tarjeta, cesantías o crédito de un banco o cooperativa: beca de contado y sin interés nuestro. Cuotas = crédito directo de Heiiu, con interés del 1,9% mensual sobre saldo, ya incluido en la tabla de abajo.', size=8.5, color=GRIS)

sec('2. Reglas del plan de pagos (no se negocian)')
bullets([
 'Separación de cupo: $300.000 hoy; se cruza con la inicial. La inicial mínima (20%) se completa máximo 3 días después. Las cuotas arrancan a los 30 días de la inicial.',
 'Máximo de cuotas: el estudiante termina de pagar antes de terminar el programa. Intensivo (10 h/semana, ~14 meses): máximo 11 cuotas. Súper intensivo (20 h/semana, ~7 meses): máximo 5. Menos cuotas siempre se puede y paga menos interés; más que el máximo, nunca.',
 'Inicial mayor, cuota menor: 30%, 40% o 50% de inicial están en la tabla. Cualquier otra combinación (otra inicial, 7 o 9 cuotas, otro paquete) la calcula coordinación con el motor en el momento; nunca de memoria.',
 'Todo pago entra a Heiiu (transferencia o tarjeta); efectivo a la sede el mismo día. El estudiante firma en la sede con coordinación: contrato de matrícula, anexo de garantía y, si va a cuotas, pagaré, carta y plan de pagos. Menor de 18: firma el acudiente.',
 'No se dice "financiación sin interés" ni "descuento". Se dice "beca del Fondo" y "crédito directo de la academia".',
 'Menores de 12 a 17: por lo general sabatino, que es otro producto (nivel suelto, sin beca del Fondo); se agenda con coordinación.'])

sec('3. Tabla de cuotas tipo: programa completo, 1,9% mensual sobre saldo')
filas = []
for nombre, base in [('Beca de alianza / Transformación', 5565440), ('Beca General', 6469824)]:
    for pct in [0.2, 0.3, 0.4, 0.5]:
        ini = round(base * pct); s = base - ini
        filas.append((nombre + ' · inicial ' + str(int(pct * 100)) + '%', m(ini)) + tuple(m(cuota(s, n)) for n in [3, 5, 6, 8, 11]))
tabla(['Beca y porcentaje de inicial', 'Inicial', '3 cuotas', '5 cuotas', '6 cuotas', '8 cuotas', '11 cuotas'], filas, [5.4, 2.2, 2.1, 2.1, 2.1, 2.1, 2.1], size=8.5)
t('Súper intensivo: hasta la columna de 5. Intensivo: hasta la de 11. Ejemplo: beca de alianza, inicial 20%, intensivo a 11 cuotas = $1.113.088 hoy y 11 cuotas de $452.348.', size=8.5, color=GRIS)

sec('4. La frase de cierre')
caja('"Tu plan queda así: hoy separas con $300.000, completas la inicial de [inicial] antes del [fecha, 3 días], y desde el [fecha, 30 días] pagas [N] cuotas de [valor]. Vas a la sede, firmas con coordinación y quedas en el grupo. ¿Lo dejamos para mañana en la mañana o en la tarde?" Y silencio. Si pide bajar la cuota: más inicial o más cuotas hasta el máximo, nunca bajar el precio.')
t('Lo que cuenta para los cortes: inicial pagada y contrato firmado con coordinación. Una separación de $300.000 sola no cuenta.', size=8.5, color=GRIS, after=0)

try: d.save(OUT); print('OK', OUT)
except PermissionError: d.save(OUT.replace('.docx', '_v2.docx')); print('ABIERTO -> _v2')
