# -*- coding: utf-8 -*-
"""Excel de registro del programa de referidos de asesores de entidades aliadas (una sola tabla, no una
pestana por asesora). Tres hojas:
  ASESORAS  una fila por codigo: entidad, oficina, nombre, celular, QR entregado, autorizacion de datos firmada.
  REGISTRO  una fila por mensaje con codigo que llega a la linea de alianzas; la llena coordinacion.
            Columna "Bono procede" se calcula sola (sin gestion humana previa, matricula pagada en <=30 dias
            del mensaje, pago recibido y retracto cumplido).
  RESUMEN   por codigo: mensajes, citas, matriculas, bonos y ranking para el premio unico. Formulas, no se toca.
Vive en la carpeta del gerente de alianzas / plan referidos (fuera de git: lleva celulares). Solo crea el archivo
si no existe, para no pisar datos ya digitados. Uso: python generadores/comercial/gen_registro_referidos_asesoras.py"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu'
CARPETA = ROOT + r'\recursos\comercial\INDUCCION_COMERCIAL\contratacion comercial\LINKEDIN DIANA GOMEZ\Candidatos\John Hinestroza\plan referidos'
OUT = os.path.join(CARPETA, 'REGISTRO_REFERIDOS_ASESORAS.xlsx')
LISTA = os.path.join(CARPETA, 'codigos_asesoras.txt')
BONO = 200000
N = 300  # filas preparadas en REGISTRO

codigos = [l.strip() for l in open(LISTA, encoding='utf-8') if l.strip() and not l.startswith('#')] if os.path.exists(LISTA) else ['C-%02d' % i for i in range(1, 11)]
HDR = PatternFill('solid', fgColor='E54C09'); HF = Font(bold=True, color='FFFFFF', size=10)
CALC = PatternFill('solid', fgColor='FFF8E7'); thin = Side(style='thin', color='BBBBBB'); BORDE = Border(left=thin, right=thin, top=thin, bottom=thin)
SINO = DataValidation(type='list', formula1='"SÍ,NO"', allow_blank=True)

def cabecera(ws, cols, anchos):
    for i, (c, w) in enumerate(zip(cols, anchos), 1):
        x = ws.cell(row=1, column=i, value=c); x.fill = HDR; x.font = HF; x.alignment = Alignment(wrap_text=True, vertical='center'); x.border = BORDE
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 42; ws.freeze_panes = 'A2'

wb = Workbook()

# ---- ASESORAS
wa = wb.active; wa.title = 'ASESORAS'
cabecera(wa, ['Código', 'Entidad', 'Oficina', 'Nombre de la asesora', 'Celular', 'Correo', 'Fecha entrega QR',
              'Autorización de datos firmada', 'Dirección de la oficina enterada', 'Observaciones'],
         [9, 18, 18, 28, 14, 26, 14, 16, 16, 30])
for r, c in enumerate(codigos, 2):
    wa.cell(row=r, column=1, value=c).font = Font(bold=True)
    for col in range(1, 11): wa.cell(row=r, column=col).border = BORDE
    wa.cell(row=r, column=7).number_format = 'DD/MM/YYYY'
wa.add_data_validation(SINO); SINO.add('H2:I%d' % (len(codigos) + 1))

# ---- REGISTRO
wr = wb.create_sheet('REGISTRO')
cols = ['#', 'Fecha del mensaje', 'Hora', 'Código', 'Celular del afiliado', 'Nombre del afiliado',
        'Gestión humana nuestra en los 30 días previos', 'Contactado (fecha)', 'Cita (fecha)', 'Quién cerró',
        'Matrícula pagada (fecha)', 'Programa', 'Valor pagado', 'Pago recibido por Heiiu (fecha)', 'Retracto cumplido',
        'Bono procede (se calcula)', 'Bono recargado (fecha)', 'Observaciones']
cabecera(wr, cols, [5, 13, 8, 9, 15, 26, 16, 13, 13, 16, 13, 16, 13, 14, 11, 12, 13, 30])
sv = DataValidation(type='list', formula1='"SÍ,NO"', allow_blank=True); wr.add_data_validation(sv)
cv = DataValidation(type='list', formula1='=ASESORAS!$A$2:$A$%d' % (len(codigos) + 1), allow_blank=True); wr.add_data_validation(cv)
for r in range(2, N + 2):
    wr.cell(row=r, column=1, value='=IF(B%d="","",ROW()-1)' % r)
    for col in (2, 8, 9, 11, 14, 17): wr.cell(row=r, column=col).number_format = 'DD/MM/YYYY'
    wr.cell(row=r, column=3).number_format = 'HH:MM'
    wr.cell(row=r, column=13).number_format = '#,##0'
    f = ('=IF(B{r}="","",IF(AND(G{r}="NO",K{r}<>"",K{r}-B{r}<=30,N{r}<>"",O{r}="SÍ"),"SÍ",'
         'IF(OR(G{r}="SÍ",AND(K{r}<>"",K{r}-B{r}>30)),"NO","")))').format(r=r)
    x = wr.cell(row=r, column=16, value=f); x.fill = CALC; x.font = Font(bold=True)
    for col in range(1, 19): wr.cell(row=r, column=col).border = BORDE
sv.add('G2:G%d' % (N + 1)); sv.add('O2:O%d' % (N + 1)); cv.add('D2:D%d' % (N + 1))

# ---- RESUMEN
ws = wb.create_sheet('RESUMEN')
cabecera(ws, ['Código', 'Entidad', 'Oficina', 'Mensajes con código', 'Citas', 'Matrículas pagadas', 'Bonos que proceden',
              'Bonos recargados', 'Total bonos $', 'Puesto (premio único)'], [9, 18, 18, 12, 9, 12, 12, 12, 14, 12])
last = len(codigos) + 1
for r, c in enumerate(codigos, 2):
    ws.cell(row=r, column=1, value=c).font = Font(bold=True)
    ws.cell(row=r, column=2, value='=IF(ASESORAS!B%d="","",ASESORAS!B%d)' % (r, r))
    ws.cell(row=r, column=3, value='=IF(ASESORAS!C%d="","",ASESORAS!C%d)' % (r, r))
    ws.cell(row=r, column=4, value='=COUNTIF(REGISTRO!$D:$D,A%d)' % r)
    ws.cell(row=r, column=5, value='=COUNTIFS(REGISTRO!$D:$D,A%d,REGISTRO!$I:$I,"<>")' % r)
    ws.cell(row=r, column=6, value='=COUNTIFS(REGISTRO!$D:$D,A%d,REGISTRO!$K:$K,"<>")' % r)
    ws.cell(row=r, column=7, value='=COUNTIFS(REGISTRO!$D:$D,A%d,REGISTRO!$P:$P,"SÍ")' % r)
    ws.cell(row=r, column=8, value='=COUNTIFS(REGISTRO!$D:$D,A%d,REGISTRO!$Q:$Q,"<>")' % r)
    ws.cell(row=r, column=9, value='=H%d*%d' % (r, BONO)); ws.cell(row=r, column=9).number_format = '#,##0'
    ws.cell(row=r, column=10, value='=IF(G%d=0,"",RANK(G%d,$G$2:$G$%d))' % (r, r, last))
    for col in range(1, 11): ws.cell(row=r, column=col).border = BORDE; ws.cell(row=r, column=col).fill = CALC
t = last + 1
ws.cell(row=t, column=1, value='TOTAL').font = Font(bold=True)
for col in range(4, 10):
    L = get_column_letter(col); x = ws.cell(row=t, column=col, value='=SUM(%s2:%s%d)' % (L, L, last)); x.font = Font(bold=True)
ws.cell(row=t, column=9).number_format = '#,##0'
ws.cell(row=t + 2, column=1, value='Bono: $%s por matrícula pagada de contado. Ranking = bonos que proceden; en empate, gana quien llegó primero (ver fechas en REGISTRO). Piloto hasta el 31/01/2027.' % format(BONO, ',').replace(',', '.')).font = Font(italic=True, color='666666')

if os.path.exists(OUT):
    print('YA EXISTE, no se pisa:', OUT)
else:
    wb.save(OUT); print('OK', OUT)
