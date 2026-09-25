# -*- coding: utf-8 -*-
"""Documento institucional para el onboarding de LaHaus: quienes somos, mision, vision, experiencia,
funcionamiento, licencias, garantia, niveles. Solo informacion de cara al cliente."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NARANJA = RGBColor(0xE5, 0x4C, 0x09); GRIS = RGBColor(0x66, 0x66, 0x66)
doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.6); s.bottom_margin = Cm(1.4); s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)
doc.styles['Normal'].font.name = 'Calibri'; doc.styles['Normal'].font.size = Pt(10.5)

def shd(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear'); el.set(qn('w:color'), 'auto'); el.set(qn('w:fill'), color); tcPr.append(el)

def sec(txt):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(3)
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(12.5); r.font.color.rgb = NARANJA

def t(txt, bold=False, italic=False):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(txt); r.font.size = Pt(10.5); r.bold = bold; r.italic = italic

def b(txt):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2); p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run('· ' + txt); r.font.size = Pt(10.5)

def tabla(filas, anchos):
    tb = doc.add_table(rows=len(filas), cols=len(filas[0])); tb.style = 'Table Grid'
    for i, fila in enumerate(filas):
        for j, v in enumerate(fila):
            c = tb.cell(i, j); c.width = Cm(anchos[j]); c.paragraphs[0].paragraph_format.space_after = Pt(0)
            r = c.paragraphs[0].add_run(v); r.font.size = Pt(9.5)
            if i == 0: r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shd(c, 'E54C09')
            elif j == 0: r.bold = True; shd(c, 'FFF8E7')

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
r = p.add_run('HEIIU ENGLISH ACADEMY — INFORMACIÓN INSTITUCIONAL'); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = NARANJA
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(8)
r = p.add_run('Global Teacher S.A.S. · NIT 900.422.478-2 · Bucaramanga · Documento para el asistente de atención · Septiembre 2026'); r.font.size = Pt(9); r.font.color.rgb = GRIS

sec('1. QUIÉNES SOMOS')
t('Heiiu English Academy es la marca comercial de Global Teacher S.A.S., una institución de educación para el trabajo y el desarrollo humano con sede en Bucaramanga. Somos una academia de inglés PRESENCIAL: creemos que el inglés se entrena hablando, en clases activas, con un profesor al frente y un grupo pequeño alrededor.')
t('El nombre viene de "hello you": hola tú. Refleja cómo tratamos a cada estudiante — por su nombre, no como un número. Grupos de máximo 16 personas, seguimiento individual y un compromiso escrito con el resultado.')
t('Nuestra historia: nacimos en 2017 como una fundación en Barro Blanco, Piedecuesta, con un deseo genuino: que las personas con menos privilegios tuvieran las oportunidades que otros reciben casi por defecto. Para poder sostener el programa, en octubre de 2018 adquirimos Global Teacher S.A.S., institución con licencia desde 2012, y desde entonces operamos como academia formal sin perder el propósito con el que empezamos. En el camino diseñamos una metodología propia basada en la repetición, el trabajo en pares y el apoyo entre estudiantes.')
t('No enseñamos inglés, transformamos futuros: pertenecer a la academia significa que cada persona que se conecta con nosotros encuentra una mano para cumplir sus metas.', italic=True)
b('Sede: Carrera 27 # 48-49, segundo piso, barrio Sotomayor, Bucaramanga.')
b('Instagram: @Heiiu_english.')

sec('2. MISIÓN')
t('Conectar a nuestros estudiantes con oportunidades laborales a nivel nacional e internacional por medio de la enseñanza del inglés como segundo idioma, y mejorar e impactar sus vidas.', italic=True)

sec('3. VISIÓN')
t('Expandirnos a nivel nacional con más programas especializados, manteniendo lo que nos define: resultados medibles y un compromiso escrito con cada estudiante.', italic=True)

sec('3B. VALORES')
b('Honestidad.')
b('Respeto por el cliente y por su inversión.')
b('Responsabilidad y compromiso con la calidad y la excelencia.')
b('Integridad: cumplimos lo que prometemos.')
b('Responsabilidad social: queremos ayudar de verdad, y vemos el futuro de nuestros estudiantes reflejado en nuestro trabajo.')
b('Inteligencia emocional: somos profesionales y adultos.')
b('Pasión.')

sec('4. EXPERIENCIA')
b('Equipo formando estudiantes desde 2017; institución con licencia de funcionamiento desde 2012.')
b('Programas de inglés de nivel A1 a B2 alineados con el Marco Común Europeo de Referencia (MCER).')
b('Convenio con Cajasan: la única academia de inglés de la ciudad con este convenio para sus afiliados. Las condiciones del convenio se explican en la asesoría personalizada.')
b('Alianzas con empresas de la región para la formación de sus equipos, y programa para colegios.')
b('Lo que nos diferencia: buscamos activamente conectar a nuestros estudiantes con oportunidades que les cambien los ingresos y la vida. Tenemos alianzas con programas internacionales como Au Pair, Work and Travel y campamentos de verano en Estados Unidos, además de bolsas de empleo en el exterior. El inglés es el requisito de entrada a todos ellos; la academia orienta y conecta, no garantiza plazas.')

sec('5. CÓMO FUNCIONA EL MÉTODO HEIIU')
t('No vendemos clases de inglés: entrenamos. El estudiante llega a un programa con estructura, medición y compromiso mutuo.', bold=True)
b('Clases inmersivas y presenciales: el estudiante habla en inglés desde el primer día. La clase se conduce en inglés y el profesor guía, corrige y hace practicar en voz alta.')
b('Grupos pequeños: máximo 16 estudiantes, para que cada uno hable en cada clase.')
b('Práctica diaria fuera del aula: el estudiante graba un video corto en inglés todos los días y lo envía a la academia. Es su portafolio de avance y el hábito que fija el idioma.')
b('Avance medido: sesión de ubicación gratuita antes de decidir (prueba en línea EF SET + entrevista oral corta; el nivel lo define la academia), evaluaciones por nivel, examen final aplicado por un evaluador externo, y certificado oficial por cada nivel aprobado.')
b('Simulaciones reales: entrevistas, presentaciones, llamadas, negociaciones — el inglés que se usa en la vida y en el trabajo, no solo gramática.')
b('Formación del carácter: cada nivel trabaja también disciplina, constancia y comunicación, porque son las habilidades que hacen que el inglés se sostenga.')
t('Jornadas según nivel y grupo vigente: mañana (8:00 a 12:00), tarde (2:00 a 6:00), franjas de 2 horas entre semana (por ejemplo 6:30 a 8:30 PM) y modalidad sabatina (8:00 AM a 12:00 M). Formatos intensivo (10 horas por semana) y súper intensivo (20 horas por semana). El horario exacto se confirma en la asesoría.')
t('Edades: programas entre semana desde los 17 años; modalidad sabatina desde los 12 años (el contrato del menor lo firma su acudiente).')

sec('6. LICENCIAS Y CERTIFICACIONES')
b('Institución de educación para el trabajo y el desarrollo humano con licencia de funcionamiento de la Secretaría de Educación desde 2012.')
b('Programas de inglés registrados ante la Secretaría de Educación: Resoluciones 1581-1131 y 2253-2025.')
b('Certificación de calidad ICONTEC: norma NTC 5555 (calidad para instituciones de formación para el trabajo) e ISO 9001 (sistema de gestión de calidad).')
b('Certificados oficiales por nivel, con validez para procesos laborales y académicos.')

sec('7. LA GARANTÍA DE APRENDIZAJE')
t('Somos la única academia de la ciudad con garantía de aprendizaje por escrito, dentro del contrato. Se cita únicamente con este texto:', bold=True)
t('"Si cumples — asistencia, tareas, evaluaciones — y no avanzas, te devolvemos el 100% del nivel. Por escrito, en el contrato."', italic=True)
t('La garantía es un trato en dos direcciones: la academia pone el método, el profesor y el respaldo; el estudiante pone el trabajo — asistencia, tareas y su práctica diaria. Por eso no prometemos que sea fácil ni rápido: prometemos que funciona si se hace el trabajo.')
t('Nuestra exigencia: somos estrictos en asistencia y puntualidad. El resultado se logra en equipo, en clase y en casa, y por eso el listening en casa es fundamental. El enfoque es conversacional, pero la gramática no se deja de lado porque es la base para comunicarse bien. Aprender un idioma exige repetir mil y una veces, y el estudiante que lo entiende desde el inicio es el que llega.')

sec('8. LOS NIVELES')
tabla([
    ['Nivel', 'Horas', 'Qué logra el estudiante'],
    ['A1 — Fundamentos', '90 h', 'Presentarse, hablar de su rutina, su familia y su trabajo; números, hora y precios; lugares de la ciudad; presente, pasado simple y planes. Cierra con la presentación oral MY WORLD.'],
    ['A2 — Vida diaria y trabajo', '110 h', 'Futuro, condicionales, presente perfecto, voz pasiva, comparativos, phrasal verbs. Situaciones: entrevista de trabajo, hotel, tienda, restaurante, servicio al cliente, reunión de equipo, primer día en un empleo. Cierra con la presentación oral MY LIFE (7-10 min).'],
    ['B1 — Comunicación con fluidez', '175 h', 'Gramática consolidada en contexto profesional más tercer condicional, deducción, discurso indirecto y tiempos perfectos. Situaciones: briefings, entrevistas, onboarding, planeación de proyectos, evaluaciones de desempeño, negociación con proveedores. Presentación intermedia MY STORY, MY GOALS, taller de pitch y presentaciones finales.'],
    ['B2 — Nivel profesional', '200 h', 'Primera mitad, vida real: restaurante, direcciones, compras, médico, arriendo, banco, aeropuerto, desacuerdos. Segunda mitad, mundo profesional: networking, CV en 3 minutos, entrevistas, reuniones, presentar una idea, cliente molesto, negociación, feedback, pitch de 90 segundos, pedir un aumento. Cierra con un simulacro de una jornada completa de trabajo en inglés.'],
], [4.6, 1.8, 10.6])
t('Total del programa completo A1 a B2: 575 horas presenciales. Cada nivel termina con examen final y certificado oficial. Un estudiante puede entrar en el nivel que le corresponda según su prueba de ubicación.')
t('Programa de lanzamiento vigente: ARRANQUE A1+A2 — los dos primeros niveles completos (200 horas), libros y certificados incluidos, cohorte que inicia el 5 de octubre de 2026.', bold=True)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(10)
r = p.add_run('Heiiu English Academy · Global Teacher S.A.S. · Bucaramanga'); r.font.size = Pt(8.5); r.font.color.rgb = GRIS

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\plataforma LaHaus\INFO_INSTITUCIONAL_HEIIU.docx'
doc.save(out); print('OK', out)
