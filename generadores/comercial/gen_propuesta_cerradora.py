# -*- coding: utf-8 -*-
"""Propuesta para cerradora high ticket — 1 pagina para compartir pantalla."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NARANJA = RGBColor(0xE5, 0x4C, 0x09)
GRIS = RGBColor(0x77, 0x77, 0x77)
CREMA = 'FFF8E7'
NARANJA_HEX = 'E54C09'

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.1); s.bottom_margin = Cm(0.9); s.left_margin = Cm(1.8); s.right_margin = Cm(1.8)
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(9.5)

def shd(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement('w:shd'); el.set(qn('w:val'), 'clear'); el.set(qn('w:fill'), color)
    tcPr.append(el)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('HEIIU ENGLISH ACADEMY — PROPUESTA COMERCIAL · CERRADORA HIGH TICKET'); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = NARANJA
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Septiembre 2026 · Confidencial'); r.font.size = Pt(9.5); r.font.color.rgb = GRIS

def sec(txt):
    p = doc.add_paragraph(); r = p.add_run(txt); r.bold = True; r.font.size = Pt(11.5); r.font.color.rgb = NARANJA
    p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(1)

def t(txt, bold=False):
    p = doc.add_paragraph(); r = p.add_run(txt); r.font.size = Pt(9.5); r.bold = bold
    p.paragraph_format.space_after = Pt(1.5)

def tab(filas):
    tb = doc.add_table(rows=len(filas), cols=len(filas[0])); tb.style = 'Table Grid'
    for i, fila in enumerate(filas):
        for j, val in enumerate(fila):
            c = tb.rows[i].cells[j]
            pp = c.paragraphs[0]; rr = pp.add_run(str(val)); rr.font.size = Pt(9.5)
            if i == 0:
                rr.bold = True; rr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shd(c, NARANJA_HEX)
    esp = doc.add_paragraph(); er = esp.add_run(''); er.font.size = Pt(2)
    esp.paragraph_format.space_after = Pt(0)

def frase(txt):
    tab = doc.add_table(rows=1, cols=1); tab.style = 'Table Grid'
    c = tab.rows[0].cells[0]; shd(c, CREMA)
    p = c.paragraphs[0]; r = p.add_run(txt); r.font.size = Pt(9.5); r.bold = True
    esp = doc.add_paragraph(); esp_r = esp.add_run(''); esp_r.font.size = Pt(3)
    esp.paragraph_format.space_after = Pt(0)

sec('EL PRODUCTO')
t('ARRANQUE A1+A2: los dos primeros niveles de inglés COMPLETOS — 200 horas presenciales, libros y certificado oficial por nivel — $2.990.000, precio de lanzamiento. Cohorte inicia el 5 DE OCTUBRE.', bold=True)
t('Y el high ticket detrás: el PROGRAMA COMPLETO (A1 a B2, 575 horas) con beca del Fondo según el perfil del cliente — caso típico (perfil General, de contado): $5.913.280. Perfiles de menores ingresos pueden acceder a becas mayores, con soportes y cupos. El cierre doble ofrece los dos caminos: las dos respuestas son sí.')

sec('POR QUÉ CIERRA (los diferenciales)')
t('· GARANTÍA DE APRENDIZAJE POR ESCRITO en el contrato — única en la ciudad: mata la objeción #1.')
t('· MÓDULO DE GRADUACIÓN a elección: sale con su PÁGINA WEB publicada y su primer mensaje de venta en inglés — o con su aplicación lista para irse (video + entrevista ensayada).')
t('· Institución con licencia desde 2012, programas registrados, certificación ICONTEC.')
t('· Precio de lanzamiento + contador de cupos REAL: la urgencia no es teatro.')

sec('LA META')
t('ABRIR LAS DOS JORNADAS — mañana (8:00-12:00) y noche (6:30-8:30) — con grupos de 13-14 estudiantes (máx. 16). Los primeros 20 cupos van a precio de lanzamiento; la cohorte inicia el 5 de octubre.', bold=True)

sec('EL TRATO')
t('· Comisión: 8% del recaudo por TODA venta de tu canal (tu contenido, tu red, y la pauta financiada). Pago contra plata efectiva en caja.', bold=True)
t('· ACELERADOR — dos puertas al 10% retroactivo sobre TODA tu serie: (a) el 30 DE SEPTIEMBRE ambas jornadas abiertas con 16 matrículas en total (mínimo 8 y 8), o (b) los 20 cupos completos antes del 5 de octubre. Por debajo: 8%.', bold=True)
t('· BONO GRUPOS LLENOS: 28 matrículas (14 + 14) al 30 de septiembre = $2.000.000 de bono adicional, encima del 10%.', bold=True)
t('· PAUTA FINANCIADA POR HEIIU sobre tus piezas, en tus canales: apruebas pieza + presupuesto con gerencia antes de publicar; tope inicial $1.000.000/mes, revisable por resultados; factura de la plataforma como soporte del gasto.')
t('· Sin exclusividad, sin horarios, sin básico. Leads que la casa te entregue: 5%.')

sec('LO QUE GANAS TÚ (con el acelerador logrado — 10%)')
tab([
 ['Escenario', 'Recaudo aprox.', 'Tu comisión'],
 ['20 Arranques', '$59.8M', '≈ $6.0M'],
 ['10 Arranques + 10 programas completos', '$89.0M', '≈ $8.9M'],
 ['Grupos llenos: 28 + bono', '$87.9M', '≈ $10.8M (incluye el bono)'],
])

sec('LA REGLA DEL CUPO 21 (automática, sin permisos)')
t('El contador es compartido y gerencia lo actualiza a diario. Los 20 cupos de lanzamiento son a $2.990.000. Del cupo 21 en adelante se sigue vendiendo ESTA MISMA cohorte (grupos hasta de 16) a la tarifa plena de los dos niveles: $3.511.000. ÚNICA excepción, escrita: pago de CONTADO completo el mismo día = se mantiene el precio de lanzamiento ($2.990.000). A cuotas: tarifa plena. Todo comisiona igual — y la urgencia gana un segundo filo: agotados los 20, el precio de lanzamiento solo existe pagando hoy de contado. Lo único prohibido: ofrecer el precio de lanzamiento en esta cohorte cuando el contador ya va en 20.')

t('REGLAS INNEGOCIABLES: sin descuentos inventados · sin promesas de empleo/visa · se dice beca del Fondo (jamás patrocinio) · toda venta se registra en planilla para comisionar.', bold=True)

sec('EL PROCESO')
t('1. ¿Aceptas? Firmamos MAÑANA 15 DE SEPTIEMBRE a primera hora (contrato de prestación con confidencialidad). 2. Firmado = manual de venta completo + planilla + piezas EN ESE MISMO MOMENTO. 3. Mañana mismo grabas tus primeros videos y montamos la primera campaña. El reloj del 30 de septiembre corre desde ya — cada día vale.', bold=True)

doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Heiiu English Academy · Carrera 27 # 48-49, Bucaramanga · Gerencia'); r.font.size = Pt(9); r.font.color.rgb = GRIS

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\INDUCCION_COMERCIAL\PROPUESTA_CERRADORA_HIGH_TICKET.docx'
doc.save(out); print('OK', out)
