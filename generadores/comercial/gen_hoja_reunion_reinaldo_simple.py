# -*- coding: utf-8 -*-
"""Hoja de reunión SIMPLE con Reinaldo (filtro 2) para gerencia: 1 página, conversación, sin entrevista. Copia de la de John con las cifras de su HV.
La ficha (FICHA_CIFRAS_CANDIDATO_ALIANZAS) recoge las cifras; la entrevista fuerte la hace el socio financiero (filtro 4).
Salida: Candidatos/HOJA_REUNION_REINALDO_2026-10-15.docx (reemplaza la de 3 páginas; la vieja queda como guía del segundo evaluador)."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAR = RGBColor(0xE5, 0x4C, 0x09); GRIS = RGBColor(0x66, 0x66, 0x66)
doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.3); s.bottom_margin = Cm(1.0); s.left_margin = Cm(1.8); s.right_margin = Cm(1.8)
doc.styles['Normal'].font.name = 'Calibri'; doc.styles['Normal'].font.size = Pt(10.5)
def shd(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear'); el.set(qn('w:color'), 'auto'); el.set(qn('w:fill'), color); tcPr.append(el)
def sec(txt):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = NAR
def t(txt, size=10.5, color=None, after=3, bold=False):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(txt); r.font.size = Pt(size); r.bold = bold
    if color: r.font.color.rgb = color
def caja(txt):
    tb = doc.add_table(rows=1, cols=1); tb.style = 'Table Grid'; c = tb.cell(0, 0); shd(c, 'FFF8E7')
    c.paragraphs[0].paragraph_format.space_after = Pt(0)
    r = c.paragraphs[0].add_run(txt); r.font.size = Pt(10.5); r.italic = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
r = p.add_run('REUNIÓN CON REINALDO — UNA CONVERSACIÓN, NO UNA ENTREVISTA'); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = NAR
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(6)
r = p.add_run('Jueves 15/10/2026 · 45 min · Uso interno de gerencia · Llevas: esta hoja, la FICHA DE TRAYECTORIA impresa y un esfero'); r.font.size = Pt(9); r.font.color.rgb = GRIS
t('Hoy tú no calificas ni repreguntas. La ficha recoge las cifras; la entrevista fuerte la hace Carlos la próxima semana con la hoja de números en la mano. Tu trabajo es escuchar y cerrar con los tres pasos.', bold=True, after=6)

sec('1. Llega (0-12 min): la ficha')
caja('"Bienvenido. Antes de conversar, llena esta ficha con calma, son 10 minutos; voy por un café y volvemos."')
t('Ahí quedan escritas sus cifras, su aspiración, su fecha de inicio y la firma de autorización de datos. No preguntas nada de eso después.', size=10, color=GRIS)

sec('2. La causa (12-17 min): hablas tú, 5 minutos')
caja('"Heiiu es una academia santandereana e independiente con una causa doble. Educativa: cada alianza que abramos forma a jóvenes y trabajadores de Santander en inglés, con valores y disciplina para emprender y liderar. Empresarial: las empresas locales no tienen que desaparecer frente al músculo financiero de las grandes cadenas. Busco a la persona que quiera dar esa pelea con nosotros."')
t('Sin nombrar competidores. Nada más.', size=10, color=GRIS)

sec('3. Una sola pregunta y escuchas (17-32 min)')
caja('"Cuéntame BBVA de principio a fin: cómo llegaste, qué hiciste y por qué saliste en mayo."   ·   Si no lo dice solo: "¿Y qué has hecho desde mayo?"   ·   Si se queda corto: "¿Y entre Bancolombia y BBVA, 2018 a 2021?"')
t('No interrumpes, no calificas, no explicas. Dos huecos de la HV que tienen que quedar contados sin que los preguntes de frente: junio 2018 a junio 2021 y mayo 2025 a hoy. Anotas aquí lo que te suene raro:', size=10, color=GRIS)
for _ in range(3): t('______________________________________________________________________________________________', after=2)

sec('4. Dos cosas que tienes que oír antes de que se vaya (32-35 min)')
t('Salen de la ficha. Las confirmas en una frase, mirando lo que escribió:')
caja('"Vi que vives en Bucaramanga y que puedes empezar en [fecha que puso], ¿cierto?"   ·   "Y con el piso de los tres primeros meses, ¿puedes arrancar?"')
t('☐ Vive aquí y puede venir todos los días      ☐ Puede arrancar con el piso      ☐ Fecha de inicio: ____________', after=4)

sec('5. Sus dudas (35-40 min)')
t('Viene de banca: lo más probable es que la ficha traiga una aspiración alta. No la comentas. Si pregunta por plata: "El básico es el mínimo legal con prestaciones, igual para todos. Lo que protege los primeros meses es el ingreso mínimo de los meses 1 a 3, y de ahí en adelante la comisión no tiene techo; los escenarios que te mandé son ilustrativos, no una promesa." No se sube el básico, no se bajan metas, no se promete el 130 ni el 150%. Nada de oferta hoy.', size=10)

sec('5-bis. Tu única pregunta de fondo (40-42 min)')
caja('"Con todo lo que ya sabes de nosotros: ¿hay algo que te haría decir que no?"   → y te callas.')
t('Lo que diga es su objeción real. No la respondes hoy: "Entiendo. Lo tengo en cuenta." (si es plata, solo el guion del paso 5). La anotas y se la pasas a Carlos. Si dice "nada": "Perfecto, entonces el siguiente paso te lo mando hoy."', size=10, color=GRIS)
t('Objeción real: ___________________________________________________________________________', after=4)

sec('6. Cierre (42-45 min): tal cual')
caja('"El proceso tiene tres pasos: esta conversación, una tarea práctica de 48 horas que te mando hoy, y una reunión con Carlos, nuestro socio financiero, que revisa los números de todo cargo comercial. Con eso decidimos. Te escribo esta tarde con la tarea."')
t('Te quedas con la ficha firmada. La fecha con Carlos se le da cuando entregue la tarea, no hoy.', size=10, color=GRIS)

sec('Después de la reunión (10 minutos, con la ficha en la mano)')
t('1. Le envías la MISMA tarea que a John (copia el mensaje que le mandaste el 14/10): 5 empresas de 40+ empleados con decisor, mensaje de apertura de máximo 120 palabras, plan de 2 semanas, celulares de referencias; entrega en 48 horas.   2. Referencias obligatorias para tu llamada: su jefe en BBVA (coordinador nacional vendor) y el coordinador regional de Bancolombia; los dos teléfonos están en su HV.   3. Le pasas a Carlos la ficha, la tarea y una hoja de números como la de John (se arma cuando entregue la tarea).   4. Me cuentas en 5 líneas: qué dijo de la salida de BBVA y de los dos huecos, qué puso de aspiración y fecha, y qué te sonó raro.', size=10)

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\INDUCCION_COMERCIAL\contratacion comercial\LINKEDIN DIANA GOMEZ\Candidatos\HOJA_REUNION_REINALDO_2026-10-15.docx'
try: doc.save(out); print('OK', out)
except PermissionError: doc.save(out.replace('.docx', '_v2.docx')); print('ABIERTO -> _v2')
