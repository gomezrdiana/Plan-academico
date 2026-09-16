# -*- coding: utf-8 -*-
"""Documento de inicio de la cerradora en formato de marca (naranja/crema)."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NARANJA = RGBColor(0xE5, 0x4C, 0x09)
GRIS = RGBColor(0x77, 0x77, 0x77)
CREMA = 'FFF8E7'
NARANJA_HEX = 'E54C09'

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.2); s.bottom_margin = Cm(1.0); s.left_margin = Cm(1.8); s.right_margin = Cm(1.8)
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(10)

def shd(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement('w:shd'); el.set(qn('w:val'), 'clear'); el.set(qn('w:fill'), color)
    tcPr.append(el)

def sec(txt):
    p = doc.add_paragraph(); r = p.add_run(txt); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = NARANJA
    p.paragraph_format.space_before = Pt(7); p.paragraph_format.space_after = Pt(2)

def t(txt, bold=False, size=10):
    p = doc.add_paragraph(); r = p.add_run(txt); r.font.size = Pt(size); r.bold = bold
    p.paragraph_format.space_after = Pt(2)
    return p

def b(txt, bold_first=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4); p.paragraph_format.space_after = Pt(2)
    r0 = p.add_run('· '); r0.bold = True; r0.font.color.rgb = NARANJA
    if bold_first:
        r1 = p.add_run(bold_first); r1.bold = True; r1.font.size = Pt(10)
    r2 = p.add_run(txt); r2.font.size = Pt(10)

def tabla(filas, anchos=None):
    tb = doc.add_table(rows=len(filas), cols=len(filas[0])); tb.style = 'Table Grid'
    for i, fila in enumerate(filas):
        for j, val in enumerate(fila):
            c = tb.rows[i].cells[j]
            pp = c.paragraphs[0]; rr = pp.add_run(str(val)); rr.font.size = Pt(10)
            if i == 0:
                rr.bold = True; rr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shd(c, NARANJA_HEX)
            elif j >= 1:
                rr.bold = True
    esp = doc.add_paragraph(); er = esp.add_run(''); er.font.size = Pt(2)
    esp.paragraph_format.space_after = Pt(0)

def caja(txt, italic=True):
    tb = doc.add_table(rows=1, cols=1); tb.style = 'Table Grid'
    c = tb.rows[0].cells[0]; shd(c, CREMA)
    p = c.paragraphs[0]; r = p.add_run(txt); r.italic = italic; r.font.size = Pt(10); r.bold = True
    esp = doc.add_paragraph(); er = esp.add_run(''); er.font.size = Pt(3)
    esp.paragraph_format.space_after = Pt(0)

# ===== PORTADILLA =====
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('ARRANQUE DE LA CERRADORA'); r.bold = True; r.font.size = Pt(17); r.font.color.rgb = NARANJA
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Qué vendemos · La meta · Tu oferta comisional — Heiiu English Academy · Septiembre 2026'); r.font.size = Pt(9.5); r.font.color.rgb = GRIS

sec('1. LO QUE VAS A VENDER')
t('ARRANQUE A1+A2: los dos primeros niveles de inglés COMPLETOS — 200 horas presenciales, libros y certificado oficial por nivel — $2.990.000, precio de lanzamiento. La cohorte inicia el 5 DE OCTUBRE, en dos jornadas: mañana (8-12) y noche (6:30-8:30).', bold=True)
t('Y detrás, el high ticket: el programa completo A1→B2 (575 horas) con beca del Fondo según el perfil. Ejemplo real: el programa completo vale $8.696.000; un cliente perfil General que paga de contado recibe beca del Fondo del 32% y queda en $5.913.280 — se ahorra más de $2.7 millones comprometiéndose hoy. El cierre doble ofrece los dos caminos: las dos respuestas son sí.')
t('Por qué este producto CIERRA (tus armas):', bold=True)
b('en el contrato — mata la objeción #1.', bold_first='La única GARANTÍA DE APRENDIZAJE POR ESCRITO de la ciudad, ')
b('página web publicada + primer mensaje de venta en inglés, o su POSTULACIÓN a un programa en el exterior (Au Pair, Work and Travel) diligenciada, con video de presentación y entrevista ensayada. Si el cliente no quiere ninguno: no pasa nada — es un bono opcional, cursa sus niveles igual, mismo precio.', bold_first='MÓDULO DE GRADUACIÓN a elección, estrenándose con esta cohorte: ')
b('Institución con licencia desde 2012, programas registrados, certificación ICONTEC. Contador REAL de cupos: la urgencia no es teatro.')
b('abono de $300.000 hoy congela cupo y precio; el desembolso o el cheque llegan en máximo 10 días hábiles.', bold_first='El crédito de financiera o cooperativa — y el cheque de CESANTÍAS — cuentan como CONTADO: ')
b('quien solo puede estudiar los sábados también se matricula, a las tarifas de la matriz vigente. Ningún cliente se va por el horario.', bold_first='MODALIDAD SABATINA: ')
b('la sesión de ubicación define su nivel, y se le vende el paquete DESDE SU NIVEL hasta B2 con beca del Fondo (A2→B2 $7.100.000 full · B1→B2 $5.185.000 full — la tarjeta de becas trae la matriz por perfil), o su nivel individual. Entra al grupo vigente de su nivel o separa cupo para la próxima apertura. Comisiona igual por canal — tickets a menudo mayores que un Arranque. Ningún cliente se va por saber inglés.', bold_first='EL QUE YA TIENE NIVEL: ')

sec('2. LA META')
t('Abrir las DOS jornadas el 5 de octubre con los 20 cupos de lanzamiento vendidos — grupos de 13-14, máximo 16. Del cupo 21 en adelante se sigue vendiendo la misma cohorte a tarifa plena ($3.511.000), con una sola excepción escrita: contado completo el mismo día mantiene el precio de lanzamiento.', bold=True)

sec('3. LO QUE PUEDES GANAR')
t('Cuánto te deja cada venta:', bold=True)
tabla([
 ['Venta', 'Al 8%', 'Al 10%'],
 ['Un Arranque ($2.990.000)', '$239.200', '$299.000'],
 ['Un programa completo (típico $5.913.280)', '$473.062', '$591.328'],
])
t('Y los escenarios completos (con el 10% logrado):', bold=True)
tabla([
 ['Escenario', 'Recaudo aprox.', 'Tu comisión'],
 ['Abres la cohorte: 20 Arranques al 5 de octubre', '$59,8 millones', '≈ $6,0 millones'],
 ['10 Arranques + 10 programas completos', '$89,0 millones', '≈ $8,9 millones'],
 ['Grupos llenos: 28 + bono', '$87,9 millones', '≈ $10,8 millones (bono incluido)'],
])

sec('4. TU OFERTA COMISIONAL')
b('por toda venta de tu canal (tus videos, tu red, tu línea).', bold_first='8% del recaudo ')
b('con una de las dos puertas: 12 matrículas en tus primeros 12 DÍAS de campaña al aire (tu día 1 es el día en que se publica tu primera campaña — la meta arranca contigo, no antes), o los 20 cupos vendidos antes del 5 de octubre.', bold_first='Sube a 10% RETROACTIVO sobre TODA tu serie ')
b('si los grupos quedan llenos: 28 matrículas antes del 5 de octubre.', bold_first='Bono de $2.000.000 ')
b('desde su plataforma: 5% cliente nuevo / 2,5% antiguo.', bold_first='Leads que la academia te asigne ')
b('Si alguien del equipo te pasa un cliente registrado y tú lo cierras, quien lo pasó gana el 1% — y al revés igual.')
b('Arranque, programas completos, niveles individuales y sabatinos. Las PUERTAS del 10% y el bono cuentan solo matrículas de la cohorte del 5 de octubre — el resto comisiona normal sin mover el contador.', bold_first='Tus porcentajes aplican a TODO el portafolio: ')

sec('5. A QUIÉN LE VENDEMOS — EL COMPROMISO DEL ESTUDIANTE')
t('Aquí no se matricula a cualquiera, y eso es una FORTALEZA de venta, no un freno:')
b('17 años o más, y sesión de ubicación ANTES de decidir (el nivel lo define la academia, no el cliente).', bold_first='Requisitos de entrada: ')
b('esto se entrena como un deporte — asistencia, tareas, audio diario. La garantía existe PORQUE exige: si el estudiante cumple su parte y no avanza, se le devuelve la plata. Sin su parte, no hay garantía.', bold_first='El trato se dice antes de firmar: ')
b('el que firma advertido no deserta en la semana 3 ni pide devolución. El que entra engañado con "fácil y divertido" se retira, reclama y habla mal. Preferimos un NO honesto hoy que una devolución en noviembre.', bold_first='Por qué se vende así: ')

sec('6. LAS DOS CAMPAÑAS (la academia financia y opera TODA la pauta)')
b('videos de la casa → los leads llegan a la plataforma de IA de la academia, que responde al instante y agenda automáticamente. Esos son leads de la casa (comisionan 5% si se te asignan).', bold_first='Campaña de la academia: ')
b('tus videos → los leads llegan a TU línea. Esos comisionan al 8-10%.', bold_first='TU campaña: ')
b('Tú no pones un peso de pauta ni la administras: la casa financia, el equipo de pauta la opera, y gerencia aprueba pieza y presupuesto antes de publicar.')

sec('7. TUS VIDEOS')
b('Abajo están los briefs: los puntos de cada pieza, dichos como te salgan, natural. No se pagan por pieza: tu pago es la comisión, y la casa pone toda la pauta.', bold_first='Las 5 piezas las grabas tú, CON TUS PALABRAS — no hay guion que leer. ')
b('una sesión de grabación, subtítulos, y al aire — en pauta, lo natural convierte más que lo producido.', bold_first='Velocidad sobre perfección: ')
b('me pasas los videos, los apruebo y salen. Más videos tuyos con los briefs = más leads tuyos.', bold_first='Antes de publicar: ')

sec('8. LA LÍNEA Y EL REGISTRO')
b('operado 100% por ti — más profesional ante el cliente y tu número personal queda libre.', bold_first='Tu línea es un WhatsApp Business de la academia ')
b('la planilla es la que liquida tus comisiones. Sin registro no hay comisión. (A futuro, si entra otra comercial al equipo, ese registro marca de quién es cada lead — hoy el canal digital es todo tuyo.)', bold_first='Todo contacto se registra en la planilla desde el primer mensaje: ')
b('un lead que espera es plata que se enfría.', bold_first='El lead caliente se atiende el mismo día: ')

sec('9. LAS REGLAS DE LA PLATA (para que la primera liquidación no tenga sorpresas)')
b('nunca a cuentas personales del vendedor. Efectivo solo EN LA SEDE, en recepción y con recibo.', bold_first='Todo pago de clientes entra SIEMPRE a las cuentas de la academia — ')
b('la SEPARACIÓN puede ser 100% digital — abono a cuentas de la academia + aceptación escrita del recibo de separación + cita de ubicación agendada en el mismo acto. La FORMALIZACIÓN es siempre en la sede: contrato de matrícula firmado en físico, a más tardar el día de la sesión de ubicación. El chat asegura el cupo; la sede firma el contrato.', bold_first='La matrícula tiene DOS momentos: ')
b('30 de septiembre y 5 de octubre (después, mensual). Se paga contra tu cuenta de cobro dentro de los 5 días hábiles siguientes, sobre pagos cuyo retracto legal ya venció. Si hay retracto o devolución de ley, la comisión se descuenta de tu siguiente liquidación.', bold_first='Liquidación por cortes: ')
b('dentro de los 5 días hábiles de retracto de ley se devuelve completo; vencido ese plazo, NO se devuelve en efectivo — queda como saldo a favor para cualquier programa. Dicho antes, es una regla; descubierto después, es un reclamo.', bold_first='El abono de separación se explica SIEMPRE al cliente: ')
b('Se dice "precio de lanzamiento" y "beca del Fondo" — JAMÁS "descuento" ni "patrocinio". Nada de promesas de visa, empleo, "fácil" o "rápido".', bold_first='Solo la matriz vigente: ')

sec('10. EL PLAN DE ARRANQUE')
t('Firma → kit completo de venta en el acto → grabas las 5 piezas → subtítulos → tu campaña al aire. Tu puerta del 10% cuenta desde el día en que tu primera campaña se publica. La única fecha fija es la cohorte: 5 de octubre.', bold=True)

sec('11. LOS 5 BRIEFS (tus palabras, esta estructura — los repasamos juntas en la entrega)')
t('REGLAS PARA TODAS LAS PIEZAS — lo único innegociable:', bold=True)
b('método + respaldo + garantía a cambio del trabajo del estudiante. Jamás prometer fácil, rápido o divertido.', bold_first='En toda pieza se vende EL TRATO, nunca la facilidad: ')
caja('La única frase LITERAL, cuando se mencione la garantía: "si cumples — asistencia, tareas, evaluaciones — y no avanzas, te devuelven la plata. Por escrito."')
b('gratis · descuento · patrocinio · "fácil" · "rápido" · "bilingüe en X meses" · promesas de empleo o visa · números de cupos inventados (con cero vendidos: "acaban de abrirse los 20 cupos").', bold_first='PROHIBIDO decir: ')
b('Todo cierre: "escríbeme al WhatsApp" (enlace prellenado). 30-45 segundos, una cámara, tu cara, luz natural.')

tabla([
 ['Pieza', 'Para quién', 'Los puntos (en tus palabras)'],
 ['1. EL DIRECTO', 'Público frío', 'Academia que te FIRMA que aprendes o te devuelven la plata · 2 niveles, 200 horas, libros y certificados por $2.990.000 · aquí se entrena en serio, por eso firman garantía · 20 cupos recién abiertos, cohorte 5 de octubre'],
 ['2. EL QUEMADO', 'Ya intentó varias veces', '¿Cuántas veces has empezado inglés? El problema fue el método: sentado, copiando, con miedo · aquí es al revés: de pie, hablando desde el día 1 · garantía firmada [frase literal] · honestidad: es exigente, se entrena como deporte — tú pones el trabajo'],
 ['3. EL EMPRENDEDOR', 'Quiere negocio', 'No es un curso: mira lo que te entregan al final · sales con TU página web publicada y tu primer mensaje de venta a un cliente real del exterior · $2.990.000 con garantía por escrito si haces tu parte — nadie más entrega eso'],
 ['4. EL QUE SE VA', 'Au Pair / trabajar afuera', '¿El inglés te frena para irte? Es el requisito #1 de todos los programas · terminas con video de presentación, postulación diligenciada y entrevista ensayada · entrenamiento en serio con garantía para el que cumple · 5 de octubre, 20 cupos'],
 ['5. MODO DICIEMBRE', 'Urgencia fin de año', '¿Otra vez vas a llegar a diciembre diciendo "el otro año sí"? · arrancando el 5 de octubre, en las fiestas ya llevas 2 meses hablando · el riesgo no es tuyo si pones tu parte [frase literal] · 20 cupos a $2.990.000; después vale más de $3.5 millones'],
])

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Heiiu English Academy · Carrera 27 # 48-49, Bucaramanga · Gerencia · Septiembre 2026'); r.font.size = Pt(8.5); r.font.color.rgb = GRIS

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\INDUCCION_COMERCIAL\CONDICIONES_ECONOMICAS_CERRADORA.docx'
doc.save(out); print('OK', out)
