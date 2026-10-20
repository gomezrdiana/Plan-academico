# -*- coding: utf-8 -*-
"""Prueba práctica del Community Manager en versión de 4 videos de 20 segundos (20/10/2026): un brief común (reglas,
entrega, evaluación) + 4 fichas, una por video; cada candidato recibe el brief con SU ficha. Sin pago por la prueba;
$100.000 si el video se publica. Salida (git, por rol): recursos/comercial/contratacion/COMMUNITY_MANAGER/PRUEBA_4_VIDEOS_CM.docx
y MENSAJES_PRUEBA_4_VIDEOS_CM_2026-10-20.docx (mensajes a los candidatos, con nombres)."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAR = RGBColor(0xE5, 0x4C, 0x09); GRIS = RGBColor(0x66, 0x66, 0x66); BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
BASE = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\contratacion\COMMUNITY_MANAGER'
FECHA_ENTREGA = 'domingo 25 de octubre'

def nuevo():
    d = Document()
    for s in d.sections:
        s.top_margin = Cm(1.2); s.bottom_margin = Cm(1.0); s.left_margin = Cm(1.6); s.right_margin = Cm(1.6)
    d.styles['Normal'].font.name = 'Calibri'; d.styles['Normal'].font.size = Pt(9.5)
    return d
def shd(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear'); el.set(qn('w:color'), 'auto'); el.set(qn('w:fill'), color); tcPr.append(el)
def mk(d):
    def sec(t):
        p = d.add_paragraph(); p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(1)
        r = p.add_run(t); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = NAR
    def t(txt, size=9.5, color=None, after=3, bold=False):
        p = d.add_paragraph(); p.paragraph_format.space_after = Pt(after)
        r = p.add_run(txt); r.font.size = Pt(size); r.bold = bold
        if color: r.font.color.rgb = color
        return p
    def bullets(items, size=9.5):
        for x in items:
            p = d.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(0); r = p.add_run(x); r.font.size = Pt(size)
    def caja(txt, size=9.5):
        tb = d.add_table(rows=1, cols=1); tb.style = 'Table Grid'; c = tb.cell(0, 0); shd(c, 'FFF8E7')
        c.paragraphs[0].paragraph_format.space_after = Pt(0); r = c.paragraphs[0].add_run(txt); r.font.size = Pt(size); r.italic = True
        d.add_paragraph().paragraph_format.space_after = Pt(0)
    def tabla(hdr, filas, anchos, size=9):
        tb = d.add_table(rows=1, cols=len(hdr)); tb.style = 'Table Grid'
        for i, h in enumerate(hdr):
            c = tb.rows[0].cells[i]; shd(c, 'E54C09'); c.paragraphs[0].paragraph_format.space_after = Pt(0)
            r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = BLANCO
        for f in filas:
            row = tb.add_row().cells
            for i, txt in enumerate(f):
                row[i].paragraphs[0].paragraph_format.space_after = Pt(0)
                r = row[i].paragraphs[0].add_run(txt); r.font.size = Pt(size); r.bold = (i == 0)
        for row in tb.rows:
            for i, w in enumerate(anchos): row.cells[i].width = Cm(w)
        d.add_paragraph().paragraph_format.space_after = Pt(0)
    return sec, t, bullets, caja, tabla

# =============================================================== BRIEF + FICHAS
d = nuevo(); sec, t, bullets, caja, tabla = mk(d)
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
r = p.add_run('HEIIU ENGLISH ACADEMY + PCG CONSTRUCTORA · COMMUNITY MANAGER DOS MARCAS'); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = NAR
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(6)
r = p.add_run('PRUEBA PRÁCTICA — UN VIDEO DE 20 SEGUNDOS, DESDE TU CASA'); r.bold = True; r.font.size = Pt(11)

sec('1. Qué es la prueba')
t('Te asignamos un tema (la ficha al final de este documento). Tú escribes el guion de un video de 20 segundos, lo grabas desde tu casa, lo editas y lo entregas junto con el guion escrito. En el cargo real, el contenido orgánico de las dos marcas (Heiiu y PCG Constructora) es híbrido: unas veces gerencia entrega el guion y otras lo escribe el CM a partir de una idea. En esta prueba lo escribes tú a partir de la idea de la ficha, y además cumples tal cual las tres instrucciones fijas que trae la ficha: así vemos las dos cosas, criterio propio y seguir instrucciones. Para este ejercicio usaremos la academia. No buscamos producción de agencia: buscamos criterio, mensaje y ejecución limpia. Puedes apoyarte en las herramientas que quieras para escribir; lo que evaluamos es que el mensaje respete las reglas y suene a una persona hablándole a otra, no a una plantilla.')
t('La prueba no tiene pago. Si la academia publica tu video en sus redes o en pauta, te paga $100.000 por la pieza, pases o no el proceso. Se publique o no, puedes usar el video en tu portafolio.', bold=True)

sec('2. El mensaje de la academia (lo que no cambia, sea cual sea el tema; la constructora tiene sus propias reglas, que se entregan aparte)')
bullets([
 'No vendemos que el inglés sea fácil, rápido ni divertido. Vendemos entrenamiento con garantía por escrito: si cumples y no avanzas, te devolvemos el nivel. Quien mira debe sentir que esto es serio y que es para quien se le mide.',
 'Hablamos de lo que la persona va a poder hacer (atender a un paciente, responder en una reunión, tener su hoja de vida en inglés), nunca de lo que le falta. El gancho no le atribuye una carencia a quien mira: nada de "¿No sabes inglés?" ni "¿Te da pena hablar?".',
 'Prohibido en voz, subtítulo o texto en pantalla: precio · "fácil" · "gratis" · "descuento" · "rápido" sin un número · "garantizado en X meses" · "consigue trabajo" · "visa" · "profesores nativos" · nombres de métodos · nombres de otras academias · nombres de programas internos · días de la semana.',
 'Cierre siempre a WhatsApp: "Escríbenos" o "Pregúntanos", con el logo de Heiiu solo en el último segundo.',
 'Todos los videos le hablan a personas desde los 17 años. No se hacen piezas dirigidas a menores.'])

sec('3. Reglas de grabación y edición')
bullets([
 'Duración entre 18 y 25 segundos. Formato 9:16, 1080x1920, MP4, vertical.',
 'Sales tú a cámara, o alguien real que te ayude: cara y voz de una persona. No se aceptan avatares, voces generadas ni imágenes creadas con IA.',
 'Plano medio, fondo limpio (una pared), sin logos de otras marcas ni ropa con marcas. Luz natural o de la habitación; sin aro de luz.',
 'Audio: micrófono de solapa o el celular a menos de 1,5 m de quien habla. Música baja o ninguna, nunca encima de la voz.',
 'Subtítulos grandes, palabra por palabra, dentro de la zona segura de Reels: ni en el 14% superior ni en el 35% inferior. Sin errores de ortografía.',
 'Textos en pantalla con los colores de marca: naranja #E54C09 y amarillo #FBC900 (el amarillo nunca como texto sobre blanco).',
 'Puedes usar una hoja, cartulina o pizarra pequeña para escribir una cifra o una palabra frente a la cámara: funciona muy bien en 20 segundos.'])

sec('4. Qué entregas y cuándo')
bullets([
 'El guion escrito, en una tabla de cuatro columnas: segundos · qué se ve · qué dice la persona (tal cual) · texto en pantalla. Se entrega antes de grabar si quieres comentarios, o junto con el video.',
 'El video terminado, nombrado CMP_[letra de tu ficha]_20s.mp4, enviado por correo electrónico al correo que te indicamos al confirmar, como adjunto o como enlace de Drive o WeTransfer, para que no se comprima. Nunca como video por WhatsApp.',
 'Fecha de entrega: ' + FECHA_ENTREGA + '. Si necesitas otra fecha, la propones al confirmar; lo que se acuerda se cumple.',
 'Dudas antes de grabar, por correo. Lo que no esté en la ficha se pregunta; preferimos la pregunta a la suposición.'])

sec('5. Cómo se evalúa (pasa quien cumple los 7)')
tabla(['Criterio', 'Qué se revisa'], [
 ('1. Gancho', 'Los primeros 2 segundos detienen el dedo sin atribuirle una carencia a quien mira. Nada prohibido.'),
 ('2. Mensaje', 'El video dice lo que pide la ficha, con el tono de la sección 2: serio, concreto, para quien se le mide. Nada prohibido en voz, subtítulo ni pantalla.'),
 ('3. Guion e instrucciones', 'Guion entregado en la tabla de cuatro columnas y ejecutado como se escribió. Las tres instrucciones fijas de la ficha cumplidas tal cual. Se nota que primero se pensó y después se grabó.'),
 ('4. Subtítulos', 'Grandes, palabra por palabra, dentro de la zona segura, sin errores de ortografía.'),
 ('5. Formato y nombre', '9:16, 1080x1920, MP4, 18 a 25 segundos, nombre exacto, enviado por correo o enlace sin comprimir.'),
 ('6. Audio e imagen', 'Voz limpia sin eco ni ruido; encuadre estable; persona real a cámara; sin aro de luz.'),
 ('7. Entrega', 'En la fecha acordada, con el guion. Las mejoras que propongas al tema se conversan antes de grabar, no se improvisan.'),
], [3.4, 13.8])

sec('6. Qué pasa después')
t('Si pasas, se firma el contrato de prestación de servicios para las dos marcas, Heiiu y PCG Constructora, en las condiciones conversadas. En cada marca el trabajo tiene dos partes: los anuncios de pauta desde guiones escritos por gerencia (en la academia, grabados con profesores y estudiantes; en la constructora, en sus proyectos), y el contenido orgánico con parrilla mensual e informe de resultados (Heiiu: Instagram, TikTok y Facebook; PCG Constructora: Instagram y Facebook). Publicar hace parte del trabajo.')

FICHAS = [
 ('A', 'TU HOJA DE VIDA EN INGLÉS, LISTA',
  'Jóvenes y profesionales desde los 17 años que ven vacantes o convocatorias que piden inglés y las pasan de largo.',
  'En Heiiu el estudiante no termina con un certificado: termina con su hoja de vida en inglés, su video de presentación grabado y su entrevista ensayada (es nuestro módulo Pasaporte, pero no digas el nombre). La idea en una frase: "El inglés no te da el puesto; te quita la excusa que hoy te lo cierra."',
  ['"Consigue trabajo", "te contratan", "en X meses". Se vende preparación, jamás empleo.', 'Nombrar LinkedIn como promesa: puede aparecer como contexto ("la vacante que viste ayer"), no como resultado garantizado.'],
  'Una hoja con la palabra "INGLÉS: AVANZADO" escrita a mano, como en una hoja de vida, y la persona diciendo: "Esto no se escribe. Se entrena."',
  'En algún momento del video se escribe a mano, frente a la cámara y legible, la cifra "575 h" (las horas del programa completo).'),
 ('B', 'EL PACIENTE QUE ESCRIBE EN INGLÉS',
  'Recepcionistas, asistentes y coordinadoras, desde los 17 años, de clínicas estéticas, odontológicas y consultorios de Bucaramanga que reciben pacientes del exterior.',
  'El paciente extranjero escribe por WhatsApp, llama para confirmar la cita, pregunta por las indicaciones después del procedimiento, y la respuesta se demora o pasa por el doctor. En Heiiu se entrena exactamente eso: cada clase tiene una simulación real (una llamada, una cita, una explicación) con el profesor corrigiendo. La idea en una frase: "La clínica atiende pacientes de afuera. La recepción también debería."',
  ['Prometer que la clínica va a recibir más pacientes o que la persona va a ganar más.', 'Nombrar clínicas o marcas reales.'],
  'La persona con el celular en la mano leyendo un mensaje en inglés (inventado, sin nombres) y levantando la mirada: "¿Y ahora quién le responde?"',
  'En algún momento del video se dice, en inglés y a cámara, una frase completa de atención al paciente (por ejemplo, confirmar una cita), con su subtítulo.'),
 ('C', 'ENFERMERÍA EN INGLÉS',
  'Enfermeras, auxiliares y estudiantes de enfermería de Bucaramanga, desde los 17 años.',
  'En un turno llega un paciente que solo habla inglés: hay que entender qué le duele, explicar un procedimiento, tranquilizar. Y las convocatorias que piden inglés existen y pasan de largo. En Heiiu se entrena con simulaciones de ese tipo: paciente, médico, familiar; el profesor como mentor que corrige y pide cuentas. La idea en una frase: "Entender al paciente no es un lujo. Es parte del cuidado."',
  ['"Visa", "trabaja en Estados Unidos", "contrato afuera" o cualquier promesa de migración o empleo.', 'Nombrar hospitales, clínicas o programas de reclutamiento.'],
  'La persona en uniforme (o con una bata sencilla) a cámara, sin escenario de hospital, diciendo la primera frase de una simulación en inglés y luego, en español: "Esto se practica. Aquí, todos los días."',
  'En algún momento del video aparece en pantalla, en texto grande, la palabra "GARANTÍA POR ESCRITO" en naranja de marca.'),
 ('D', 'ASÍ SE APRENDE AQUÍ',
  'Cualquier persona desde los 17 años que ya intentó aprender inglés y no pasó del nivel básico.',
  'Una clase en Heiiu no es sentarse a copiar: desde el primer día se habla, se practica, uno se equivoca, se frustra, y el profesor, que es un mentor, corrige y lo repite hasta que sale. Todos los días el estudiante graba un video corto en inglés y lo recibe de vuelta con una mejora. Se dice de frente que no es fácil, y por eso hay garantía por escrito. La idea en una frase: "Aquí te vas a frustrar. Y vas a salir hablando."',
  ['Nombres de métodos o de enfoques pedagógicos. Se dice "clases inmersivas, activas y prácticas", nada más.', '"Fácil", "divertido", "sin esfuerzo", "aprende jugando".', 'Mostrar o nombrar estudiantes reales (se graba desde tu casa; la academia no aparece).'],
  'La persona escribe en una hoja "FÁCIL" y lo tacha; debajo escribe "INTENSO". Y a cámara: "Lo primero no te lo vamos a decir. Lo segundo, sí."',
  'En algún momento del video se dice, tal cual: "Todos los días grabas un video corto en inglés y tu profesor te lo devuelve con una mejora."'),
]
d.add_page_break()
sec('7. Las fichas (cada candidato recibe una)')
for letra, titulo, publico, idea, no, gancho, fijo in FICHAS:
    t('FICHA ' + letra + ' · ' + titulo, bold=True, size=10.5, after=1)
    tabla(['', ''], [
     ('A quién le habla', publico),
     ('Qué tiene que decir (la idea; las palabras son tuyas)', idea),
     ('Lo que no se dice en esta ficha, además de la lista general', ' · '.join(no)),
     ('Un gancho posible (puedes proponer otro)', gancho),
     ('Tres instrucciones fijas (se cumplen tal cual)', '1. ' + fijo + ' · 2. La frase final, dicha tal cual: "Pregúntanos qué fecha te queda. Escríbenos." · 3. El logo de Heiiu aparece solo en el último segundo; antes, nunca.'),
     ('Archivo', 'CMP_' + letra + '_20s.mp4'),
    ], [4.6, 12.6])
out = BASE + r'\PRUEBA_4_VIDEOS_CM.docx'
try: d.save(out); print('OK', out)
except PermissionError: d.save(out.replace('.docx', '_v2.docx')); print('ABIERTO -> _v2')

# =============================================================== MENSAJES
d = nuevo(); sec, t, bullets, caja, tabla = mk(d)
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(4)
r = p.add_run('MENSAJES A CANDIDATOS CM — PRUEBA DE 4 VIDEOS (20/10/2026)'); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = NAR
t('Uso interno. Adjuntar a cada uno PRUEBA_4_VIDEOS_CM.pdf y decirle su letra. Entrega: ' + FECHA_ENTREGA + '. Sin pago por la prueba; $100.000 si se publica. Se envía por correo a los cuatro a la vez, con el PDF adjunto; las preguntas pendientes van en el mismo correo, sin condicionar la prueba. Asignación: A Kimberly · B Astrid · C Sebastián · D Luna.', size=9, color=GRIS)
FIRMA = '\n\nDiana Gómez\nGerente General, Heiiu English Academy'
MSG = [
 ('KIMBERLY (ficha A)', 'Hola Kimberly, te pido disculpas por la demora: arrancamos un cargo nuevo esta semana y se me corrió todo. Gracias por la respuesta tan completa. Cumples con lo que necesitamos en ubicación, equipo, disponibilidad y proceso creativo.

Sobre la tarifa: el inicio es de $1.800.000 por las dos marcas, con revisión al tercer mes. Si se cumplen los mínimos acordados, desde el cuarto mes pasa a $2.000.000. Preferimos que el aumento vaya atado a resultados, igual que todo en la academia.

El siguiente paso es una prueba práctica corta: un video de 20 segundos desde tu casa, con el tema de la ficha A del documento adjunto y las reglas de mensaje de la academia. Te soy transparente: la prueba no tiene pago, pero si publicamos tu video te pagamos $100.000 por la pieza, pases o no. Ahí verás que nuestro mensaje es distinto al habitual: no vendemos que el inglés sea fácil, vendemos entrenamiento con garantía por escrito.

La entrega es el ' + FECHA_ENTREGA + ', a este mismo correo (adjunto o enlace). Las dudas, también por aquí. Si ya no te interesa, me lo dices con confianza.' + FIRMA),
 ('LUNA (ficha D)', 'Hola Luna, gracias por tu postulación, por el portafolio y por el video. Nos gustó tu respuesta sobre los testimonios: conversación cercana y preguntas sencillas es exactamente como los hacemos aquí.

El siguiente paso es una prueba práctica corta: un video de 20 segundos desde tu casa, con el tema de la ficha D del documento adjunto y las reglas de mensaje de la academia. Te soy transparente: la prueba no tiene pago, pero si publicamos tu video te pagamos $100.000 por la pieza, pases o no. La entrega es el ' + FECHA_ENTREGA + ', a este mismo correo (adjunto o enlace); las dudas, también por aquí.

Y en el mismo correo de la entrega, por favor, tu hoja de vida con fechas (desde y hasta) de cada trabajo y dos referencias: nombre, cargo y teléfono de un jefe o cliente para quien hayas hecho contenido o manejado redes.' + FIRMA),
 ('ASTRID (ficha B)', 'Hola Astrid, gracias por tu postulación, por el portafolio y por los videos. Se nota oficio.

El siguiente paso es una prueba práctica corta: un video de 20 segundos desde tu casa, con el tema de la ficha B del documento adjunto y las reglas de mensaje de la academia. Te soy transparente: la prueba no tiene pago, pero si publicamos tu video te pagamos $100.000 por la pieza, pases o no. La entrega es el ' + FECHA_ENTREGA + ', a este mismo correo (adjunto o enlace); las dudas, también por aquí.

Y dos preguntas de frente, para responder en el mismo correo, porque de eso depende que el cargo te sirva: parte de lo que se graba es en las clases de la noche (6:30 a 8:30) y los sábados en la mañana, y tu disponibilidad es hasta mediodía: ¿podrías cubrir al menos dos noches por semana en la academia? Y hoy manejas tres cuentas; este cargo cubre dos marcas y pide presencia en la academia varias veces por semana: ¿cómo lo ves en tiempo?

Una aclaración sobre los testimonios: aquí siempre salen con la persona real, su cara y su voz, porque es lo que respalda la garantía que firmamos; no usamos avatares de IA para eso. En la prueba también: sales tú o alguien real.' + FIRMA),
 ('SEBASTIÁN (ficha C)', 'Hola Sebastián, retomo tu postulación y te pido disculpas por la demora. El siguiente paso es una prueba práctica corta: un video de 20 segundos desde tu casa, con el tema de la ficha C del documento adjunto y las reglas de mensaje de la academia. Te soy transparente: la prueba no tiene pago, pero si publicamos tu video te pagamos $100.000 por la pieza, pases o no. La entrega es el ' + FECHA_ENTREGA + ', a este mismo correo (adjunto o enlace); las dudas, también por aquí.

En el mismo correo de la entrega, por favor: las fechas (desde y hasta) de tus trabajos en Star Market, Even Company y La Gente de la Cámara; dos referencias (nombre, cargo y teléfono de un jefe o cliente para quien hayas trabajado en contenido o redes); y enlaces a cuentas que hayas manejado o a videos tuyos publicados.' + FIRMA),
]
for titulo, cuerpo in MSG:
    sec(titulo)
    for parr in cuerpo.split('\n\n'):
        t(parr, after=4)
out = BASE + r'\MENSAJES_PRUEBA_4_VIDEOS_CM_2026-10-20.docx'
try: d.save(out); print('OK', out)
except PermissionError: d.save(out.replace('.docx', '_v2.docx')); print('ABIERTO -> _v2')
