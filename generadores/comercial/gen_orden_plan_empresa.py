# -*- coding: utf-8 -*-
"""ORDEN PLAN EMPRESA (1 pagina): tarifa corporativa por volumen (3-7 / 8-15 / 16-30 cupos). La cierra la ejecutiva de alianzas; firma gerencia.
Salida: recursos/comercial/B2B/ORDEN_PLAN_EMPRESA_v1.docx"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

GRIS = RGBColor(0x77, 0x77, 0x77)
d = Document()
for s in d.sections:
    s.top_margin = Cm(1.3); s.bottom_margin = Cm(1.1); s.left_margin = Cm(1.9); s.right_margin = Cm(1.9)
d.styles['Normal'].font.name = 'Calibri'; d.styles['Normal'].font.size = Pt(9.5)

def shd(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear'); el.set(qn('w:color'), 'auto'); el.set(qn('w:fill'), color); tcPr.append(el)

def t(txt, bold=False, after=3, size=9.5, align=None, color=None):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    if align: p.alignment = align
    r = p.add_run(txt); r.bold = bold; r.font.size = Pt(size)
    if color: r.font.color.rgb = color

def cl(n, tit, txt):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f'{n}. {tit}. '); r.bold = True; r.font.size = Pt(9.5)
    r2 = p.add_run(txt); r2.font.size = Pt(9.5)

t('ORDEN DE SERVICIO — PLAN EMPRESA HEIIU', bold=True, size=12.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
t('Formato estándar (3 a 30 cupos) · Heiiu English Academy — Global Teacher S.A.S., NIT 900.422.478-2 · v1 · octubre de 2026', size=8.5, color=GRIS, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

t('LA EMPRESA ______________________________________________, NIT ____________________, representada por ______________________________________, cargo ____________________, contrata con GLOBAL TEACHER S.A.S. (Heiiu English Academy) el programa de inglés para sus colaboradores, en las condiciones estándar siguientes:', after=6)

tb = d.add_table(rows=5, cols=4); tb.style = 'Table Grid'
hdr = ['Programa completo A1 → B2 (575 h)', 'Cupos', 'Valor por cupo (según el tramo)', 'Total']
for j, h in enumerate(hdr):
    c = tb.rows[0].cells[j]; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9); shd(c, 'FFF8E7')
filas = [['Tramo 1: de 3 a 7 cupos (grupos regulares)', '____', '$4.782.800', '$ ______________'],
         ['Tramo 2: de 8 a 15 cupos (grupos regulares, o grupo cerrado en la sede desde 12)', '____', '$4.350.000', '$ ______________'],
         ['Tramo 3: de 16 a 30 cupos (grupo cerrado en la sede o en la empresa)', '____', '$3.950.000', '$ ______________'],
         ['TOTAL (un solo tramo por orden)', '____', '', '$ ______________']]
for i, f in enumerate(filas, 1):
    for j, v in enumerate(f):
        r = tb.rows[i].cells[j].paragraphs[0].add_run(v); r.font.size = Pt(9); r.bold = (i == 4)
d.add_paragraph().paragraph_format.space_after = Pt(2)

cl(1, 'QUÉ INCLUYE', 'Los cupos contratados, con el mismo programa, materiales, certificado por nivel y garantía de aprendizaje por escrito que cualquier estudiante. Tramo 1 y 2: en los grupos regulares de la academia (jornadas vigentes); desde 12 cupos puede abrirse grupo cerrado en la sede. Tramo 3: grupo cerrado en la sede o en las instalaciones de LA EMPRESA, de lunes a viernes entre 7:00 am y 8:00 pm, con horario fijo acordado por escrito. Cada colaborador firma su contrato de matrícula y su anexo de garantía; LA EMPRESA es quien paga.')
cl(2, 'FORMA DE PAGO (marcar una)', '☐ De contado contra factura, antes del inicio de clases.   ☐ Inicial contra factura (30% en el tramo 1; 40% en los tramos 2 y 3) y el saldo en tres (3) cuotas mensuales iguales, la última antes de terminar el primer nivel.   ☐ Copago: LA EMPRESA paga $ ______________ por cupo y el colaborador paga el saldo con su propio contrato de matrícula.')
cl(3, 'COLABORADORES', 'LA EMPRESA designa por escrito a los colaboradores beneficiarios antes del inicio del nivel. Un cupo puede reasignarse a otro colaborador una sola vez, antes de la cuarta clase del nivel. Los cupos no usados no se devuelven en dinero; quedan como saldo a favor de LA EMPRESA durante doce (12) meses.')
cl(4, 'CONDICIONES DEL PROGRAMA', 'Aplican a cada colaborador las condiciones académicas del contrato de matrícula (asistencia, tareas, evaluaciones y puntualidad). LA EMPRESA recibe reporte de asistencia y cumplimiento de sus colaboradores; las notas y resultados se comunican solo al estudiante por coordinación. La garantía de aprendizaje opera a favor de quien pagó el cupo.')
cl(5, 'LO QUE NO INCLUYE EN ESTE FORMATO', 'Más de treinta (30) cupos, horarios fuera de la franja indicada, plazos distintos de los del numeral 2 y cualquier valor distinto de los tres tramos. Esas condiciones las cotiza por separado la gerencia de la academia.')
cl(6, 'PERFECCIONAMIENTO', 'Esta orden queda en firme con la firma del representante legal de la academia y el pago inicial. Vigencia de los precios: treinta (30) días desde la firma de LA EMPRESA.')

t('Contacto en LA EMPRESA: ______________________________  Cargo: ____________________  Teléfono / correo: ______________________________', after=6)
t('Firmada en Bucaramanga el ____ / ____ / 2026', after=10)
t('POR LA EMPRESA                                                                          POR HEIIU — GLOBAL TEACHER S.A.S.', after=14)
t('_____________________________                                          _____________________________', after=1)
t('Nombre: ________________________                                        DIANA MARCELA GÓMEZ RANGEL', after=1)
t('C.C. ______________________________                                       Representante Legal', after=4)
t('Presentado por: ______________________________ (ejecutiva de alianzas Heiiu) · Fecha: ____ / ____ / 2026', size=8.5, color=GRIS)

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\B2B\ORDEN_PLAN_EMPRESA_v1.docx'
d.save(out); print('OK', out)
