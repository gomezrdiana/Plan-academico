# -*- coding: utf-8 -*-
"""Borrador contrato docente v3: hora catedra por actas + gestion de alianzas. Base: v2 firmado (27/04/2026)."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

GRIS = RGBColor(0x77, 0x77, 0x77)
ROJO = RGBColor(0xC0, 0x00, 0x00)

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(2.2); s.bottom_margin = Cm(2.0); s.left_margin = Cm(2.4); s.right_margin = Cm(2.4)
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(10.5)

def titulo(txt):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(13)

def aviso(txt):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(10); r.font.color.rgb = ROJO

def t(txt, bold=False):
    p = doc.add_paragraph(); r = p.add_run(txt); r.bold = bold; r.font.size = Pt(10.5)
    p.paragraph_format.space_after = Pt(5)
    return p

def clau(nombre, cuerpo):
    p = doc.add_paragraph()
    r = p.add_run(nombre + ' '); r.bold = True; r.font.size = Pt(10.5)
    r2 = p.add_run(cuerpo); r2.font.size = Pt(10.5)
    p.paragraph_format.space_after = Pt(5)

def num(n, txt, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(str(n) + '. ' + txt); r.font.size = Pt(10.5); r.bold = bold
    p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.space_after = Pt(3)

titulo('CONTRATO DE PRESTACIÓN DE SERVICIOS (HORA CÁTEDRA + GESTIÓN DE ALIANZAS) — VERSIÓN 2026 v3')
aviso('BORRADOR 14/09/2026 — PENDIENTE VISTO BUENO DEL ABOGADO ANTES DE FIRMA · Cambios frente al v2 firmado: objeto dual, vigencia por actas, cesión de obra creada, remuneración de alianzas, sin facultad de representación')
doc.add_paragraph()

t('Entre los suscritos a saber: DIANA MARCELA GÓMEZ RANGEL, mayor de edad, identificada con cédula de ciudadanía No. 63.552.709 expedida en Bucaramanga, quien actúa en nombre de GLOBAL TEACHER S.A.S., sociedad legalmente constituida, con NIT 900.422.478-2, con domicilio principal en la ciudad de Bucaramanga, según consta en el Certificado de Existencia y Representación Legal, quien para efectos del presente contrato se denominará EL CONTRATANTE, y [NOMBRE COMPLETO DEL DOCENTE], mayor de edad, vecino(a) de la ciudad de Bucaramanga, identificado(a) con cédula de ciudadanía No. [___________], quien en adelante se denominará EL CONTRATISTA, hemos acordado celebrar el presente contrato de prestación de servicios así:')

clau('CLÁUSULA PRIMERA. OBJETO:', 'El objeto del presente contrato lo constituye la prestación de servicios profesionales por parte del CONTRATISTA en dos frentes: (a) como docente HORA CÁTEDRA en los cursos de inglés que le sean asignados; y (b) como GESTOR DE ALIANZAS institucionales dentro de la iniciativa Go Above and Beyond (programas Business Partner, Student Partner y School Partner).')

clau('CLÁUSULA SEGUNDA. OBJETO ESPECÍFICO Y VALOR CONTRACTUAL:', '(a) DOCENCIA: el CONTRATISTA dictará los cursos que se le asignen mediante ACTA INTERNA DE PROGRAMACIÓN ACADÉMICA suscrita por las partes, en la cual se indicará para cada curso: nivel, número de horas, horario y fechas. El valor de la hora cátedra es de DIECISIETE MIL PESOS COP ($17.000) [verificar valor vigente], pagadero por mensualidades vencidas de acuerdo con las horas efectivamente impartidas. (b) GESTIÓN COMERCIAL — FACULTATIVA: EL CONTRATISTA PODRÁ, si así lo decide, comercializar los programas de la institución usando exclusivamente los materiales y condiciones aprobados por gerencia. Esta actividad no constituye obligación contractual ni su omisión genera incumplimiento. Cuando la ejerza, la remuneración será por TRES canales, siempre sobre el recaudo efectivo a caja: (i) leads entregados por la institución: CINCO por ciento (5%) del recaudo de cliente nuevo y 2,5% de cliente antiguo; (ii) canal con PAUTA FINANCIADA POR LA INSTITUCIÓN y ejecutada por EL CONTRATISTA: OCHO por ciento (8%) del recaudo de cliente nuevo, con entrega de reportes de campaña e inversión a gerencia; (iii) CANAL PROPIO AUTOFINANCIADO (su pauta, su red, su gestión, sus recursos): DIEZ por ciento (10%) del recaudo de cliente nuevo, sin más carga que el registro de la venta. ATRIBUCIÓN DE CANAL: la pauta del canal (ii) se ejecuta exclusivamente desde la cuenta publicitaria de la institución (que asume el gasto y conserva visibilidad plena de campañas e inversión); la pauta del canal (iii) se ejecuta desde cuentas y medios de pago propios del CONTRATISTA. Cada campaña usa su propio enlace de contacto, de modo que el canal de cada lead queda determinado por la cuenta publicitaria que originó el anuncio. Los leads de red o gestión personal del CONTRATISTA pertenecen al canal (iii). PARA RECLAMAR CUALQUIER COMISIÓN, el contacto debe estar registrado en la planilla comercial de la institución; en caso de concurrencia entre vendedores, la comisión corresponde a quien registró primero el contacto. Sin recaudo no hay comisión.')
t('PARÁGRAFO I. El presente contrato se mantiene vigente por el término de UN (1) AÑO contado desde su suscripción, prorrogable por acuerdo escrito. La terminación o culminación de un curso individual NO termina el contrato: solo cierra el acta correspondiente. Cualquiera de las partes podrá terminar el contrato con PREAVISO ESCRITO de TREINTA (30) días calendario; tratándose de cursos en ejecución, EL CONTRATISTA acompañará la transición del curso durante ese período conforme a las guías de la institución.', bold=True)
t('PARÁGRAFO II. Si las actividades académicas se extienden más allá de las horas inicialmente pactadas en el acta (reprogramaciones, refuerzos, recuperación de festivos, evaluaciones adicionales, etc.), las horas adicionales se pagarán al mismo valor de hora cátedra establecido, previo acuerdo escrito entre las partes.')
t('PARÁGRAFO III. Cualquier cambio de horario, suspensión de clase o reprogramación deberá ser comunicado a la administración con la antelación que esta determine en sus políticas internas.')
t('PARÁGRAFO IV — SIN FACULTAD DE REPRESENTACIÓN: en la gestión de alianzas, EL CONTRATISTA presenta, gestiona y entrega información, pero NO tiene facultad para firmar convenios, otorgar becas o descuentos, pactar condiciones económicas ni comprometer jurídicamente a GLOBAL TEACHER S.A.S. Todo convenio lo suscribe únicamente el representante legal.', bold=True)

clau('CLÁUSULA TERCERA. PERFECCIONAMIENTO:', 'El presente contrato se entiende perfeccionado desde la fecha de suscripción.')

clau('CLÁUSULA CUARTA: OBLIGACIONES DEL CONTRATISTA:', '')
obls = [
 'Poner al servicio de GLOBAL TEACHER SAS sus conocimientos y capacidad de trabajo en el desempeño de las funciones propias del oficio mencionado y en las labores anexas y complementarias del mismo, de conformidad con las órdenes e instrucciones que se le impartan.',
 'Guardar en el desempeño de sus funciones y fuera de ellas la discreción y la confidencialidad que exige la lealtad que debe a EL CONTRATANTE.',
 'Prestar sus servicios como docente en las clases de inglés asignadas según los horarios disponibles por ambas partes.',
 'Respetar y no usar el Material Protegido (definido en la CLÁUSULA QUINTA) por fuera de las instalaciones de GLOBAL TEACHER SAS, ya que se asume total confidencialidad y derechos de autor del material usado.',
 'No cortejar o iniciar ningún tipo de relación sentimental con algún estudiante del Instituto de Idiomas GLOBAL TEACHER SAS.',
 'Acatar las políticas de seguridad de la institución.',
 'Responder por los elementos, herramientas, máquinas, muebles, enseres, materiales, materias primas, etc., de trabajo que EL CONTRATANTE le entregue para el desempeño de su cargo.',
 'Conservar y entregar oportunamente todos los documentos, informes, razones, envíos que, con destino a EL CONTRATANTE, haya recibido de terceros en el ejercicio de sus funciones.',
 'Guardar absoluta reserva sobre los hechos, documentos físicos y/o electrónicos, informaciones y, en general, sobre todos los asuntos y materias que lleguen a su conocimiento por causa o con ocasión de su contrato, lo estipulado en la Ley 1266 de 2008 de Habeas Data y las normas que la reglamenten o modifiquen.',
 'Cumplir con todas las funciones principales y secundarias anexas en este contrato.',
 'Y todas las demás que sean complementarias, conexas de sus actividades que le imparta el jefe inmediato o quien haga sus veces.',
 'Portar siempre la dotación dada por EL CONTRATANTE y mantener siempre una imagen impecable de la empresa.',
 'Cumplir puntualmente con los horarios de entrada a la sede de Global Teacher o a cualquier evento programado por la empresa. De manera específica, EL CONTRATISTA se compromete a presentarse en la sede al menos DIEZ (10) MINUTOS ANTES de la hora de inicio de cada clase, con el fin de revisar la guía del día, preparar el tablero, los materiales y el ritual de apertura, y recibir cualquier indicación de coordinación.',
 'Asistir a los eventos programados de Seguridad y Salud en el Trabajo y cualquier evento relacionado con el objetivo de la empresa.',
 'Realizar los aportes al Sistema General de Seguridad Social y Parafiscales de Ley y presentar la planilla de pago como requisito para el pago mensual.',
 'Cumplir íntegramente con los procedimientos pedagógicos establecidos por GLOBAL TEACHER S.A.S. para cada clase, incluyendo la ejecución de la guía del día tal como es entregada por coordinación, los rituales de apertura y cierre, el seguimiento de tareas, y la entrega de reportes. EL CONTRATISTA se compromete a NO improvisar contenido fuera de la guía ni a modificar el método pedagógico sin autorización escrita de coordinación.',
 'Asistir puntualmente a todas las clases asignadas. Cualquier inasistencia, retraso o necesidad de reprogramación deberá ser notificada a coordinación con la antelación que defina la política interna. La inasistencia injustificada o reiterada será causal de revisión del contrato conforme a la Cláusula Décima Segunda.',
 'Entregar reporte escrito y firmado al final de cada clase, conforme al formato vigente entregado por coordinación, incluyendo: asistencia, tareas verificadas, errores observados y entregables del día, junto con el registro fotográfico del paquete de evidencia que defina coordinación. Este reporte es la base del sistema de auditoría académica de la institución.',
 'Cumplir con el protocolo de manejo de errores y evaluaciones: los errores observados en clase se registran de forma anónima en el papel físico de errores que se entrega a coordinación al final del día. El registro nominal con citas literales se conserva en la libreta privada del profesor exclusivamente para coordinación. EL CONTRATISTA NO comunica a los estudiantes resultados de evaluaciones, midterms, finales, criterios de calificación ni notas; esa comunicación es competencia exclusiva de coordinación.',
 'Respetar la propiedad intelectual de las metodologías y rituales pedagógicos de GLOBAL TEACHER S.A.S. No mencionar a estudiantes los nombres técnicos de las metodologías aplicadas. La identificación operativa de los bloques pedagógicos se utiliza únicamente al interior de la institución y se considera Material Protegido conforme a la Cláusula Quinta.',
 'Mantener absoluta confidencialidad sobre la información de los estudiantes (datos personales, evaluaciones, comentarios de coordinación, planes individuales) conforme a la Ley 1581 de 2012 de Protección de Datos Personales.',
 'EN LA GESTIÓN DE ALIANZAS: usar únicamente los materiales, cartas y condiciones aprobados por gerencia; no prometer empleo, visas, prácticas, contratos de aprendizaje ni resultados; no negociar precios, becas ni condiciones por fuera de las escritas; registrar ante gerencia cada gestión realizada (empresa/colegio, contacto, estado, próximo paso); y tratar las listas de contactos y condiciones comerciales como Material Protegido.',
 'NO COMERCIALIZAR programas, cursos o servicios propios o de terceros (incluidos programas de estudios en el exterior) a estudiantes, exalumnos, acudientes o bases de datos de la institución, salvo acuerdo escrito de remisión firmado por el representante legal. Toda pieza publicitaria del canal propio del CONTRATISTA que use la marca o los programas de la institución requiere aprobación escrita de gerencia ANTES de su publicación. El aula jamás se usa para actividad comercial.',
]
for i, o in enumerate(obls, 1):
    num(i, o, bold=(i == 22))
t('PARÁGRAFO V. Las modificaciones, adiciones al presente contrato se harán de mutuo acuerdo entre las partes contratantes, previo documento que lo justifique suscrito entre el contratante y el docente.')

clau('CLÁUSULA QUINTA: DEFINICIÓN DE "MATERIAL PROTEGIDO", PROPIEDAD INTELECTUAL Y CESIÓN DE OBRA:', 'Para efectos del presente contrato, se entiende por "Material Protegido" todo material físico o digital que GLOBAL TEACHER S.A.S. ha desarrollado, adaptado o curado para su modelo pedagógico y comercial, incluyendo de forma enunciativa pero no taxativa: a) guías de clase impresas o digitales; b) manuales del profesor; c) anexos pedagógicos; d) plantillas de reportes, autochequeos, calendarios de virtudes y formatos de evaluación; e) sistemas y rituales pedagógicos y cualquier bloque, ritual, secuencia, formato o procedimiento pedagógico desarrollado por la institución; f) la secuencia operativa de los bloques de clase, su orden, sus tiempos y la manera específica en que se hilan entre sí; g) las metodologías académicas integradas en el modelo, su selección, adaptación y combinación; h) bases estructurales y planes de clase de todos los niveles; i) bases de datos de estudiantes, métricas pedagógicas y reportes de progreso; j) los materiales comerciales y de alianzas (cartas, hojas de condiciones, presentaciones, listas de contactos empresariales e institucionales, matrices de precios y becas); k) cualquier otro material o información proporcionado por EL CONTRATANTE durante la vigencia del contrato. CESIÓN: todo material, documento, carta, presentación o pieza — pedagógica o comercial — que EL CONTRATISTA cree, redacte o adapte en ejecución del presente contrato se entiende realizado por encargo de GLOBAL TEACHER S.A.S., y sus derechos patrimoniales de autor pertenecen exclusivamente a GLOBAL TEACHER S.A.S., conforme a la Ley 23 de 1982 y la Ley 1915 de 2018. EL CONTRATISTA reconoce expresamente que la totalidad del Material Protegido es propiedad intelectual exclusiva de GLOBAL TEACHER S.A.S.')

clau('CLÁUSULA SEXTA: VIGENCIA DE LAS OBLIGACIONES DE CONFIDENCIALIDAD Y PROPIEDAD INTELECTUAL POST-CONTRATO:', 'Las obligaciones de confidencialidad, derechos de autor y propiedad intelectual establecidas en las CLÁUSULAS CUARTA y QUINTA permanecerán vigentes durante un período mínimo de CINCO (5) AÑOS posteriores a la terminación del presente contrato, independientemente de la causa de terminación. Durante este período, EL CONTRATISTA se obliga a: no utilizar la metodología ni los sistemas de la institución fuera de GLOBAL TEACHER S.A.S.; no reproducir, replicar, enseñar, transferir, capacitar a terceros, vender, licenciar, publicar ni divulgar el Material Protegido; no adaptarlo para presentarlo como propio o de un tercero; no utilizarlo como base para ningún producto, curso, academia, plataforma o servicio educativo propio o de un tercero; y no nombrar a estudiantes ni a terceros los sistemas internos de la institución. Esta vigencia post-contractual aplica únicamente a las obligaciones de confidencialidad y propiedad intelectual sobre el Material Protegido y NO constituye una cláusula de no competencia post-empleo. EL CONTRATISTA conserva plena libertad de ejercer su profesión como docente de inglés en otras instituciones, siempre que NO utilice el Material Protegido aquí definido.')

clau('CLÁUSULA SÉPTIMA: DEVOLUCIÓN Y ELIMINACIÓN DE MATERIAL AL TERMINAR EL CONTRATO:', 'A la terminación del contrato por cualquier causa, EL CONTRATISTA se obliga a devolver físicamente todo el Material Protegido en un plazo máximo de tres (3) días hábiles, a eliminar de todos sus dispositivos personales la totalidad de los archivos digitales del Material Protegido, y a certificar por escrito el cumplimiento de lo anterior en el mismo plazo. El incumplimiento activará la cláusula penal de la CLÁUSULA OCTAVA.')

clau('CLÁUSULA OCTAVA: CLÁUSULA PENAL POR VIOLACIÓN DE CONFIDENCIALIDAD Y PROPIEDAD INTELECTUAL:', 'El incumplimiento de las obligaciones de confidencialidad, propiedad intelectual y devolución de material dará derecho a EL CONTRATANTE a reclamar perjuicios y a exigir el pago de una cláusula penal escalonada según la gravedad: violación LEVE hasta CINCO (5) SMLMV; violación MODERADA hasta QUINCE (15) SMLMV; violación GRAVE hasta TREINTA (30) SMLMV, por cada violación documentada; a ejercer las acciones penales que correspondan (Ley 1273 de 2009, Ley 1915 de 2018 y demás normas aplicables) y a solicitar medidas cautelares. La calificación de gravedad la determinará EL CONTRATANTE con base en evidencia documentada, sin perjuicio del derecho del CONTRATISTA a controvertirla judicialmente.')

clau('CLÁUSULA NOVENA: OBLIGACIONES DE LA EMPRESA:', 'Pagar al contratista el valor del presente contrato por mensualidades vencidas; facilitar los espacios físicos y los materiales requeridos; garantizar el debido proceso ante cualquier circunstancia que involucre la conducta del docente; hacer la liquidación que corresponda al término del contrato; entregar formalmente copia firmada de este contrato y dejar constancia escrita de cada pieza de Material Protegido entregada (con fecha y descripción); y liquidar oportunamente la remuneración de alianzas pactada contra los ingresos efectivos a caja.')

clau('CLÁUSULA DÉCIMA: SUPERVISIÓN Y RESPONSABILIDAD:', 'La supervisión de la labor docente será ejercida por el Director(a) Académico(a); la supervisión de la gestión de alianzas será ejercida directamente por gerencia. Se llevará registro escrito de la entrega de Material Protegido y de su devolución, y de las gestiones de alianzas reportadas.')

clau('CLÁUSULA DÉCIMA PRIMERA: CAPACIDAD CONTRACTUAL:', 'EL CONTRATISTA declara, bajo la gravedad del juramento, que tiene capacidad jurídica para contratar y que no se encuentra incurso(a) en ninguna inhabilidad o incompatibilidad.')

clau('CLÁUSULA DÉCIMA SEGUNDA: CAUSALES DE TERMINACIÓN:', 'Mutuo acuerdo; unilateralmente por parte de GLOBAL TEACHER SAS cuando EL CONTRATISTA incumpla cualquiera de las obligaciones de este contrato (la violación de las cláusulas de confidencialidad, propiedad intelectual o devolución de material será causal de terminación inmediata, sin perjuicio de las acciones legales); fuerza mayor o caso fortuito; vencimiento del término establecido; o imposibilidad física, jurídica o profesional del CONTRATISTA.')

clau('CLÁUSULA DÉCIMA TERCERA: DOMICILIO CONTRACTUAL.', 'Para todos los efectos legales, se tendrá como domicilio contractual la ciudad de Bucaramanga, Colombia.')

clau('CLÁUSULA DÉCIMA CUARTA: ACEPTACIÓN.', 'EL CONTRATISTA declara haber leído íntegramente el presente contrato, comprenderlo en su totalidad, y aceptar todas y cada una de sus cláusulas, en particular las relativas a confidencialidad, propiedad intelectual, cesión de obra, devolución de material y cláusula penal.')

doc.add_paragraph()
t('Para constancia se firma en la ciudad de BUCARAMANGA, el día ______ de ____________ de 2026.')
for _ in range(2): doc.add_paragraph()
t('EL INSTITUTO'); doc.add_paragraph()
t('__________________________________')
t('DIANA MARCELA GÓMEZ RANGEL · GLOBAL TEACHER S.A.S. · NIT 900.422.478-2')
doc.add_paragraph()
t('EL CONTRATISTA'); doc.add_paragraph()
t('__________________________________')
t('[NOMBRE] · C.C. [___________]')

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\contratos\BORRADOR_CONTRATO_DOCENTE_ALIANZAS_v3.docx'
doc.save(out); print('OK', out)
