# -*- coding: utf-8 -*-
"""Mensaje a Jaider (LaHaus) para montar la reactivacion de octubre: plantilla, botones, a quien, exclusiones,
que hace el agente con cada boton, fecha de envio. Sin datos de leads. Salida (git):
recursos/comercial/plataforma LaHaus/reactivacion 23 oct 2026.docx (nombre puesto por gerencia)"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\plataforma LaHaus\reactivacion 23 oct 2026.docx'
NAR = RGBColor(0xE5, 0x4C, 0x09); GRIS = RGBColor(0x66, 0x66, 0x66)
d = Document()
for s in d.sections:
    s.top_margin = Cm(1.2); s.bottom_margin = Cm(1.0); s.left_margin = Cm(1.6); s.right_margin = Cm(1.6)
d.styles['Normal'].font.name = 'Calibri'; d.styles['Normal'].font.size = Pt(10)

def p(t, bold=False, size=10, color=None, after=4, italic=False):
    x = d.add_paragraph(); x.paragraph_format.space_after = Pt(after)
    r = x.add_run(t); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    return x
def h(t):
    x = d.add_paragraph(); x.paragraph_format.space_before = Pt(6); x.paragraph_format.space_after = Pt(2)
    r = x.add_run(t); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = NAR

p('Jaider, de acuerdo con la reactivación y con aprovechar la fecha. Te dejo todo en un solo mensaje para que quede configurada. Igual quedo atenta a tus comentarios y sugerencias.\n\n'
  'Para seguir probando qué nos funciona hice dos imágenes y dos plantillas para repartir alternando en la lista, así la prueba del flujo no se mezcla. Se mide por plantilla: botones, respuesta por texto y citas.')
p('PLANTILLA A (imagen A):', True, after=2)
p('Hola {{cName}}, soy Louissiana, la asistente de Heiiu. Hace un tiempo nos escribiste por el inglés y la conversación se quedó ahí.\n\n'
  'Este Halloween, lo que debería darte miedo no es el inglés: es perder la oportunidad. La entrevista, el viaje o el cliente de afuera llegan igual, con o sin él.\n\n'
  'En noviembre abre grupo nuevo y todavía hay cupo, con garantía de aprendizaje por escrito. ¿Para qué lo quieres: viajar, trabajar o estudiar? (Si es para tu negocio o para otra persona, escríbenos.)', size=10, after=6)
p('PLANTILLA B (imagen B):', True, after=2)
p('Hola {{cName}}, soy Louissiana, la asistente de Heiiu. Hace un tiempo nos escribiste por el inglés y la conversación se quedó ahí.\n\n'
  'Este Halloween, que no te dé miedo hablar inglés: aquí se habla desde la primera clase, en grupos de máximo 16 y con un profesor al lado que te corrige ahí mismo.\n\n'
  'En noviembre abre grupo nuevo y todavía hay cupo, con garantía de aprendizaje por escrito. ¿Para qué lo quieres: viajar, trabajar o estudiar? (Si es para tu negocio o para otra persona, escríbenos.)', size=10, after=4)
p('Botones: [Viajar] · [Trabajar] · [Estudiar]', True)
p('"Responde BAJA para no recibir más mensajes."', size=9.5, color=GRIS)
p('Las dos llevan los mismos botones y el mismo pie. Sin nombre de programa, sin precio, sin "descuento", "gratis" ni "fácil". La imagen va arriba, como en los ejemplos que me mandaste. Tres cosas de configuración: (a) si el nombre del contacto trae @, números, correo o más de tres palabras, la plantilla sale sin nombre ("Hola, soy Louissiana…"), como quedó pedido el 7/10; (b) si el lead pregunta la fecha exacta, el agente no la inventa: "la fecha y la jornada te las confirma la asesora en la cita"; (c) confírmame que la base pasa por el Registro de Números Excluidos antes del envío?', size=9.5, color=GRIS, after=6)

h('2. Qué hace el agente con cada botón')
p('• Cualquier botón (Viajar, Trabajar o Estudiar) ES la respuesta al paso 1 ("¿para qué necesitas el inglés?"): el agente no la vuelve a preguntar, sigue con el paso 2, respeta la prueba A/B por último dígito del teléfono (par = 7 pasos, impar = flujo corto) y agenda la cita. '
  'Si el lead pregunta la fecha exacta, el agente responde: "La fecha y la jornada te las confirma la asesora en la cita; por eso la agendamos ya", y sigue al paso pendiente. No inventa fechas.')
p('• Si escribe "más adelante", "ahora no" o parecido: despedida corta y queda marcado para la reactivación de noviembre (la de la semana del 23). No se le insiste esta semana.')
p('• "Ya no", "no me interesa" o BAJA: despedida y se excluye de todas las reactivaciones.')
p('• Si responde con texto en vez de botón (negocio, para un hijo, una pregunta), el agente lo trata como respuesta al paso 1 y sigue. Si pregunta precio dos veces o "cuánto mensual", notificación a nuestra asesora, como ya está configurado.', after=6)

h('3. A quién se le envía (en este orden) y a quién no')
p('• Sí: todos los leads de la plataforma en estado cerrado o sin respuesta, incluidos los de la reactivación del 6 de octubre que no contestaron y los de pauta de octubre que se enfriaron.')
p('• No: leads con cita agendada o asistida, los que dijeron que no o pidieron BAJA, y cualquier conversación con mensajes en los últimos 7 días (esa la termina el agente, no se pisa).')
p('• Además te adjunto un archivo con 73 teléfonos que se excluyen de este envío: son leads que conversaron y se enfriaron, y los está trabajando por llamada nuestro gerente de alianzas desde hoy.')
p('• De la reserva del bot anterior (unos 1.500 leads que todavía no están en la plataforma) te subo hoy los 300 más calientes para que salgan en este mismo envío, si se puede como me dijiste.', after=6)

h('4. Envío y rescate')
p('• Hora: viernes 23 a las 9:00 a. m. Domingo no: nuestra regla de contacto es la de la ley, lunes a viernes de 7 a. m. a 7 p. m. y sábados hasta las 3 p. m., nunca domingos ni festivos. El fin de semana lo cubre el agente, que agenda él mismo, y nuestro equipo hace barridos el viernes y el sábado; lo que responda y no agende queda en una lista para llamarlo el lunes. Si Meta no aprueba las plantillas a tiempo, se corre al lunes 26 a las 9:00; no se manda en la tarde.')
p('• Nuestra asesora y nuestro gerente de alianzas hacen barridos el viernes a las 12:00 y a las 6:00 p. m. y el sábado a las 10:00 a. m. y a la 1:00 p. m., en el filtro de contacto humano y en las conversaciones que se quedaron a mitad de flujo. Lo que se cae a mitad de camino lo llama uno de ellos el mismo día.')
p('• El lunes miraríamos a ver cuántos respondieron por botón y por texto y cuántas citas, separado por plantilla (A y B) y aparte de pauta.', after=6)

h('5. Una regla nueva, que sale de un caso que probé en otra empresa')
p('Esto no nos ha pasado con nuestro asistente; me pasó a mí como cliente con el agente de otra empresa y se me quedó en la cabeza. Cuando el agente recibe algo que no está en el flujo (dos preguntas a la vez, un cálculo, una petición rara), nunca reinicia la conversación ni vuelve a pedir el nombre. Responde "eso te lo resuelve la asesora, te la paso ahora", manda la notificación y conserva el contexto. Un lead que ya dio toda su información y recibe "¿con quién tengo el gusto?" se pierde. Lo revisamos en la próxima reunión con el resto.')
p('Gracias. Confírmame cuando quede montada y la hora de salida.', after=0)

try: d.save(OUT); print('OK', OUT)
except PermissionError: d.save(OUT.replace('.docx', '_v2.docx')); print('ABIERTO -> _v2')
