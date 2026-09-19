# -*- coding: utf-8 -*-
"""Contrato de prestacion de servicios comerciales (rol: vendedora independiente)
+ 2 anexos economicos (cerradora high ticket / asesora comercial).
El contrato es ESTANDAR por rol; la plata vive solo en el anexo. BORRADOR pendiente abogado."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

GRIS = RGBColor(0x77, 0x77, 0x77)

RAIZ = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\contratos'

def nuevo_doc():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.6); s.bottom_margin = Cm(1.4); s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(10.5)
    return doc

def titulo(doc, txt):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(13)

def aviso(doc, txt):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = GRIS

def t(doc, txt, bold=False):
    p = doc.add_paragraph(); r = p.add_run(txt); r.font.size = Pt(10.5); r.bold = bold
    p.paragraph_format.space_after = Pt(4)

def clau(doc, nombre, cuerpo):
    p = doc.add_paragraph()
    r = p.add_run(nombre + ' '); r.bold = True; r.font.size = Pt(10.5)
    if cuerpo:
        r2 = p.add_run(cuerpo); r2.font.size = Pt(10.5)
    p.paragraph_format.space_after = Pt(4)

def num(doc, n, txt, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(f'{n}. '); r.bold = True; r.font.size = Pt(10.5)
    r2 = p.add_run(txt); r2.font.size = Pt(10.5); r2.bold = bold
    p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.space_after = Pt(3)

def firmas(doc):
    doc.add_paragraph(); doc.add_paragraph()
    t(doc, 'Para constancia se firma en Bucaramanga, a los ____ días del mes de _____________ de 2026.')
    doc.add_paragraph()
    t(doc, 'EL CONTRATANTE:                                                        EL (LA) CONTRATISTA:')
    doc.add_paragraph()
    t(doc, '_____________________________                                _____________________________')
    t(doc, 'DIANA MARCELA GÓMEZ RANGEL                              Nombre: ________________________')
    t(doc, 'C.C. 63.552.709 de Bucaramanga                                 C.C. ______________________________')
    t(doc, 'Representante Legal — GLOBAL TEACHER S.A.S.')

# ============================================================
# 1) CONTRATO ESTANDAR — PRESTACION DE SERVICIOS COMERCIALES
# ============================================================
doc = nuevo_doc()
titulo(doc, 'CONTRATO DE PRESTACIÓN DE SERVICIOS COMERCIALES INDEPENDIENTES — VERSIÓN 2026 v1')
aviso(doc, 'Contrato estándar por rol: la remuneración específica se pacta en el ANEXO ECONÓMICO, que hace parte integral de este contrato')
doc.add_paragraph()

t(doc, 'Entre los suscritos a saber: DIANA MARCELA GÓMEZ RANGEL, mayor de edad, identificada con cédula de ciudadanía No. 63.552.709 expedida en Bucaramanga, quien actúa en nombre de GLOBAL TEACHER S.A.S., sociedad legalmente constituida, con NIT 900.422.478-2, con domicilio principal en la ciudad de Bucaramanga, según consta en el Certificado de Existencia y Representación Legal, quien para efectos del presente contrato se denominará EL CONTRATANTE, y [NOMBRE COMPLETO], mayor de edad, identificado(a) con cédula de ciudadanía No. [___________], quien en adelante se denominará EL (LA) CONTRATISTA, hemos acordado celebrar el presente contrato de prestación de servicios así:')

clau(doc, 'CLÁUSULA PRIMERA. OBJETO:', 'El objeto del presente contrato lo constituye la prestación independiente de servicios comerciales por parte del (de la) CONTRATISTA, consistentes en la promoción y venta de los programas educativos de GLOBAL TEACHER S.A.S., utilizando exclusivamente los materiales, precios, becas y condiciones aprobados por gerencia y vigentes al momento de cada oferta.')

clau(doc, 'CLÁUSULA SEGUNDA. NATURALEZA DEL VÍNCULO:', 'El presente es un contrato de prestación de servicios de naturaleza civil/comercial, celebrado entre partes independientes. EL (LA) CONTRATISTA ejecuta el objeto con autonomía técnica y administrativa, sin subordinación, sin horario impuesto y sin exclusividad, respondiendo por ENTREGABLES y RESULTADOS conforme a este contrato y sus anexos. Nada en este contrato constituye relación laboral, y EL (LA) CONTRATISTA asume sus aportes al Sistema General de Seguridad Social como trabajador(a) independiente, presentando la planilla de pago como requisito para cada liquidación.')

clau(doc, 'CLÁUSULA TERCERA. REMUNERACIÓN:', 'La remuneración es EXCLUSIVAMENTE POR COMISIÓN sobre el recaudo efectivo a caja, en los porcentajes, aceleradores, bonos y cortes de liquidación pactados en el ANEXO ECONÓMICO suscrito por las partes, el cual hace parte integral de este contrato. Reglas invariables de toda comisión: (i) TODO pago de clientes ingresa única y directamente a las cuentas oficiales de GLOBAL TEACHER S.A.S. — EL (LA) CONTRATISTA no está autorizado(a) a recibir dinero de clientes en efectivo ni en cuentas propias, y hacerlo constituye causal de terminación inmediata; (ii) sin recaudo no hay comisión — cobrar también es vender; (iii) para reclamar cualquier comisión, el contacto debe estar registrado en la planilla comercial de la institución desde el primer contacto; en caso de concurrencia entre vendedores, la comisión corresponde a quien registró primero; (iv) las liquidaciones se pagan contra cuenta de cobro dentro de los CINCO (5) días hábiles siguientes a cada corte, respecto de pagos cuyo término legal de retracto ya esté vencido; (v) LA COMISIÓN SIGUE A LA PLATA: si un pago comisionado es objeto de retracto o devolución conforme a la ley, la comisión correspondiente se descuenta de la siguiente liquidación.')

clau(doc, 'CLÁUSULA CUARTA. CANALES Y ATRIBUCIÓN:', 'Las ventas se atribuyen por el PUNTO DE LLEGADA del lead: los contactos que llegan a los canales personales del (de la) CONTRATISTA — incluida la línea de WhatsApp que la institución le asigne para su operación conforme al ANEXO ECONÓMICO, que para efectos de atribución cuenta como canal propio — pertenecen a su canal propio; los que llegan a los demás flujos de captación de la institución (bot, redes y líneas generales de la casa) son de la casa, y solo comisionan como "leads entregados" cuando la institución los asigne expresamente en planilla. Los porcentajes de cada canal se pactan en el ANEXO ECONÓMICO. PAUTA: la institución podrá financiar y/o ejecutar campañas publicitarias sobre piezas del (de la) CONTRATISTA, previa aprobación escrita de pieza y presupuesto; la atribución de esas campañas se define en el ANEXO ECONÓMICO.')

clau(doc, 'CLÁUSULA QUINTA. REGLAS DE VENTA — LÍNEAS ROJAS:', 'En toda gestión comercial EL (LA) CONTRATISTA se obliga a: ofrecer únicamente los precios, becas y planes de pago de la matriz vigente aprobada por gerencia, sin descuentos ni condiciones inventadas; referirse siempre a "beca del Fondo" y jamás a patrocinios; no prometer empleo, visas, resultados garantizados distintos del texto exacto de la garantía institucional, ni facilidad o rapidez del programa; citar la garantía únicamente con su texto aprobado; someter a aprobación escrita de gerencia toda pieza publicitaria que use la marca o los programas de la institución ANTES de su publicación; y registrar cada gestión en la planilla comercial. SIN FACULTAD DE REPRESENTACIÓN: EL (LA) CONTRATISTA presenta, gestiona y entrega información, pero NO puede firmar convenios, otorgar becas o descuentos por fuera de la matriz, pactar condiciones económicas especiales ni comprometer jurídicamente a GLOBAL TEACHER S.A.S.; todo convenio lo suscribe únicamente el representante legal.')

clau(doc, 'CLÁUSULA SEXTA. ENTREGABLES:', 'EL (LA) CONTRATISTA reportará su gestión conforme a los formatos vigentes de la institución. Cuando las partes suscriban un ANEXO DE ACTIVIDAD (acuerdo de entregables mínimos), este hace parte integral del contrato y su incumplimiento reiterado y documentado constituye incumplimiento contractual conforme a la cláusula de terminación.')

clau(doc, 'CLÁUSULA SÉPTIMA. MATERIAL PROTEGIDO, PROPIEDAD INTELECTUAL Y CESIÓN DE OBRA:', 'Se entiende por "Material Protegido" todo material físico o digital que GLOBAL TEACHER S.A.S. ha desarrollado, adaptado o curado para su modelo pedagógico y comercial, incluyendo de forma enunciativa: kits y manuales de venta; guiones; planillas y formatos comerciales; matrices de precios y becas; presentaciones y cartas de alianzas; listas de contactos empresariales e institucionales; bases de datos de estudiantes, exalumnos y prospectos; estrategias, campañas y piezas; y cualquier otro material o información proporcionado por EL CONTRATANTE durante la vigencia del contrato. CESIÓN: toda pieza, guion, video, documento o material — comercial o publicitario — que EL (LA) CONTRATISTA cree, grabe, redacte o adapte en ejecución del presente contrato se entiende realizado por encargo de GLOBAL TEACHER S.A.S., y sus derechos patrimoniales de autor pertenecen exclusivamente a GLOBAL TEACHER S.A.S., conforme a la Ley 23 de 1982 y la Ley 1915 de 2018. EL (LA) CONTRATISTA reconoce expresamente que la totalidad del Material Protegido es propiedad intelectual exclusiva de GLOBAL TEACHER S.A.S.')

clau(doc, 'CLÁUSULA OCTAVA. DATOS PERSONALES:', 'Los datos de leads, prospectos, clientes y estudiantes gestionados en ejecución de este contrato son de titularidad y responsabilidad de GLOBAL TEACHER S.A.S. y se tratan conforme a la Ley 1581 de 2012. EL (LA) CONTRATISTA los usará únicamente para la gestión comercial de los programas de la institución, no los copiará a bases propias, y a la terminación del contrato cesará todo uso y los eliminará de sus dispositivos.')

clau(doc, 'CLÁUSULA NOVENA. NO COMERCIALIZACIÓN DE TERCEROS Y NO CONFLICTO:', 'EL (LA) CONTRATISTA no comercializará programas, cursos o servicios propios o de terceros (incluidos programas de estudios o trabajo en el exterior) a estudiantes, exalumnos, acudientes, prospectos o bases de datos de la institución, salvo acuerdo escrito de remisión firmado por el representante legal. Durante la vigencia del contrato, EL (LA) CONTRATISTA no promocionará ni venderá programas de formación en inglés de instituciones competidoras. Esta obligación NO limita su libertad de trabajar en otros sectores o productos no competidores.')

clau(doc, 'CLÁUSULA DÉCIMA. CONFIDENCIALIDAD POST-CONTRATO:', 'Las obligaciones de confidencialidad y propiedad intelectual de las cláusulas SÉPTIMA, OCTAVA y NOVENA (esta última en lo relativo a bases de datos) permanecerán vigentes durante un período mínimo de CINCO (5) AÑOS posteriores a la terminación del contrato, independientemente de la causa. Durante este período EL (LA) CONTRATISTA no reproducirá, divulgará, venderá ni usará el Material Protegido ni las bases de datos, ni los adaptará para presentarlos como propios o de terceros. Esto NO constituye cláusula de no competencia post-contrato: EL (LA) CONTRATISTA conserva plena libertad de ejercer su actividad comercial en cualquier sector, siempre que no utilice el Material Protegido aquí definido.')

clau(doc, 'CLÁUSULA DÉCIMA PRIMERA. DEVOLUCIÓN DE MATERIAL:', 'A la terminación del contrato por cualquier causa, EL (LA) CONTRATISTA devolverá físicamente todo el Material Protegido en un plazo máximo de tres (3) días hábiles, eliminará de sus dispositivos la totalidad de los archivos digitales del Material Protegido y de las bases de datos, y certificará por escrito el cumplimiento en el mismo plazo.')

clau(doc, 'CLÁUSULA DÉCIMA SEGUNDA. CLÁUSULA PENAL:', 'El incumplimiento de las obligaciones de confidencialidad, propiedad intelectual, datos personales y devolución de material dará derecho a EL CONTRATANTE a reclamar perjuicios y a exigir el pago de una cláusula penal escalonada según la gravedad: violación LEVE hasta CINCO (5) SMLMV; violación MODERADA hasta QUINCE (15) SMLMV; violación GRAVE hasta TREINTA (30) SMLMV, por cada violación documentada; a ejercer las acciones penales que correspondan (Ley 1273 de 2009, Ley 1915 de 2018 y demás normas aplicables) y a solicitar medidas cautelares. La calificación de gravedad la determinará EL CONTRATANTE con base en evidencia documentada, sin perjuicio del derecho del (de la) CONTRATISTA a controvertirla judicialmente.')

clau(doc, 'CLÁUSULA DÉCIMA TERCERA. VIGENCIA Y TERMINACIÓN:', 'El presente contrato rige por UN (1) AÑO desde su suscripción, prorrogable por acuerdo escrito. SIN PERMANENCIA: por tratarse de una prestación de servicios independiente, CUALQUIERA de las partes podrá darlo por terminado EN CUALQUIER MOMENTO, de manera inmediata y sin necesidad de preaviso, mediante simple comunicación escrita (incluido mensaje de datos), sin que ello genere indemnización, sanción o pago alguno distinto de las comisiones ya causadas sobre recaudos efectivos anteriores a la terminación, que se liquidan en el corte siguiente conforme a la cláusula TERCERA. La terminación no extingue las obligaciones de confidencialidad, propiedad intelectual, datos personales y devolución de material (cláusulas SÉPTIMA a DÉCIMA SEGUNDA). También termina por mutuo acuerdo, fuerza mayor, vencimiento del término o imposibilidad del (de la) CONTRATISTA.')

clau(doc, 'CLÁUSULA DÉCIMA CUARTA. DOMICILIO CONTRACTUAL:', 'Para todos los efectos legales, se tendrá como domicilio contractual la ciudad de Bucaramanga, Colombia.')

clau(doc, 'CLÁUSULA DÉCIMA QUINTA. ACEPTACIÓN:', 'EL (LA) CONTRATISTA declara haber leído íntegramente el presente contrato y su ANEXO ECONÓMICO, comprenderlos en su totalidad, y aceptar todas y cada una de sus cláusulas, en particular las relativas a manejo de dineros, confidencialidad, propiedad intelectual, cesión de obra, datos personales, devolución de material y cláusula penal.')

firmas(doc)
doc.save(RAIZ + r'\CONTRATO_PRESTACION_COMERCIAL_v1.docx')
print('OK contrato')

# ============================================================
# 2) ANEXO ECONOMICO — CERRADORA HIGH TICKET
# ============================================================
doc = nuevo_doc()
titulo(doc, 'ANEXO ECONÓMICO No. ___ — ROL: CERRADORA HIGH TICKET · CAMPAÑA ARRANQUE A1+A2')
aviso(doc, 'Hace parte integral del Contrato de Prestación de Servicios Comerciales suscrito entre las partes · Septiembre 2026')
doc.add_paragraph()
t(doc, 'Las partes pactan la siguiente remuneración, que reemplaza cualquier acuerdo verbal anterior:')
num(doc, 1, 'CANAL PROPIO — 8%: comisión del OCHO por ciento (8%) del recaudo efectivo a caja por toda venta originada en el canal propio del (de la) CONTRATISTA: su contenido, su red personal y las campañas de pauta que la institución financie sobre sus piezas. La pauta la financia y la ejecuta la institución desde sus cuentas publicitarias, sobre piezas del (de la) CONTRATISTA aprobadas por gerencia antes de publicar, dirigiendo los leads a la línea de WhatsApp Business de la institución operada por el (la) CONTRATISTA; esos leads son canal propio. Aplica a todo el portafolio (Arranque, programas completos, niveles individuales y sabatinos).', bold=True)
num(doc, 2, 'ACELERADOR — 10% RETROACTIVO: la comisión sube al DIEZ por ciento (10%) sobre la TOTALIDAD de la serie de ventas del (de la) CONTRATISTA en esta campaña si se cumple cualquiera de las dos metas: (a) DOCE (12) matrículas de la cohorte del 5 de octubre dentro de los primeros DOCE (12) DÍAS de campaña al aire (el día 1 es la fecha de publicación de la primera campaña, registrada en la planilla); o (b) los veinte (20) cupos de lanzamiento vendidos antes del 5 DE OCTUBRE de 2026. No cumplida ninguna, aplica el 8%. Las metas cuentan solo matrículas de la cohorte del 5 de octubre.', bold=True)
num(doc, 3, 'BONO GRUPOS LLENOS: veintiocho (28) matrículas (14 y 14 por jornada) antes del 5 DE OCTUBRE de 2026 = bono único de DOS MILLONES DE PESOS ($2.000.000), adicional al acelerador.')
num(doc, 4, 'LEADS ENTREGADOS POR LA INSTITUCIÓN: 5% del recaudo de cliente nuevo y 2,5% de cliente antiguo, solo sobre leads asignados expresamente en planilla. Todo lead asignado se atiende el mismo día y queda registrado en la planilla desde el primer contacto. PASE CRUZADO: quien pase un cliente registrado a otro miembro del equipo recibe el UNO por ciento (1%) del recaudo si este lo cierra — en ambas direcciones, siempre con el pase anotado en planilla.')
num(doc, 5, 'PAUTA: presupuesto inicial de hasta UN MILLÓN DE PESOS ($1.000.000) mensuales para la campaña del (de la) CONTRATISTA, aprobado pieza a pieza por gerencia y ajustable según el costo por matrícula de cada campaña. El gasto de pauta lo asume la institución y no se descuenta de las comisiones. Las piezas que el (la) CONTRATISTA grabe para su propia campaña no se remuneran por pieza: su retribución es la comisión de este anexo.')
num(doc, 6, 'CORTES DE LIQUIDACIÓN: 30 de septiembre y 5 de octubre de 2026; en adelante, cortes mensuales al último día del mes. Pago conforme a la cláusula TERCERA del contrato (cuenta de cobro, 5 días hábiles, retracto vencido, clawback).')
num(doc, 7, 'VENTAS DEL CUPO 21 EN ADELANTE: comisionan igual, a la tarifa plena vigente, con la única excepción escrita de pago de contado completo el mismo día (precio de lanzamiento). Prohibido ofrecer precio de lanzamiento con el contador en 20, salvo esa excepción.')
firmas(doc)
doc.save(RAIZ + r'\ANEXO_ECONOMICO_CERRADORA_v1.docx')
print('OK anexo cerradora')

# ============================================================
# 3) ANEXO ECONOMICO — ASESORA COMERCIAL (PROSPECCION)
# ============================================================
doc = nuevo_doc()
titulo(doc, 'ANEXO ECONÓMICO No. ___ — ROL: ASESORA COMERCIAL · PROSPECCIÓN')
aviso(doc, 'Hace parte integral del Contrato de Prestación de Servicios Comerciales suscrito entre las partes')
doc.add_paragraph()
t(doc, 'Las partes pactan la siguiente remuneración, que reemplaza cualquier acuerdo verbal anterior:')
num(doc, 1, 'COMISIÓN: CINCO por ciento (5%) del recaudo efectivo a caja por venta a cliente nuevo, y DOS PUNTO CINCO por ciento (2,5%) por venta a cliente antiguo. Aplica a toda venta gestionada por el (la) CONTRATISTA e inscrita en planilla, incluidas las ventas corporativas y de alianzas que gestione.', bold=True)
num(doc, 2, 'FRENTE DE TRABAJO: prospección en los canales definidos en el plan comercial vigente (base propia de la institución, exalumnos, referidos, redes, empresas y colegios). Los leads de los flujos de captación de la casa (bot y líneas de la institución) solo comisionan cuando gerencia los asigne expresamente en planilla.')
num(doc, 3, 'ANEXO DE ACTIVIDAD: el acuerdo de entregables mínimos suscrito entre las partes hace parte integral del contrato (cláusula SEXTA). Su incumplimiento reiterado y documentado es causal de terminación conforme a la cláusula DÉCIMA TERCERA.', bold=True)
num(doc, 4, 'CORTES DE LIQUIDACIÓN: mensuales, al último día de cada mes. Pago conforme a la cláusula TERCERA del contrato (cuenta de cobro, 5 días hábiles, retracto vencido, clawback).')
num(doc, 5, 'APOYOS NO SALARIALES: cualquier auxilio de conectividad o transporte que la institución otorgue es voluntario, temporal, no constitutivo de remuneración y condicionado al cumplimiento del ANEXO DE ACTIVIDAD del período.')
firmas(doc)
doc.save(RAIZ + r'\ANEXO_ECONOMICO_ASESORA_v1.docx')
print('OK anexo asesora')
