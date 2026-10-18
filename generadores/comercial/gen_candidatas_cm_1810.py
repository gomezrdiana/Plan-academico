# -*- coding: utf-8 -*-
"""Lectura de las dos postulaciones nuevas al cargo de Community Manager (18/10/2026) y mensajes de siguiente paso
(misma prueba práctica pagada del brief). Salida: recursos/comercial/contratacion/COMMUNITY_MANAGER/CANDIDATAS_CM_SIGUIENTE_PASO_2026-10-18.docx"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAR = RGBColor(0xE5, 0x4C, 0x09); GRIS = RGBColor(0x66, 0x66, 0x66)
d = Document(); d.styles['Normal'].font.name = 'Calibri'; d.styles['Normal'].font.size = Pt(10)
for s in d.sections: s.left_margin = s.right_margin = Cm(1.8); s.top_margin = Cm(1.5); s.bottom_margin = Cm(1.3)
def shd(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear'); el.set(qn('w:color'), 'auto'); el.set(qn('w:fill'), color); tcPr.append(el)
def p(t, bold=False, color=None, size=10, after=4):
    x = d.add_paragraph(); r = x.add_run(t); r.bold = bool(bold); r.font.size = Pt(size); x.paragraph_format.space_after = Pt(after)
    if color is not None: r.font.color.rgb = color

p('COMMUNITY MANAGER — DOS POSTULACIONES NUEVAS (18/10/2026)', True, NAR, 13)
p('Uso interno de gerencia. Lectura de lo que enviaron y mensajes de siguiente paso. La prueba práctica es la misma del brief (guion "Siete meses", desde casa, $100.000), para comparar con la rúbrica de 7 criterios.', False, GRIS, 9.5)

tb = d.add_table(rows=1, cols=3); tb.style = 'Table Grid'
for i, h in enumerate(['', 'Luna', 'Astrid']):
    c = tb.rows[0].cells[i]; shd(c, 'E54C09'); r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
filas = [
 ('Qué envió', 'No envió hoja de vida: una presentación, un portafolio en Canva y un video en Drive.', 'Mensaje completo, portafolio web, dos videos publicados (Instagram y YouTube). Comunicadora social.'),
 ('Disponibilidad', 'Lunes a sábado, mañana, tarde y noche. Vive en Bucaramanga (San Francisco).', 'Lunes a sábado solo en la mañana hasta mediodía. Vive en Lagos, Floridablanca.'),
 ('Equipo y programas', 'iPhone 17 Pro Max; CapCut, Canva, herramientas de IA.', 'Samsung S24 FE, micrófono y cámara Nikon Coolpix; CapCut, Canva, imágenes con ChatGPT o Gemini.'),
 ('Hoy trabaja para', 'No lo dice.', 'Tres cuentas activas (motos, celulares, una marca local). Sería una cuarta cuenta, no un cargo de dedicación.'),
 ('Cómo lograría un testimonio natural', 'Conversación cercana, preguntas sencillas, sin memorizar respuestas. Es exactamente lo que pedimos.', 'Sonrisa, mostrar ejemplos, y ofrecer salir solo con voz o con un avatar de IA. El avatar va contra la regla: los testimonios son con la persona real, cara y voz.'),
 ('Lo que falta ver', 'Hoja de vida con fechas, empleadores y referencias; abrir el portafolio y el video.', 'Cómo graba en la noche si solo tiene mañanas; si escribe guiones; referencias de las tres cuentas.'),
 ('Lectura', 'La que mejor encaja en disponibilidad y equipo. Falta todo lo verificable.', 'Perfil más formado, pero disponibilidad corta y el reflejo de resolver con IA lo que debe ser real.'),
]
for a, b, c3 in filas:
    row = tb.add_row().cells
    for i, txt in enumerate((a, b, c3)):
        r = row[i].paragraphs[0].add_run(txt); r.font.size = Pt(9); r.bold = (i == 0)
for row in tb.rows:
    row.cells[0].width = Cm(3.2); row.cells[1].width = Cm(7.0); row.cells[2].width = Cm(7.2)
d.add_paragraph()
p('Recomendación: la prueba práctica a las dos, con el mismo brief y el mismo plazo, y se comparan las entregas con la rúbrica. A Luna se le pide antes la hoja de vida con referencias; a Astrid se le pregunta de frente por las noches y por las tres cuentas. Ninguna decisión por el mensaje de postulación: la decisión la da la prueba.', True)

p('MENSAJE PARA LUNA (adjuntar BRIEF_PRUEBA_PRACTICA_CM)', True, NAR, 11)
p("""Hola Luna, gracias por tu postulación, por el portafolio y por el video. Nos gustó tu respuesta sobre los testimonios: conversación cercana y preguntas sencillas es exactamente como los hacemos aquí.

Para avanzar te pedimos dos cosas:
1. Tu hoja de vida con fechas (desde y hasta) de cada trabajo, y dos referencias: nombre, cargo y teléfono de un jefe o cliente para quien hayas hecho contenido o manejado redes.
2. ¿Escribes los guiones de los videos que produces? Si tienes uno, compártenoslo.

Y el siguiente paso, que puedes empezar ya: una prueba práctica pagada de $100.000. Ejecutas desde tu casa uno de nuestros guiones de anuncio, con las mismas reglas de entrega que usa nuestro videógrafo. Te adjunto el brief completo; ahí verás que nuestro mensaje es distinto al habitual: no vendemos que el inglés sea fácil ni divertido, vendemos entrenamiento con garantía por escrito. Se paga igual si pasas o no. ¿Me confirmas si te animas y para cuándo la tendrías?

Diana Gómez
Gerente General, Heiiu English Academy""")

p('MENSAJE PARA ASTRID (adjuntar BRIEF_PRUEBA_PRACTICA_CM)', True, NAR, 11)
p("""Hola Astrid, gracias por tu postulación, por el portafolio y por los videos. Se nota oficio.

Antes de avanzar, dos preguntas de frente, porque de eso depende que el cargo te sirva:
1. Parte de lo que se graba es en las clases de la noche (6:30 a 8:30) y los sábados en la mañana. Tu disponibilidad es hasta mediodía: ¿podrías cubrir al menos dos noches por semana en la academia?
2. Hoy manejas tres cuentas. Este cargo cubre dos marcas y pide presencia en la academia varias veces por semana. ¿Cómo lo ves en tiempo?

Una aclaración sobre los testimonios: aquí siempre salen con la persona real, su cara y su voz, porque es lo que respalda la garantía que firmamos; no usamos avatares de IA para eso. Lo que sí hacemos es acompañar al estudiante hasta que se sienta cómodo.

Si las dos respuestas encajan, el siguiente paso es una prueba práctica pagada de $100.000: ejecutar desde casa uno de nuestros guiones de anuncio, con las mismas reglas de entrega que usa nuestro videógrafo. Te adjunto el brief completo. Se paga igual si pasas o no. ¿Me confirmas?

Diana Gómez
Gerente General, Heiiu English Academy""")

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\contratacion\COMMUNITY_MANAGER\CANDIDATAS_CM_SIGUIENTE_PASO_2026-10-18.docx'
try: d.save(out); print('OK', out)
except PermissionError: d.save(out.replace('.docx', '_v2.docx')); print('ABIERTO -> _v2')
