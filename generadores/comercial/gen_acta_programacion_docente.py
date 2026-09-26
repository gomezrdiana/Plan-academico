# -*- coding: utf-8 -*-
"""Acta interna de programacion academica: media hoja por curso, firmada por las partes, que desarrolla el contrato docente v3.1."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

GRIS = RGBColor(0x77, 0x77, 0x77)
doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.6); s.bottom_margin = Cm(1.4); s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)
doc.styles['Normal'].font.name = 'Calibri'; doc.styles['Normal'].font.size = Pt(10.5)

def shd(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear'); el.set(qn('w:color'), 'auto'); el.set(qn('w:fill'), color); tcPr.append(el)

def t(txt, bold=False, after=4, size=10.5):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(txt); r.font.size = Pt(size); r.bold = bold

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('ACTA INTERNA DE PROGRAMACIÓN ACADÉMICA No. ______'); r.bold = True; r.font.size = Pt(13)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Desarrolla la Cláusula Segunda literal (a) del Contrato de Prestación de Servicios Profesionales — Docente Hora Cátedra v3.1 · Una acta por curso'); r.font.size = Pt(9); r.font.color.rgb = GRIS
doc.add_paragraph()

t('Entre GLOBAL TEACHER S.A.S. (Heiiu English Academy), NIT 900.422.478-2, representada por DIANA MARCELA GÓMEZ RANGEL, y el (la) docente ______________________________________, C.C. ________________, en desarrollo del contrato de prestación de servicios suscrito el ____ / ____ / ______, se acuerda la siguiente programación:')

tb = doc.add_table(rows=8, cols=2); tb.style = 'Table Grid'
filas = [
    ('Nivel y grupo', ''),
    ('Modalidad / jornada', '□ Intensivo (4 h)   □ Nocturno (2 h)   □ Sabatino   □ Otro: ____________'),
    ('Días y horario', ''),
    ('Fecha de inicio', '____ / ____ / ______'),
    ('Fecha estimada de cierre', '____ / ____ / ______'),
    ('Total de horas del curso', '__________ horas'),
    ('Valor de la hora cátedra', '$ 17.000 (según contrato)'),
    ('Material Protegido entregado con esta acta', 'Guías de clase del nivel (solo la hoja del día, según política) · formato de reporte · otros: ______________'),
]
for i, (a, b) in enumerate(filas):
    c0, c1 = tb.rows[i].cells
    c0.width = Cm(5.5); c1.width = Cm(11.5)
    r0 = c0.paragraphs[0].add_run(a); r0.bold = True; r0.font.size = Pt(10); shd(c0, 'FFF8E7')
    r1 = c1.paragraphs[0].add_run(b); r1.font.size = Pt(10)

doc.add_paragraph()
t('Condiciones de esta acta:', bold=True, after=3)
for n, txt in enumerate([
    'Esta acta no modifica el contrato: define el curso, el horario y las horas de este encargo. Su cierre no termina el contrato, que sigue vigente para los cursos siguientes.',
    'El horario aquí acordado es una condición del servicio contratado. Cualquier cambio, suspensión o reprogramación se acuerda por escrito con coordinación con la antelación de las políticas del servicio.',
    'Las horas se pagan por mensualidades vencidas, contra cuenta de cobro y planilla de seguridad social, según las horas efectivamente dictadas y registradas en los reportes de clase.',
    'Las horas adicionales (refuerzos, recuperaciones, evaluaciones adicionales) se pagan al mismo valor, previo acuerdo escrito.',
    'Al cierre del curso, el (la) docente devuelve el Material Protegido recibido y las partes firman el cierre de esta acta.',
], 1):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f'{n}. '); r.bold = True; p.add_run(txt)

doc.add_paragraph()
t('Firmada en Bucaramanga el ____ / ____ / ______', after=10)
t('EL CONTRATANTE                                                        EL (LA) CONTRATISTA', after=14)
t('_____________________________                                _____________________________', after=2)
t('DIANA MARCELA GÓMEZ RANGEL                              Nombre: ________________________', after=2)
t('GLOBAL TEACHER S.A.S.                                              C.C. ______________________________', after=10)
p = doc.add_paragraph(); r = p.add_run('CIERRE DE LA ACTA (al terminar el curso):  fecha ____ / ____ / ______  ·  horas dictadas ________  ·  material devuelto  □ Sí   ·   Firmas: ______________  /  ______________'); r.font.size = Pt(9); r.font.color.rgb = GRIS

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\contratos\ACTA_PROGRAMACION_ACADEMICA_DOCENTE.docx'
doc.save(out); print('OK', out)
