# -*- coding: utf-8 -*-
"""Hoja de una pagina para quien llama y vende desde gerencia: lo que no se puede saltar ni olvidar.
Fuente de doctrina: KIT_DE_VENTA_ASESORAS v2, ESTRATEGIA_VENTAS_2026, CONDICIONES_ECONOMICAS_CERRADORA."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NARANJA = RGBColor(0xE5, 0x4C, 0x09)
GRIS = RGBColor(0x66, 0x66, 0x66)

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.0); s.bottom_margin = Cm(0.9); s.left_margin = Cm(1.4); s.right_margin = Cm(1.4)
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(9)

def shd(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement('w:shd'); el.set(qn('w:val'), 'clear'); el.set(qn('w:color'), 'auto'); el.set(qn('w:fill'), color)
    tcPr.append(el)

def titulo(txt, sub):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = NARANJA
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(4)
    r = p.add_run(sub); r.font.size = Pt(8.5); r.font.color.rgb = GRIS

def sec(txt):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(1)
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(10); r.font.color.rgb = NARANJA

def item(txt, bold_prefix=None):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.left_indent = Cm(0.3)
    if bold_prefix:
        r = p.add_run(bold_prefix + ' '); r.bold = True; r.font.size = Pt(9)
    r = p.add_run(txt); r.font.size = Pt(9)

def tabla(filas, anchos):
    t = doc.add_table(rows=len(filas), cols=len(filas[0])); t.style = 'Table Grid'
    for i, fila in enumerate(filas):
        for j, v in enumerate(fila):
            c = t.cell(i, j); c.width = Cm(anchos[j])
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            r = c.paragraphs[0].add_run(v); r.font.size = Pt(8.5)
            if i == 0:
                r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shd(c, 'E54C09')
            elif j == 0:
                r.bold = True; shd(c, 'FFF8E7')

titulo('HOJA DE LLAMADAS — GERENCIA',
       'Una página · lo que no se salta y no se olvida en ninguna llamada ni asesoría · Uso interno · Septiembre 2026')

sec('1. PARA QUÉ ES LA LLAMADA')
item('La llamada vende la ASESORÍA, no el programa. Termina de una de tres formas: cita agendada con fecha y hora · cupo separado con abono · descartado con motivo anotado. Nunca "me llama luego".', 'Meta:')
item('"Hola [nombre], soy Diana, la directora de Heiiu. Vi que te interesaste en el inglés con nosotros — ¿tienes dos minutos? Te hago tres preguntas y te digo si esto es para ti."', 'Apertura:')
item('¿Has estudiado inglés antes y cómo te fue? · ¿Para qué lo necesitas: montar algo tuyo, irte a trabajar afuera, tu trabajo actual? · ¿Eres afiliado a Cajasan? (se anota siempre).', 'Las 3 preguntas:')
item('"Te agendo la asesoría el [día] a las [hora]. Antes, haz este test gratis de 50 minutos: efset.org/ef-set-50 — llegamos sabiendo tu nivel exacto." Sin cita agendada en la llamada, la llamada no valió.', 'Cierre de la llamada:')

sec('2. QUÉ VENDES — Y CÓMO SE DICE')
item('No vendemos inglés: vendemos entrenamiento con garantía firmada + un destino (emprender, irse, ascender). Jamás prometer fácil, rápido ni divertido. Jamás prometer nativos, empleo ni visas.', 'El Pacto:')
item('garantía por escrito · avance medido (EF SET, audios, reportes) · institución con licencia e ICONTEC · formación del carácter.', 'Las 4 anclas:')
item('"Beca del Fondo" — NUNCA "descuento" ni "patrocinador". Yo no decido la beca: "el Fondo la evalúa y audita las fechas". El porcentaje no se negocia: una excepción lo tumba para todas las familias.', 'Vocabulario:')
item('0 vendidos: "acaban de abrirse los 20 cupos" · 1–4: "los cupos ya se están asignando" · 5+: el número real. Jamás inventar el número.', 'El contador:')

sec('3. LOS NÚMEROS QUE HAY QUE SABERSE')
tabla([
    ['Producto', 'Precio', 'Lo que se dice'],
    ['ARRANQUE A1+A2 (200 h)', '$2.990.000 · 20 cupos · inicia 5 de octubre', 'Precio de lanzamiento, sin beca. Jornadas 8–12 AM y 6:30–8:30 PM. Grupos de 13–14, máx 16. Desde el cupo 21: $3.511.000 salvo contado el mismo día.'],
    ['PROGRAMA COMPLETO A1→B2', '$8.696.000 full · beca por perfil', 'Transformación (estrato 1–2 con soporte) 45% · General 32% · Ejecutivo (paga empresa) 15%. General de contado: $5.913.280. Cuotas: motor de cuotas, nunca de memoria.'],
    ['Niveles sueltos', 'A1 $1.596.000 · A2 $1.915.000 · B1 $2.842.000 · B2 $2.343.000', 'Para quien "ya sabe algo": sesión de ubicación con EF SET previo, y entra al nivel que le corresponde.'],
], [4.2, 5.2, 8.6])
item('el que ya intentó y fracasó, o el que compara precio → Arranque. El que va por la meta completa → programa con beca. Las dos respuestas son SÍ.', 'Cierre doble:')

sec('4. LA PLATA — REGLAS QUE NO SE ROMPEN')
item('transferencia, débito, o efectivo SOLO en la sede con recibo. Nada de efectivo por fuera ni cuentas personales.', 'Contado:')
item('cesantías y crédito de un tercero (Comultrasan, banco, financiera) = CONTADO: abono de $300.000 hoy, desembolso en máximo 10 días hábiles y antes del inicio. Si no llega: paga por otro medio o el abono queda como saldo a favor. Nunca devolución en efectivo.', 'Cuenta como contado:')
item('cuenta como contado, pero el valor sube 5% por el costo de la transacción — se dice ANTES de pasarla.', 'Tarjeta de crédito:')
item('el último recurso: 1,9% mensual sobre saldo, inicial mínimo 20%, siempre con pagaré firmado.', 'Crédito directo Heiiu:')
item('toda beca sale por escrito con vencimiento a 72 horas. El abono de $300.000 congela cupo y precio 7 días. Si venció, venció.', 'Vigencia:')

sec('5. MATRÍCULA EN DOS MOMENTOS')
item('abono a cuentas de la academia + el cliente escribe "ACEPTO las condiciones del recibo No. X" con foto de cédula + cita de ubicación agendada en el mismo acto. Con eso el cupo está separado.', 'Separación (por chat, válida):')
item('contrato firmado en físico en la sede, a más tardar el primer día de clases. "El chat asegura el cupo; la sede firma el contrato."', 'Formalización:')

sec('6. LO QUE NO SE PUEDE OLVIDAR — CHECKLIST')
item('□ Registrar el contacto en la planilla desde la primera llamada (canal, qué se ofreció, próximo paso, motivo real si no cierra).   □ Preguntar Cajasan.   □ Pedir el EF SET antes de la cita.   □ Anunciar las dos jornadas y la regla del grupo mínimo ANTES de firmar.   □ La frase del expediente: garantía y audio diario se firman juntos.   □ Una frase de cierre y SILENCIO: el primero que habla, concede.   □ Si no cierra hoy: oferta por escrito con vencimiento + abono ofrecido + seguimiento pautado (mismo día resumen → +24h mensaje → +48h llamada → día 3 "vence hoy").')

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(3)
r = p.add_run('Fuente: Kit de Venta v2 · Estrategia de Ventas 2026 · Lo que no está aquí, está en el Kit. Lo que no está en el Kit, se pregunta antes de prometerlo.')
r.font.size = Pt(7.5); r.font.color.rgb = GRIS

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\INDUCCION_COMERCIAL\HOJA_LLAMADAS_GERENCIA.docx'
doc.save(out); print('OK', out)
