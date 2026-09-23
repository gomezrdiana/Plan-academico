# -*- coding: utf-8 -*-
"""Borrador contrato docente v3: hora catedra por actas + gestion de alianzas. Base: v2 firmado (27/04/2026)."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

GRIS = RGBColor(0x77, 0x77, 0x77)
ROJO = RGBColor(0xC0, 0x00, 0x00)

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(2.2); s.bottom_margin = Cm(2.0); s.left_margin = Cm(2.4); s.right_margin = Cm(2.4)
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(10.5)

def titulo(txt):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(13)

def aviso(txt):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(10); r.font.color.rgb = ROJO

def t(txt, bold=False):
    p = doc.add_paragraph(); r = p.add_run(txt); r.bold = bold; r.font.size = Pt(10.5)
    p.paragraph_format.space_after = Pt(5)
    return p

def clau(nombre, cuerpo):
    p = doc.add_paragraph()
    r = p.add_run(nombre + ' '); r.bold = True; r.font.size = Pt(10.5)
    r2 = p.add_run(cuerpo); r2.font.size = Pt(10.5)
    p.paragraph_format.space_after = Pt(5)

def num(n, txt, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(str(n) + '. ' + txt); r.font.size = Pt(10.5); r.bold = bold
    p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.space_after = Pt(3)

titulo('CONTRATO DE PRESTACIÓN DE SERVICIOS PROFESIONALES — DOCENTE HORA CÁTEDRA — VERSIÓN 2026 v3.1')
aviso('Contrato estándar para docentes hora cátedra · La habilitación comercial, cuando aplique, se pacta en anexo aparte')
doc.add_paragraph()

t('Entre los suscritos a saber: DIANA MARCELA GÓMEZ RANGEL, mayor de edad, identificada con cédula de ciudadanía No. 63.552.709 expedida en Bucaramanga, quien actúa en nombre de GLOBAL TEACHER S.A.S., sociedad legalmente constituida, con NIT 900.422.478-2, con domicilio principal en la ciudad de Bucaramanga, según consta en el Certificado de Existencia y Representación Legal, quien para efectos del presente contrato se denominará EL CONTRATANTE, y [NOMBRE COMPLETO DEL DOCENTE], mayor de edad, vecino(a) de la ciudad de Bucaramanga, identificado(a) con cédula de ciudadanía No. [___________], quien en adelante se denominará EL CONTRATISTA, hemos acordado celebrar el presente contrato de prestación de servicios así:')

clau('CLÁUSULA PRIMERA. OBJETO:', 'El objeto del presente contrato lo constituye la prestación independiente de servicios profesionales de docencia de inglés, en la modalidad de HORA CÁTEDRA, en los cursos que las partes acuerden mediante actas de programación académica. De manera accesoria y facultativa, EL CONTRATISTA podrá referir a la institución personas interesadas en sus programas, en los términos de la Cláusula Segunda literal (b).')

clau('CLÁUSULA SEGUNDA. OBJETO ESPECÍFICO Y VALOR CONTRACTUAL:', '(a) DOCENCIA: EL CONTRATISTA dictará los cursos que las partes acuerden mediante ACTA INTERNA DE PROGRAMACIÓN ACADÉMICA suscrita por ambas, en la cual se indicará para cada curso: nivel, número de horas, horario y fechas. El valor de la hora cátedra es de DIECISIETE MIL PESOS COP ($17.000), pagadero por mensualidades vencidas contra cuenta de cobro y planilla de seguridad social, de acuerdo con las horas efectivamente impartidas, con las retenciones de ley a que haya lugar. (b) REFERIDOS — FACULTATIVO: EL CONTRATISTA PODRÁ, si así lo decide, referir a la institución personas interesadas en sus programas. Esta actividad no constituye obligación contractual ni su omisión genera incumplimiento. Por cada referido que se matricule, EL CONTRATISTA recibirá el DOS por ciento (2%) del recaudo efectivo a caja de ese estudiante, siempre que: (i) el contacto haya sido registrado por EL CONTRATISTA en la planilla comercial de la institución antes de cualquier gestión de venta; (ii) la asesoría y el cierre los realice el equipo comercial de la institución; y (iii) en caso de concurrencia, el contacto no haya sido registrado antes por otra persona. Sin recaudo no hay comisión, y si un pago comisionado es objeto de retracto o devolución, la comisión se descuenta de la siguiente liquidación. (c) HABILITACIÓN COMERCIAL: la conducción de asesorías y el cierre de ventas por parte del CONTRATISTA, con sus comisiones correspondientes, solo procede cuando las partes suscriban un ANEXO DE HABILITACIÓN COMERCIAL, previa inducción en los materiales y condiciones aprobados por gerencia. Sin dicho anexo, EL CONTRATISTA no ofrece precios, becas ni condiciones a ningún interesado: lo remite al equipo comercial.')
t('PARÁGRAFO I. El presente contrato se mantiene vigente por el término de UN (1) AÑO contado desde su suscripción, prorrogable por acuerdo escrito. La terminación o culminación de un curso individual NO termina el contrato: solo cierra el acta correspondiente. Cualquiera de las partes podrá terminar el contrato con PREAVISO ESCRITO de TREINTA (30) días calendario; tratándose de cursos en ejecución, EL CONTRATISTA acompañará la transición del curso durante ese período conforme a las guías de la institución.', bold=True)
t('PARÁGRAFO II. Si las actividades académicas se extienden más allá de las horas inicialmente pactadas en el acta (reprogramaciones, refuerzos, recuperación de festivos, evaluaciones adicionales, etc.), las horas adicionales se pagarán al mismo valor de hora cátedra establecido, previo acuerdo escrito entre las partes.')
t('PARÁGRAFO III. Cualquier cambio de horario, suspensión de clase o reprogramación deberá ser comunicado a la administración con la antelación que esta determine en sus políticas internas.')
t('PARÁGRAFO IV — NATURALEZA DEL VÍNCULO Y SIN FACULTAD DE REPRESENTACIÓN: el presente es un contrato de prestación de servicios de naturaleza civil, celebrado entre partes independientes. EL CONTRATISTA ejecuta el servicio con autonomía técnica y administrativa, sin subordinación ni exclusividad, dentro de los estándares de calidad y los horarios de los cursos acordados en cada acta, que son condiciones propias del servicio contratado y no órdenes de un empleador. EL CONTRATISTA NO tiene facultad para firmar convenios, otorgar becas o descuentos, pactar condiciones económicas ni comprometer jurídicamente a GLOBAL TEACHER S.A.S.', bold=True)

clau('CLÁUSULA TERCERA. PERFECCIONAMIENTO:', 'El presente contrato se entiende perfeccionado desde la fecha de suscripción.')

clau('CLÁUSULA CUARTA: OBLIGACIONES DEL CONTRATISTA:', '')
obls = [
 'Prestar el servicio de docencia con sus conocimientos y capacidad profesional, conforme a los estándares de calidad pedagógica de GLOBAL TEACHER S.A.S. descritos en este contrato y en las guías de clase que hacen parte del servicio contratado.',
 'Guardar en el desempeño de sus funciones y fuera de ellas la discreción y la confidencialidad que exige la lealtad que debe a EL CONTRATANTE.',
 'Prestar sus servicios como docente en las clases de inglés asignadas según los horarios disponibles por ambas partes.',
 'Respetar y no usar el Material Protegido (definido en la CLÁUSULA QUINTA) por fuera de las instalaciones de GLOBAL TEACHER SAS, ya que se asume total confidencialidad y derechos de autor del material usado.',
 'No cortejar o iniciar ningún tipo de relación sentimental con algún estudiante del Instituto de Idiomas GLOBAL TEACHER SAS.',
 'Acatar las políticas de seguridad de la institución.',
 'Responder por los elementos, herramientas, máquinas, muebles, enseres, materiales, materias primas, etc., de trabajo que EL CONTRATANTE le entregue para el desempeño de su cargo.',
 'Conservar y entregar oportunamente todos los documentos, informes, razones, envíos que, con destino a EL CONTRATANTE, haya recibido de terceros en el ejercicio de sus funciones.',
 'Guardar absoluta reserva sobre los hechos, documentos físicos y/o electrónicos, informaciones y, en general, sobre todos los asuntos y materias que lleguen a su conocimiento por causa o con ocasión de su contrato, lo estipulado en la Ley 1266 de 2008 de Habeas Data y las normas que la reglamenten o modifiquen.',
 'Cumplir con las obligaciones descritas en este contrato y con las condiciones del servicio acordadas en cada acta de programación.',
 'Presentarse a las clases con una presentación personal acorde con el entorno educativo, y usar en clase los distintivos institucionales que la academia le facilite, cuando los facilite.',
 'Iniciar cada clase a la hora acordada en el acta. Como condición del servicio, EL CONTRATISTA llegará a la sede al menos DIEZ (10) MINUTOS ANTES del inicio de cada clase para revisar la guía del día, preparar el tablero y los materiales, y recibir la información de coordinación necesaria para la sesión.',
 'Participar en las actividades del Sistema de Gestión de Seguridad y Salud en el Trabajo que la ley exige para contratistas, y en las reuniones de coordinación académica que se acuerden entre las partes.',
 'Realizar los aportes al Sistema General de Seguridad Social y Parafiscales de Ley y presentar la planilla de pago como requisito para el pago mensual.',
 'Ejecutar el servicio conforme a la guía del día y los protocolos pedagógicos de GLOBAL TEACHER S.A.S., que constituyen la especificación técnica del servicio contratado: rituales de apertura y cierre, seguimiento de tareas y entrega de reportes. Las modificaciones al contenido o al método se acuerdan por escrito con coordinación.',
 'Dictar todas las clases acordadas en el acta. Cualquier imposibilidad, retraso o necesidad de reprogramación se notifica a coordinación con la antelación acordada en las políticas del servicio. La inasistencia injustificada o reiterada constituye incumplimiento del contrato conforme a la Cláusula Décima Tercera.',
 'Entregar reporte escrito y firmado al final de cada clase, conforme al formato vigente entregado por coordinación, incluyendo: asistencia, tareas verificadas, errores observados y entregables del día, junto con el registro fotográfico del paquete de evidencia que defina coordinación. Este reporte es la base del sistema de auditoría académica de la institución.',
 'Cumplir con el protocolo de manejo de errores y evaluaciones: los errores observados en clase se registran de forma anónima en el papel físico de errores que se entrega a coordinación al final del día. El registro nominal con citas literales se conserva en la libreta privada del profesor exclusivamente para coordinación. EL CONTRATISTA NO comunica a los estudiantes resultados de evaluaciones, midterms, finales, criterios de calificación ni notas; esa comunicación es competencia exclusiva de coordinación.',
 'Respetar la propiedad intelectual de las metodologías y rituales pedagógicos de GLOBAL TEACHER S.A.S. No mencionar a estudiantes los nombres técnicos de las metodologías aplicadas. La identificación operativa de los bloques pedagógicos se utiliza únicamente al interior de la institución y se considera Material Protegido conforme a la Cláusula Quinta.',
 'Mantener absoluta confidencialidad sobre la información de los estudiantes (datos personales, evaluaciones, comentarios de coordinación, planes individuales) conforme a la Ley 1581 de 2012 de Protección de Datos Personales.',
 'EN LOS REFERIDOS: no prometer empleo, visas, prácticas, resultados, precios, becas ni condiciones; registrar el contacto en la planilla comercial y remitirlo al equipo comercial; y tratar las listas de contactos y las condiciones comerciales como Material Protegido.',
 'NO COMERCIALIZAR programas, cursos o servicios propios o de terceros (incluidos programas de estudios en el exterior) a estudiantes, exalumnos, acudientes o bases de datos de la institución, salvo acuerdo escrito de remisión firmado por el representante legal. Toda pieza publicitaria del canal propio del CONTRATISTA que use la marca o los programas de la institución requiere aprobación escrita de gerencia ANTES de su publicación. El aula jamás se usa para actividad comercial.',
]
for i, o in enumerate(obls, 1):
    num(i, o, bold=(i in (22, 23)))
t('PARÁGRAFO V. Las modificaciones, adiciones al presente contrato se harán de mutuo acuerdo entre las partes contratantes, previo documento que lo justifique suscrito entre el contratante y el docente.')

clau('CLÁUSULA QUINTA: DEFINICIÓN DE "MATERIAL PROTEGIDO", PROPIEDAD INTELECTUAL Y CESIÓN DE OBRA:', 'Para efectos del presente contrato, se entiende por "Material Protegido" todo material físico o digital que GLOBAL TEACHER S.A.S. ha desarrollado, adaptado o curado para su modelo pedagógico y comercial, incluyendo de forma enunciativa pero no taxativa: a) guías de clase impresas o digitales; b) manuales del profesor; c) anexos pedagógicos; d) plantillas de reportes, autochequeos, calendarios de virtudes y formatos de evaluación; e) sistemas y rituales pedagógicos y cualquier bloque, ritual, secuencia, formato o procedimiento pedagógico desarrollado por la institución; f) la secuencia operativa de los bloques de clase, su orden, sus tiempos y la manera específica en que se hilan entre sí; g) las metodologías académicas integradas en el modelo, su selección, adaptación y combinación; h) bases estructurales y planes de clase de todos los niveles; i) bases de datos de estudiantes, métricas pedagógicas y reportes de progreso; j) los materiales comerciales y de alianzas (cartas, hojas de condiciones, presentaciones, listas de contactos empresariales e institucionales, matrices de precios y becas); k) cualquier otro material o información proporcionado por EL CONTRATANTE durante la vigencia del contrato. CESIÓN: todo material, documento, carta, presentación o pieza — pedagógica o comercial — que EL CONTRATISTA cree, redacte o adapte en ejecución del presente contrato se entiende realizado por encargo de GLOBAL TEACHER S.A.S., y sus derechos patrimoniales de autor pertenecen exclusivamente a GLOBAL TEACHER S.A.S., conforme a la Ley 23 de 1982 y la Ley 1915 de 2018. EL CONTRATISTA reconoce expresamente que la totalidad del Material Protegido es propiedad intelectual exclusiva de GLOBAL TEACHER S.A.S.')

clau('CLÁUSULA SEXTA: VIGENCIA DE LAS OBLIGACIONES DE CONFIDENCIALIDAD Y PROPIEDAD INTELECTUAL POST-CONTRATO:', 'Las obligaciones de confidencialidad, derechos de autor y propiedad intelectual establecidas en las CLÁUSULAS CUARTA y QUINTA permanecerán vigentes durante un período mínimo de CINCO (5) AÑOS posteriores a la terminación del presente contrato, independientemente de la causa de terminación. Durante este período, EL CONTRATISTA se obliga a: no utilizar la metodología ni los sistemas de la institución fuera de GLOBAL TEACHER S.A.S.; no reproducir, replicar, enseñar, transferir, capacitar a terceros, vender, licenciar, publicar ni divulgar el Material Protegido; no adaptarlo para presentarlo como propio o de un tercero; no utilizarlo como base para ningún producto, curso, academia, plataforma o servicio educativo propio o de un tercero; y no nombrar a estudiantes ni a terceros los sistemas internos de la institución. Esta vigencia post-contractual aplica únicamente a las obligaciones de confidencialidad y propiedad intelectual sobre el Material Protegido y NO constituye una cláusula de no competencia post-empleo. EL CONTRATISTA conserva plena libertad de ejercer su profesión como docente de inglés en otras instituciones, siempre que NO utilice el Material Protegido aquí definido.')

clau('CLÁUSULA SÉPTIMA: DEVOLUCIÓN Y ELIMINACIÓN DE MATERIAL AL TERMINAR EL CONTRATO:', 'A la terminación del contrato por cualquier causa, EL CONTRATISTA se obliga a devolver físicamente todo el Material Protegido en un plazo máximo de tres (3) días hábiles, a eliminar de todos sus dispositivos personales la totalidad de los archivos digitales del Material Protegido, y a certificar por escrito el cumplimiento de lo anterior en el mismo plazo. El incumplimiento activará la cláusula penal de la CLÁUSULA OCTAVA.')

clau('CLÁUSULA OCTAVA: CLÁUSULA PENAL POR VIOLACIÓN DE CONFIDENCIALIDAD Y PROPIEDAD INTELECTUAL:', 'El incumplimiento de las obligaciones de confidencialidad, propiedad intelectual y devolución de material dará derecho a EL CONTRATANTE a reclamar perjuicios y a exigir el pago de una cláusula penal escalonada según la gravedad: violación LEVE hasta CINCO (5) SMLMV; violación MODERADA hasta QUINCE (15) SMLMV; violación GRAVE hasta TREINTA (30) SMLMV, por cada violación documentada; a ejercer las acciones penales que correspondan (Ley 1273 de 2009, Ley 1915 de 2018 y demás normas aplicables) y a solicitar medidas cautelares. La calificación de la gravedad se hará entre las partes de buena fe, con base en evidencia documentada; a falta de acuerdo, la determinará el juez competente.')

clau('CLÁUSULA NOVENA: OBLIGACIONES DE LA EMPRESA:', 'Pagar al contratista el valor del presente contrato por mensualidades vencidas; facilitar los espacios físicos y los materiales requeridos; escuchar al CONTRATISTA antes de tomar cualquier decisión sobre incumplimientos; hacer el cierre de cuentas que corresponda al término del contrato; entregar formalmente copia firmada de este contrato y dejar constancia escrita de cada pieza de Material Protegido entregada (con fecha y descripción); y liquidar oportunamente las comisiones por referidos contra los ingresos efectivos a caja.')

clau('CLÁUSULA DÉCIMA: SUPERVISIÓN Y RESPONSABILIDAD:', 'La verificación del cumplimiento del servicio docente la realiza coordinación académica; la de los referidos, gerencia. Se llevará registro escrito de la entrega de Material Protegido y de su devolución.')

clau('CLÁUSULA DÉCIMA PRIMERA: AUTORIZACIÓN DE USO DE IMAGEN Y VOZ:', 'EL CONTRATISTA autoriza a GLOBAL TEACHER S.A.S. a captar, fijar y usar su imagen y su voz en fotografías, videos y grabaciones realizadas durante las clases, eventos y actividades de la institución, y a publicarlas en los canales de la academia (redes sociales, sitio web, material comercial y pedagógico) con fines institucionales y publicitarios, sin remuneración adicional, durante la vigencia del contrato y hasta CINCO (5) AÑOS después de su terminación respecto del material producido durante la vigencia. La institución se abstendrá de usar la imagen del CONTRATISTA de manera que afecte su honra o buen nombre. Esta autorización se otorga conforme a la Ley 1581 de 2012.')

clau('CLÁUSULA DÉCIMA SEGUNDA: CAPACIDAD CONTRACTUAL:', 'EL CONTRATISTA declara, bajo la gravedad del juramento, que tiene capacidad jurídica para contratar y que no se encuentra incurso(a) en ninguna inhabilidad o incompatibilidad.')

clau('CLÁUSULA DÉCIMA TERCERA: CAUSALES DE TERMINACIÓN:', 'Mutuo acuerdo; unilateralmente por parte de GLOBAL TEACHER SAS cuando EL CONTRATISTA incumpla cualquiera de las obligaciones de este contrato (la violación de las cláusulas de confidencialidad, propiedad intelectual o devolución de material será causal de terminación inmediata, sin perjuicio de las acciones legales); fuerza mayor o caso fortuito; vencimiento del término establecido; o imposibilidad física, jurídica o profesional del CONTRATISTA.')

clau('CLÁUSULA DÉCIMA CUARTA: DOMICILIO CONTRACTUAL.', 'Para todos los efectos legales, se tendrá como domicilio contractual la ciudad de Bucaramanga, Colombia.')

clau('CLÁUSULA DÉCIMA QUINTA: ACEPTACIÓN.', 'EL CONTRATISTA declara haber leído íntegramente el presente contrato, comprenderlo en su totalidad, y aceptar todas y cada una de sus cláusulas, en particular las relativas a confidencialidad, propiedad intelectual, cesión de obra, devolución de material, cláusula penal y uso de imagen.')

doc.add_paragraph()
t('Para constancia se firma en la ciudad de BUCARAMANGA, el día ______ de ____________ de 2026.')
for _ in range(2): doc.add_paragraph()
t('EL CONTRATANTE'); doc.add_paragraph()
t('__________________________________')
t('DIANA MARCELA GÓMEZ RANGEL · GLOBAL TEACHER S.A.S. · NIT 900.422.478-2')
doc.add_paragraph()
t('EL CONTRATISTA'); doc.add_paragraph()
t('__________________________________')
t('[NOMBRE] · C.C. [___________]')

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\contratos\CONTRATO_DOCENTE_HORA_CATEDRA_v3_1.docx'
doc.save(out); print('OK', out)
