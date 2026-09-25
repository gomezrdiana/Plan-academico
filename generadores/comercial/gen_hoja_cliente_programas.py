# -*- coding: utf-8 -*-
"""Hoja para el CLIENTE (la envia la asesora despues de la cita): los dos caminos lado a lado, una pagina, formato de marca."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NARANJA = RGBColor(0xE5, 0x4C, 0x09); GRIS = RGBColor(0x66, 0x66, 0x66); BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.1); s.bottom_margin = Cm(0.9); s.left_margin = Cm(1.5); s.right_margin = Cm(1.5)
doc.styles['Normal'].font.name = 'Calibri'; doc.styles['Normal'].font.size = Pt(9.5)

def shd(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear'); el.set(qn('w:color'), 'auto'); el.set(qn('w:fill'), color); tcPr.append(el)

def run(p, txt, size=9.5, bold=False, color=None, italic=False):
    r = p.add_run(txt); r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = color
    return r

def cell_par(cell, txt, size=9.5, bold=False, color=None, align=None, space=2, first=False):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    p.paragraph_format.space_after = Pt(space); p.paragraph_format.space_before = Pt(0)
    if align: p.alignment = align
    run(p, txt, size, bold, color)
    return p

def bullet(cell, bold_txt, txt):
    p = cell.add_paragraph(); p.paragraph_format.space_after = Pt(1.5); p.paragraph_format.left_indent = Cm(0.25)
    run(p, '• ' + bold_txt, 9.5, True); run(p, txt, 9.5)

# ---- Encabezado
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
run(p, 'HEIIU ENGLISH ACADEMY', 10, True, GRIS)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
run(p, 'Dos caminos para hablar inglés de verdad', 18, True, NARANJA)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(6)
run(p, 'Presencial en Bucaramanga · grupos de máximo 16 · la única academia de la ciudad con garantía de aprendizaje por escrito', 9.5, False, GRIS)

# ---- Tabla de dos columnas
tb = doc.add_table(rows=2, cols=2); tb.style = 'Table Grid'; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
for j in range(2):
    tb.columns[j].width = Cm(9.0)
    for i in range(2): tb.cell(i, j).width = Cm(9.0)

# Encabezados
h1, h2 = tb.cell(0, 0), tb.cell(0, 1)
shd(h1, 'E54C09'); shd(h2, 'E54C09')
cell_par(h1, 'CAMINO 1 · ARRANQUE A1 + A2', 12, True, BLANCO, WD_ALIGN_PARAGRAPH.CENTER, 0, first=True)
cell_par(h1, 'Empieza hoy y decide después', 9, False, BLANCO, WD_ALIGN_PARAGRAPH.CENTER, 2)
cell_par(h2, 'CAMINO 2 · PROGRAMA COMPLETO A1 → B2', 12, True, BLANCO, WD_ALIGN_PARAGRAPH.CENTER, 0, first=True)
cell_par(h2, 'Comprométete con la meta completa', 9, False, BLANCO, WD_ALIGN_PARAGRAPH.CENTER, 2)

# Columna 1: Arranque
c = tb.cell(1, 0); shd(c, 'FFF8E7')
cell_par(c, '$2.990.000', 20, True, NARANJA, WD_ALIGN_PARAGRAPH.CENTER, 0, first=True)
cell_par(c, 'Precio de lanzamiento · 20 cupos · cohorte del 5 de octubre', 8.5, False, GRIS, WD_ALIGN_PARAGRAPH.CENTER, 6)
bullet(c, '200 horas presenciales: ', 'los dos primeros niveles completos.')
bullet(c, 'Libros incluidos ', 'y certificado oficial por cada nivel aprobado.')
bullet(c, 'Hablas desde el primer día: ', 'clases activas, en inglés, con simulaciones de la vida real: presentarte, tu trabajo, llamadas, entrevistas.')
bullet(c, 'Jornadas: ', 'mañana 8:00 a 12:00 (dos meses y medio) o noche 6:30 a 8:30 PM (cinco meses).')
bullet(c, 'Garantía por escrito ', 'en tu contrato.')
bullet(c, 'Módulo de Graduación incluido, a tu elección al terminar: ', '')
p = c.add_paragraph(); p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.space_after = Pt(1)
run(p, 'Emprendedor: ', 9, True); run(p, 'sales con tu página web publicada, tu pitch en inglés y tu primer mensaje de venta enviado a un cliente real.', 9)
p = c.add_paragraph(); p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.space_after = Pt(3)
run(p, 'Pasaporte: ', 9, True); run(p, 'sales con tu video de presentación, tu hoja de vida en inglés y la entrevista ensayada para programas como Au Pair o Work and Travel.', 9)
bullet(c, 'Al terminar, ', 'si quieres seguir a B1 y B2, continúas con el Fondo de Becas.')

# Columna 2: Completo
c = tb.cell(1, 1); shd(c, 'FFFFFF')
cell_par(c, 'Beca del Fondo según tu perfil', 14, True, NARANJA, WD_ALIGN_PARAGRAPH.CENTER, 0, first=True)
cell_par(c, 'Precio pleno $8.696.000 · tu beca se define en la asesoría', 8.5, False, GRIS, WD_ALIGN_PARAGRAPH.CENTER, 6)
bullet(c, '575 horas presenciales: ', 'los cuatro niveles, de cero a nivel profesional.')
bullet(c, 'Libros incluidos ', 'y certificado oficial por cada nivel aprobado.')
bullet(c, 'Todo lo del Arranque, ', 'y además B1 y B2: presentaciones, reuniones, negociación, entrevistas difíciles, pitch de 90 segundos, pedir un aumento, manejar un cliente molesto.')
bullet(c, 'Garantía por escrito ', 'en cada nivel.')
bullet(c, 'Módulo de Graduación al terminar A2 ', '(Emprendedor o Pasaporte).')
bullet(c, 'Refuerzo PRO al terminar B2: ', 'pitch avanzado, negociación real en inglés y tu video del antes y después.')
bullet(c, 'La beca más alta: ', 'si eres afiliado a Cajasan, entras directo a la beca más alta del Fondo con tu carné.')
bullet(c, 'Un solo compromiso, un solo precio: ', 'comprar los cuatro niveles juntos cuesta más de un millón menos que comprarlos uno a uno.')

# ---- Garantía (caja)
doc.add_paragraph().paragraph_format.space_after = Pt(2)
g = doc.add_table(rows=1, cols=1); g.style = 'Table Grid'; g.alignment = WD_TABLE_ALIGNMENT.CENTER
gc = g.cell(0, 0); gc.width = Cm(18.0); shd(gc, 'FFF8E7')
cell_par(gc, 'LA GARANTÍA HEIIU, TAL COMO QUEDA EN TU CONTRATO', 10, True, NARANJA, WD_ALIGN_PARAGRAPH.CENTER, 1, first=True)
p = gc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(1)
run(p, '"Si cumples — asistencia, tareas, evaluaciones — y no avanzas, te devolvemos el 100% del nivel. Por escrito, en el contrato."', 10.5, True, None, True)
p = gc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(2)
run(p, 'Es un trato: nosotros ponemos el método, el profesor y el respaldo. Tú pones el trabajo: asistir, hacer las tareas y grabar tu video corto de práctica todos los días. No te prometemos que sea fácil. Te prometemos que funciona si haces tu parte.', 9, False, GRIS)

# ---- Formas de pago y siguiente paso
doc.add_paragraph().paragraph_format.space_after = Pt(2)
t2 = doc.add_table(rows=1, cols=2); t2.style = 'Table Grid'; t2.alignment = WD_TABLE_ALIGNMENT.CENTER
a, b = t2.cell(0, 0), t2.cell(0, 1); a.width = Cm(9.0); b.width = Cm(9.0)
cell_par(a, 'CÓMO PUEDES PAGAR', 10, True, NARANJA, None, 2, first=True)
for bt, tx in [('Contado: ', 'transferencia, tarjeta débito o crédito (sin recargo), o efectivo en la sede.'),
               ('Cesantías o crédito de tu cooperativa o banco: ', 'cuentan como pago de contado.'),
               ('A cuotas: ', 'con una inicial y mensualidades; el plan exacto se arma en tu asesoría.'),
               ('Afiliado a Cajasan: ', 'condiciones especiales por nuestro convenio, la única academia de inglés de la ciudad que lo tiene.')]:
    bullet(a, bt, tx)
cell_par(b, 'TU SIGUIENTE PASO', 10, True, NARANJA, None, 2, first=True)
for bt, tx in [('Asegura tu cupo con $300.000: ', 'te congela el cupo y el precio por 7 días, y se cruza con tu primer pago.'),
               ('Tienes 5 días hábiles ', 'para cambiar de opinión con devolución total, por ley.'),
               ('La matrícula se formaliza en la sede: ', 'Carrera 27 # 48-49, segundo piso, Sotomayor.'),
               ('Cohorte del Arranque: ', '5 de octubre. Los 20 cupos de lanzamiento se asignan en orden de llegada.')]:
    bullet(b, bt, tx)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(5)
run(p, 'Tu asesora: Paula Saenz · WhatsApp 315 547 0657 · Instagram @Heiiu_english · Heiiu English Academy, Bucaramanga', 8.5, False, GRIS)

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\INDUCCION_COMERCIAL\HOJA_CLIENTE_DOS_CAMINOS.docx'
doc.save(out); print('OK', out)
