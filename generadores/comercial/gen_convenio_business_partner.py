# -*- coding: utf-8 -*-
"""Convenio Heiiu Business Partner (1 pagina, formato estandar que cierra el docente habilitado y firma gerencia).
Salida: recursos/comercial/B2B/CONVENIO_HEIIU_BUSINESS_PARTNER_v1.docx"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

GRIS = RGBColor(0x77, 0x77, 0x77)
d = Document()
for s in d.sections:
    s.top_margin = Cm(1.3); s.bottom_margin = Cm(1.1); s.left_margin = Cm(1.9); s.right_margin = Cm(1.9)
d.styles['Normal'].font.name = 'Calibri'; d.styles['Normal'].font.size = Pt(9.5)

def t(txt, bold=False, after=3, size=9.5, align=None, color=None):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    if align: p.alignment = align
    r = p.add_run(txt); r.bold = bold; r.font.size = Pt(size)
    if color: r.font.color.rgb = color

def cl(n, tit, txt):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f'{n}. {tit}. '); r.bold = True; r.font.size = Pt(9.5)
    r2 = p.add_run(txt); r2.font.size = Pt(9.5)

t('CONVENIO DE ALIANZA — HEIIU BUSINESS PARTNER', bold=True, size=12.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
t('Iniciativa Go Above and Beyond · Heiiu English Academy — Global Teacher S.A.S., NIT 900.422.478-2 · Formato estándar v1 · octubre de 2026', size=8.5, color=GRIS, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

t('Entre GLOBAL TEACHER S.A.S. (Heiiu English Academy), representada por su representante legal, y la empresa ______________________________________________, NIT ____________________, representada por ______________________________________, cargo ____________________ (LA EMPRESA ALIADA), se acuerda:', after=6)

cl(1, 'OBJETO', 'LA EMPRESA ALIADA se vincula como Heiiu Business Partner para que sus colaboradores accedan a los beneficios de este convenio en los programas de inglés de la academia. El convenio no tiene costo para LA EMPRESA ALIADA ni le genera obligaciones económicas.')
cl(2, 'BENEFICIOS PARA LOS COLABORADORES', '(a) Todo colaborador que presente carta de postulación de LA EMPRESA ALIADA accede a la columna Transformación del Fondo de Becas Heiiu (45% de contado o 36% a cuotas sobre la tarifa de lista del programa completo; proporcional en los demás paquetes), conforme al Reglamento del Fondo, artículo 3.5. (b) BECA [EMPRESA]: hasta tres colaboradores por trimestre calendario, elegidos por LA EMPRESA ALIADA, reciben el 45% aunque paguen a cuotas, más los libros del programa y prioridad de cupo; escalera: el cupo base es de tres becas por trimestre; sube a cinco para el trimestre siguiente cuando todos los becados activos de LA EMPRESA ALIADA cumplieron (asistencia mínima del 80% y pagos al día); si alguno incumple, vuelve a tres; nunca pasa de cinco ni se acumula. (c) Los beneficios no se acumulan entre sí ni con otras becas: se aplica el mejor al que el colaborador califique. (d) El beneficio puede extenderse al cónyuge o a un hijo desde los 12 años del colaborador (menor de edad con acudiente), con la misma carta de postulación. (e) Si LA EMPRESA ALIADA tiene fundación o programa de bienestar o responsabilidad social, las Becas [Empresa] pueden asignarse a beneficiarios de esa fundación o programa, incluidas las familias de la comunidad que atiende, con carta de postulación firmada por la fundación. Las becas asignadas a familiares o a beneficiarios de la fundación cuentan dentro del mismo cupo trimestral y cumplen las mismas condiciones.')
cl(3, 'BENEFICIOS PARA LA EMPRESA ALIADA', 'Entrega de las Becas [Empresa] con certificado a nombre de LA EMPRESA ALIADA (en la empresa o en la sede); reporte periódico del avance académico de sus colaboradores becados (asistencia y cumplimiento, nunca notas); y un diagnóstico de nivel de inglés para su equipo, hasta diez personas, sin costo. Opcionalmente, acceso a la vitrina Heiiu Student Partner de egresados destacados, sin obligación de contratación.')
cl(4, 'COMPROMISOS DE LA EMPRESA ALIADA', 'Designar un contacto de recursos humanos; emitir las cartas de postulación de los colaboradores que elija; y difundir el beneficio internamente. LA EMPRESA ALIADA no recibe ni gestiona pagos: cada colaborador contrata y paga su programa directamente con la academia, bajo el contrato de matrícula y la garantía de aprendizaje por escrito de Heiiu.')
cl(5, 'LO QUE NO INCLUYE', 'Este convenio no otorga exclusividad, no constituye intermediación laboral ni promete empleo, visas ni resultados a ningún colaborador. Si LA EMPRESA ALIADA desea pagar directamente el programa de sus colaboradores, aplica el Plan Empresa: tarifa corporativa por cupo, de 3 a 7 cupos con el formato estándar (ORDEN PLAN EMPRESA), y de 8 en adelante o con condiciones distintas, por acuerdo con la gerencia de la academia.')
cl(6, 'DATOS Y CONFIDENCIALIDAD', 'La academia trata los datos de los colaboradores conforme a la Ley 1581 de 2012, únicamente para los fines de este convenio. El reporte de avance a LA EMPRESA ALIADA se limita a asistencia y cumplimiento, con autorización del colaborador en su contrato de matrícula.')
cl(7, 'VIGENCIA', 'Doce (12) meses desde la firma, renovables por acuerdo escrito. Cualquiera de las partes puede terminarlo con aviso escrito de treinta (30) días, sin afectar los beneficios ya otorgados a colaboradores matriculados. Se perfecciona con la firma del representante legal de la academia.')

t('Contacto de recursos humanos designado: ______________________________  Cargo: ____________________  Teléfono / correo: ______________________________', after=6)
t('Firmado en Bucaramanga el ____ / ____ / 2026', after=10)
t('POR LA EMPRESA ALIADA                                                              POR HEIIU — GLOBAL TEACHER S.A.S.', after=14)
t('_____________________________                                          _____________________________', after=1)
t('Nombre: ________________________                                        DIANA MARCELA GÓMEZ RANGEL', after=1)
t('C.C. ______________________________                                       Representante Legal', after=4)
t('Presentado por: ______________________________ (docente habilitado Heiiu) · Fecha de presentación: ____ / ____ / 2026', size=8.5, color=GRIS)

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\B2B\CONVENIO_HEIIU_BUSINESS_PARTNER_v1.docx'
d.save(out); print('OK', out)
