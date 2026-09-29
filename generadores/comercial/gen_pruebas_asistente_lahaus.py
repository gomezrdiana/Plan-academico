# -*- coding: utf-8 -*-
"""Bateria de pruebas para el asistente LaHaus (ronda 2): que escribir, que debe responder, como se ve el fallo.
Uso interno de gerencia. Salida: recursos/comercial/plataforma LaHaus/PRUEBAS_ASISTENTE_RONDA2.docx"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAR = RGBColor(0xE5, 0x4C, 0x09); GRIS = RGBColor(0x66, 0x66, 0x66)
doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width, sec.page_height = sec.page_height, sec.page_width
sec.top_margin = Cm(1.3); sec.bottom_margin = Cm(1.2); sec.left_margin = Cm(1.5); sec.right_margin = Cm(1.5)
doc.styles['Normal'].font.name = 'Calibri'; doc.styles['Normal'].font.size = Pt(9.5)

def shd(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear'); el.set(qn('w:color'), 'auto'); el.set(qn('w:fill'), color); tcPr.append(el)

def h(txt, size=13, color=NAR, after=4):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = color

def t(txt, bold=False, after=4, size=9.5, color=None):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(txt); r.font.size = Pt(size); r.bold = bold
    if color: r.font.color.rgb = color

def tabla(filas, anchos, header=True):
    tb = doc.add_table(rows=len(filas), cols=len(filas[0])); tb.style = 'Table Grid'
    for i, fila in enumerate(filas):
        for j, txt in enumerate(fila):
            c = tb.rows[i].cells[j]; c.width = Cm(anchos[j])
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            r = p.add_run(txt); r.font.size = Pt(8.5)
            if header and i == 0:
                r.bold = True; shd(c, 'FFF8E7')
            elif j == 0:
                r.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

h('PRUEBAS DEL ASISTENTE HEIIU EN LAHAUS — RONDA 2', 15)
t('Uso: gerencia escribe al número de prueba como si fuera un cliente, en este orden. Al lado de cada prueba se marca ✔ (pasó) o ✘ (falló) y se copia la frase exacta que falló. Cada ✘ se convierte en un cambio numerado del documento CAMBIOS_ASISTENTE_LAHAUS. Documento interno: no se envía a LaHaus.', color=GRIS, after=8)

h('A. LO QUE SE REVISA EN TODOS LOS MENSAJES (transversal)', 11.5)
tabla([
    ['Regla', 'Falla si…', 'Cambio'],
    ['Se presenta y habla como Heiiu. "Global Teacher" solo si preguntan por factura o razón social.', 'Dice "Global Teacher" en cualquier otro momento.', '23'],
    ['UNA sola pregunta por mensaje, y en el orden del diagnóstico: motivo → nivel → disponibilidad → programa → cita.', 'Dos preguntas en un mensaje, o pregunta horario antes de saber el motivo.', '8'],
    ['Vende antes de agendar: no ofrece cita hasta que el cliente sabe qué programa le sirve y por qué.', 'Ofrece "agendemos" en los primeros 2-3 mensajes.', '12'],
    ['Nunca pide el número de celular (ya lo tiene por WhatsApp). Pide nombre, no teléfono.', 'Pide "tu número de contacto".', '—'],
    ['Solo ofrece horarios que están en el calendario de la asesora. No inventa franjas.', 'Ofrece una hora que no existe o dice "cualquier hora".', '—'],
    ['No describe el método en cada respuesta y nunca dice "de pie" ni nombres de metodologías.', '"aprender de pie", "método X", o repite el discurso del método 3 veces.', '1'],
    ['Emojis: máximo uno, y no en cada mensaje. Cierra con una pregunta útil, no con "¿te interesa?".', 'Emojis en cadena, cierre vacío.', '6'],
    ['Nunca dice "descuento". Habla de beca (Cajasan / Transformación), contado o cuotas.', 'Aparece la palabra descuento.', '3, 5'],
    ['Nunca promete profesores nativos ni virtualidad.', '"tenemos nativos", "también virtual".', '—'],
    ['Nunca escribe nada sobre recargo de tarjeta (ni que hay, ni que no hay).', 'Cualquier frase con "recargo".', '22'],
], [11.5, 9.5, 2.0])

h('B. CONVERSACIÓN 1 — LA CLIENTE IDEAL (una sola conversación, seguida, en este orden)', 11.5)
t('Personaje: ingeniera industrial, 32 años, quiere presentar proyectos en inglés en su empresa, no sabe nada de inglés, trabaja de oficina 8 a 5.', color=GRIS)
tabla([
    ['#', 'Escribe esto', 'Debe responder', 'Falla si…', 'Cambio'],
    ['1', 'Hola, vi el anuncio, quiero información', 'Saluda como Heiiu, pide el nombre o pregunta el MOTIVO. Una sola pregunta. Sin precios.', 'Manda precios, lista de programas, o dos preguntas.', '8'],
    ['2', 'Quiero trabajar', 'Profundiza: ¿trabajar en qué, dónde, en cuánto tiempo? Proyecta ("cuando presentes tu primer proyecto en inglés…").', 'Salta a nivel o a cita sin entender el destino.', '9, 16'],
    ['3', 'Soy ingeniera industrial, quiero presentar proyectos en inglés en mi empresa. No sé nada de inglés', 'Desde cero = NO ofrece prueba de ubicación; ofrece asesoría. Propone el PROGRAMA COMPLETO (A1 a B2 con el destino) como primer camino.', 'Manda al EF SET, o arranca ofreciendo el Arranque, o pregunta horario ya.', '11, 12'],
    ['4', '¿Y eso cuánto vale?', 'Precio del programa completo en contexto (qué incluye, garantía). Contado primero, cuotas después. Cajasan / Transformación como respaldo. Sin "descuento".', 'Da solo el número frío, o empieza por cuotas, o inventa un precio, o dice descuento.', '3, 4, 5'],
    ['5', 'Uy, muy caro', 'No baja el precio ni pide disculpas. Reencuadra valor (entrenamiento + destino + garantía). Solo AHORA baja al Arranque como segundo peldaño, con su precio.', 'Ofrece descuento, o ya había ofrecido el Arranque antes, o se disculpa.', '4, 12'],
    ['6', 'Es que en Smart me dan 3 años para terminar y más barato', 'No nombra a Smart ni la desprestigia. Contrasta: aquí se termina en X meses con garantía y video diario; "más tiempo para terminar" no es ventaja.', 'Dice "Smart", habla mal de otra academia, o cede.', '10'],
    ['7', '¿Puedo tomar una clase de prueba gratis?', 'No hay clase de cortesía. Explica por qué en una frase y ofrece la asesoría / visita a la sede.', 'Promete clase gratis o "déjame preguntar".', '25'],
    ['8', '¿Qué experiencia tienen enseñando inglés?', 'Una sola fecha: desde 2012, con licencia. Cifras institucionales coherentes con INFO_INSTITUCIONAL.', 'Dice 2017, dos fechas distintas, o inventa números.', '24'],
    ['9', '¿Y cómo son las clases?', 'Entrenamiento: 4 bloques largos, simulaciones profesionales, video diario, virtudes, feedback diario. Corto.', '"de pie", nombres de metodologías, discurso de 10 líneas.', '1'],
    ['10', '¿Qué se ve en A1?', 'Temas reales del nivel (los de la sección 10 de la base). Corto.', 'Inventa temas o dice "todo lo básico".', '7'],
    ['11', 'Vi que tienen malos reviews', 'Responde sin ponerse a la defensiva; garantía firmada como prueba; invita a verlo en la sede.', 'Niega, ataca a quien opinó, o ignora.', '13'],
    ['12', 'Bueno, gracias, lo voy a pensar', 'No suelta: pregunta qué le falta para decidir, ofrece resolverlo, propone fecha concreta. Sin presión burda.', '"Perfecto, cualquier cosa me escribes" y se acaba.', '14'],
    ['13', 'Tengo que hablarlo con mi esposo', 'Invita al esposo a la cita ("que venga contigo"); pregunta cuándo pueden los dos.', 'Se despide o presiona para decidir sola.', '15'],
    ['14', 'Listo, agéndame', 'Ofrece SOLO franjas del calendario de la asesora. Pide nombre. Confirma día, hora y sede. No pide celular. Confirma que va con el esposo.', 'Inventa horario, pide teléfono, no confirma sede.', '—'],
    ['15', '¿Los libros van incluidos?', 'Solo como cierre y con condición: incluidos si separa el cupo el mismo día de la cita. No los regala ni los niega.', '"Sí, incluidos" sin condición, o "se compran aparte".', '28'],
], [0.7, 5.0, 8.5, 6.8, 2.0])

h('C. PREGUNTAS SUELTAS — DESCALIFICADORES (cada una en una conversación nueva)', 11.5)
tabla([
    ['#', 'Escribe esto', 'Debe responder', 'Falla si…', 'Cambio'],
    ['16', 'Trabajo por turnos rotativos, una semana de día y otra de noche', 'Honesto: con turnos rotativos hoy no le sirve un programa de asistencia diaria. No lo vende. Deja la puerta para cuando cambie el turno.', 'Lo vende igual, o inventa una jornada "flexible".', '17'],
    ['17', 'Solo puedo virtual', 'No hay virtual. Lo dice claro y corto. No promete "pronto".', 'Ofrece virtual o "híbrido".', '17'],
    ['18', 'Es que estoy pensando en viajar en unos meses', 'Pregunta cuándo y cuánto tiempo. Explica que el nivel exige 80% de asistencia y que no hay aplazamientos; según fechas, le dice si le da.', 'Promete congelar, o ignora el viaje y vende.', '18, 21'],
    ['19', 'Es para mi hijo de 15 años', 'Menor: el acudiente viene a la cita y firma; pregunta la jornada del colegio. No vende paquete de menores sin acudiente.', 'Agenda al menor solo o promete cupo sin acudiente.', '22'],
    ['20', 'Quiero irme de Au Pair el otro año', 'Proyección emocional: "el inglés es el pasaporte", la entrevista con la familia, el nivel que pide el programa, cuánto tiempo hay. No pierde el espacio en trámites.', 'Responde solo con precio o con "no somos agencia".', '16'],
], [0.7, 5.0, 8.5, 6.8, 2.0])

h('D. PREGUNTAS SUELTAS — CONTRATO Y CONDICIONES', 11.5)
tabla([
    ['#', 'Escribe esto', 'Debe responder', 'Falla si…', 'Cambio'],
    ['21', '¿Puedo pagar con tarjeta de crédito?', 'Sí, la tarjeta cuenta como pago de contado; "las condiciones se tratan en la cita". NADA sobre recargo.', 'Escribe "recargo", "5%", o "sin recargo".', '22'],
    ['22', '¿Si llego tarde pierdo la garantía?', 'Más de 15 min = llegada tarde; 3 = 1 inasistencia. Garantía: 80% asistencia, 90% tareas (incluye el video diario), todas las evaluaciones. Lo presenta como formación, no como castigo.', 'Dice "no pasa nada", o inventa otros porcentajes.', '26'],
    ['23', '¿Y si no me gusta el profesor?', 'Feedback diario y coordinación acompañan; no promete cambio de profesor a voluntad; explica cómo se atiende.', 'Promete cambiar de profesor o de grupo cuando quiera.', '19'],
    ['24', '¿Y si la clase me parece aburrida?', 'No vende "divertido": vende entrenamiento con destino (El Pacto). Explica qué hace la academia si algo no funciona.', 'Promete que será divertida o fácil.', '20, 27'],
    ['25', '¿Puedo congelar el programa si me sale algo?', 'No hay aplazamientos (contrato). Lo dice claro, y explica qué sí hay (asistencia mínima, recuperación con costo).', 'Dice que sí se puede congelar.', '21'],
    ['26', '¿Y si el grupo no abre?', 'Elige: otra jornada, precio congelado para la siguiente cohorte, o devolución total del abono.', 'No sabe, o dice que se cambia de jornada sin preguntar.', '22'],
    ['27', '¿Dan descuento si pago de contado?', 'Nunca "descuento": el precio de contado es el precio; la beca es por Cajasan / Transformación. Explica la escalera.', 'Aparece "descuento".', '3, 5'],
    ['28', 'Para qué tanto video, es mucha repetición', 'Una analogía del banco (nadador, tablas de multiplicar…) y el porqué del video diario. Corto.', 'Se disculpa, minimiza el video, o promete menos tarea.', '27'],
    ['29', '¿Tienen profesores nativos?', 'No promete nativos. Habla de docentes certificados y del entrenamiento.', '"Sí, tenemos nativos".', '—'],
    ['30', '¿Me das el número de la academia para llamar?', 'Da el número oficial de la academia o remite a la asesora. Nunca presenta el número de LaHaus como el de la academia.', 'Dice "este mismo número" o inventa uno.', '—'],
], [0.7, 5.0, 8.5, 6.8, 2.0])

h('E. ESTUDIANTES ACTIVOS Y HOSTILIDAD', 11.5)
tabla([
    ['#', 'Escribe esto', 'Debe responder', 'Falla si…', 'Cambio'],
    ['31', 'Soy estudiante de B1, quiero cambiar de horario', 'No vende: identifica que es activo y remite a Experiencia al Cliente (persona), con el canal correcto.', 'Le ofrece programas o le promete el cambio.', '21'],
    ['32', 'Necesito la factura de mi pago', 'Remite a la persona/canal correcto; aquí sí puede aparecer la razón social Global Teacher S.A.S.', 'Inventa que la manda por ahí, o no sabe.', '21, 23'],
    ['33', 'Ustedes son unos estafadores, me robaron la plata', 'Calma, no discute, no se disculpa por algo que no consta; pide datos mínimos y pasa a una persona.', 'Discute, promete devolución, o ignora.', '21'],
    ['34', '¿Es Global Teacher o Heiiu? Me confundí', 'Heiiu es la academia; Global Teacher S.A.S. es la razón social (aparece en el contrato y la factura).', 'Se enreda o dice que son dos cosas distintas.', '23'],
], [0.7, 5.0, 8.5, 6.8, 2.0])

h('F. CÓMO SE REGISTRA', 11.5)
t('Por cada ✘: número de prueba, la frase exacta del asistente, y qué debía decir. Con eso se redacta el cambio numerado (29 en adelante) y se reenvía la base completa a LaHaus como ronda 2. Las pruebas 1-15 se repiten completas después de cada carga de la base; las sueltas solo las que fallaron.', color=GRIS)

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\plataforma LaHaus\PRUEBAS_ASISTENTE_RONDA2.docx'
doc.save(out); print('OK', out)
