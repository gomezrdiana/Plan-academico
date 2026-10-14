# -*- coding: utf-8 -*-
"""ANEXO AL KIT DE VENTA (15/10/2026): El Pacto versión 15/10 + la perfilación ANTES de informar, como estructura de la llamada.
Para la asesora comercial (por rol, sin nombres). Se suma al Kit de Venta del 25/09 sin reeditarlo (regla: dudas vivas, kits congelados).
Fuentes: decisiones de gerencia 13/10 (flujo del asistente de 7 pasos, El Pacto v3, cita por llamada), investigaciones de
precalificación, cadencia y formato de cita (plataforma LaHaus/), Kit de Venta v2 (sección 1 y duda 5).
Salida: recursos/comercial/INDUCCION_COMERCIAL/ANEXO_KIT_PACTO_Y_PERFILACION.docx"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NARANJA = RGBColor(0xE5, 0x4C, 0x09); GRIS = RGBColor(0x66, 0x66, 0x66)
doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.5); s.bottom_margin = Cm(1.3); s.left_margin = Cm(1.9); s.right_margin = Cm(1.9)
doc.styles['Normal'].font.name = 'Calibri'; doc.styles['Normal'].font.size = Pt(10.5)

def shd(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear'); el.set(qn('w:color'), 'auto'); el.set(qn('w:fill'), color); tcPr.append(el)
def sec(txt):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(3)
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = NARANJA
def t(txt, bold=False, size=10.5, color=None, after=4):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(txt); r.font.size = Pt(size); r.bold = bold
    if color: r.font.color.rgb = color
    return p
def item(txt, prefix=None):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2); p.paragraph_format.left_indent = Cm(0.4)
    if prefix:
        r = p.add_run(prefix + ' '); r.bold = True; r.font.size = Pt(10.5)
    r = p.add_run(txt); r.font.size = Pt(10.5)
def frase(txt):
    tb = doc.add_table(rows=1, cols=1); tb.style = 'Table Grid'; c = tb.cell(0, 0); shd(c, 'FFF8E7')
    c.paragraphs[0].paragraph_format.space_after = Pt(0)
    r = c.paragraphs[0].add_run(txt); r.font.size = Pt(10.5); r.italic = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
def tabla(filas, anchos, fs=9.5):
    tb = doc.add_table(rows=len(filas), cols=len(filas[0])); tb.style = 'Table Grid'
    for i, fila in enumerate(filas):
        for j, v in enumerate(fila):
            c = tb.cell(i, j); c.width = Cm(anchos[j]); c.paragraphs[0].paragraph_format.space_after = Pt(0)
            r = c.paragraphs[0].add_run(v); r.font.size = Pt(fs)
            if i == 0: r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shd(c, 'E54C09')
            elif j == 0: r.bold = True; shd(c, 'FFF8E7')
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
r = p.add_run('ANEXO AL KIT DE VENTA — EL PACTO Y LA PERFILACIÓN ANTES DE INFORMAR'); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = NARANJA
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(8)
r = p.add_run('Para la asesora comercial · Heiiu English Academy · 15 de octubre de 2026 · Uso interno · Se suma al Kit de Venta del 25/09; donde choque con él, manda este anexo'); r.font.size = Pt(9); r.font.color.rgb = GRIS

sec('1. QUÉ CAMBIA DESDE HOY (en cinco líneas)')
item('Antes de dar CUALQUIER información (programa, precio, beca, jornadas) se perfila al cliente con seis preguntas, en este orden. La información llega al final, cuando ya sabes qué le sirve.', '1.')
item('El Pacto se le dice a TODOS, antes de informar, con el texto de la sección 3. La charla larga de "no fuiste tú, fue el método" sigue siendo solo para el cliente quemado (duda 5 del Kit) y va dentro del mismo paso.', '2.')
item('El asistente del chat hace estas mismas preguntas desde hoy. Lo que el cliente ya respondió lo ves en el panel: en la llamada lo CONFIRMAS, no lo repites. Si el cliente llegó por formulario o por otra vía, haces las seis completas.', '3.')
item('La asesoría es por LLAMADA de WhatsApp por defecto, desde la plataforma (queda grabada y transcrita). Videollamada solo si el cliente la pide o si decide con alguien más; presencial solo si él lo pide. La firma del contrato siempre es presencial.', '4.')
item('Mensajes y llamadas por iniciativa nuestra: lunes a viernes de 7:00 a.m. a 7:00 p.m. y sábados de 8:00 a.m. a 3:00 p.m.; nunca domingos ni festivos; máximo un contacto al día por cliente. Lo que el cliente escriba se responde a cualquier hora.', '5.')

sec('2. LA ESTRUCTURA DE LA LLAMADA (20 a 30 minutos)')
t('El principio: preguntar más que hablar. El que informa primero vende el programa; el que perfila primero vende el plan que le sirve a ESA persona y cierra.', bold=True)
tabla([
 ['Min.', 'Paso', 'Qué dices (texto)', 'Qué anotas'],
 ['0-1', 'Apertura', '"Hola [nombre], soy [nombre], de Heiiu. Agendaste tu asesoría para esta hora: ¿tienes 20 minutos? Te hago unas preguntas para armarte el plan que te sirve y después te cuento todo."', 'Que está en condiciones de hablar. Si no: nueva hora, escrita por él.'],
 ['1-3', '1. Motivo y destino', '"¿Para qué necesitas el inglés? ¿Para cuándo?" Si es genérico: "¿En qué trabajas o qué proyecto tienes en la mira?"', 'Motivo: trabajo · irse · hijo · ascender · emprender. Fecha límite si la hay.'],
 ['3-4', '2. Para quién y quién decide', '"¿El curso es para ti? ¿Alguien más decide contigo?"', 'Decisor. Si es para otro adulto, esa persona debe estar en la llamada o en otra.'],
 ['4-7', '3. Nivel y experiencia', '"¿Has estudiado inglés antes? ¿Cómo te fue?" Si ya estudió: "¿Y qué pasó esa vez? ¿Por qué lo dejaste?"', 'Cero / ya sabe algo (→ sesión de ubicación). Marca QUEMADO si fracasó antes.'],
 ['7-8', '4. Presencial y jornada', '"Las clases son presenciales en Bucaramanga: mañana de 8:00 a 12:00 o noche de 6:30 a 8:30. ¿Cuál te queda?" y "¿Tu horario es fijo o tienes flexibilidad?"', 'Jornada · FLEXIBLE / SOLO MAÑANA / SOLO NOCHE. Turnos rotativos o no presencial: ruta personalizado o despedida honesta.'],
 ['8-11', '5. EL PACTO', 'Texto de la sección 3, completo, y la pregunta: "¿Te le mides a ese ritmo?" Al quemado, antes de la pregunta, la absolución (duda 5 del Kit).', 'PACTO SÍ / DUDA / NO. Con NO o "busco algo más suave": despedida honesta y fin. No se informa precio.'],
 ['11-13', '6. Forma de pago', '"Para llevarte el plan que te sirve: ¿cómo prefieres invertir: de contado (sirven cesantías, tarjeta o crédito de tu banco) o con una inicial del 20% y cuotas mensuales?" y "¿Eres afiliado a Cajasan?"', 'CONTADO / CUOTAS / ALGO MÁS BAJO · Cajasan sí/no. "Algo más bajo" NO es despedida: es ruta de menor ticket (Arranque, nivel suelto, sabatino).'],
 ['13-25', '7. SOLO AHORA: la información', 'Kit, sección 1, tal cual: programa completo con SU beca primero, Arranque segundo (cierre doble) · las 4 anclas · módulos de graduación · frase del expediente (garantía y video diario se firman juntos) · libros como gancho de cierre.', 'Su número exacto (tarjeta de becas o motor de cuotas). Nunca una cuota inventada.'],
 ['25-30', '8. Cierre', 'Una frase de cierre y silencio. Salidas: matriculado o cupo separado · lista de espera CON FECHA · no claro con motivo.', 'Separación digital (duda 8 del Kit): comprobante + cédula + "ACEPTO" + sesión de ubicación agendada. Formalización en la sede.'],
], [1.2, 2.6, 8.0, 5.4], fs=9)
t('Si el cliente pide el precio antes del paso 7: "Te lo digo completo en dos minutos, cuando sepa qué te sirve; déjame hacerte dos preguntas más." Si insiste, la referencia de la sección 4, una sola vez, y sigues con las preguntas.', bold=True)

sec('3. EL PACTO: TEXTO EXACTO (versión del 15/10/2026, para todos, antes de informar)')
frase('"Antes de contarte precios, te cuento cómo es esto, para que decidas con todo claro. Nuestro objetivo es que salgas hablando inglés; si no se logra, algo está mal: o tú dejaste de entrenar, o nosotros no estamos cumpliendo. Esto es un entrenamiento en equipo, como un gimnasio o aprender a nadar: desde el primer día hablas, tu profesor identifica dónde puedes mejorar y lo repetimos hasta que sale. Por eso grabas un video corto todos los días, y tu profesor te lo devuelve con una mejora. A cambio, la garantía por escrito: si cumples tu parte y no avanzas, te devolvemos el 100% del nivel. ¿Te le mides a ese ritmo?"')
item('"Son unos minutos al día, y tu profesor te responde cada uno. Es la parte que más te hace avanzar. ¿Te le mides?"', 'Si duda ("¿todos los días?"):')
item('"Te agradezco la sinceridad. Este programa es exigente a propósito, y prefiero decírtelo ahora a que lo descubras en la segunda semana. Si más adelante quieres entrenarte con propósito, ¡escríbenos!" Fin de la llamada; se marca FRÍO con el motivo. Sin reintentos.', 'Si dice "busco algo más suave":')
item('Antes de la pregunta "¿te le mides?", la absolución de la duda 5 del Kit: "si llevas años intentándolo y no has podido, el problema nunca fuiste tú: fue el método, años sentado copiando del tablero, con miedo a hablar. Aquí es al revés: de pie, hablando desde el día uno". Solo a él. Al que llega limpio no se le receta un trauma que no tiene.', 'Al cliente QUEMADO:')
t('Por qué se dice antes de informar y no al final: el que acepta el ritmo antes de oír el precio compra el entrenamiento, no una promesa; el que firma advertido no deserta en la semana 3 ni pide devolución; y la garantía solo es creíble si el trato se dijo completo. Lo que NUNCA se dice: fácil, rápido, divertido, "en tres meses hablas", nativos, empleo, visa.', size=10, color=GRIS)

sec('4. EL PRECIO: CUÁNDO Y CÓMO')
item('Nunca antes del paso 7 por iniciativa tuya. Ningún total "de lista" suelto, ninguna cuota de memoria.', 'Regla:')
item('"Te doy la referencia: el programa completo A1 a B2 es de $8.696.000 de lista, y con la beca del Fondo, hasta el 45% según tu perfil, baja bastante. Se paga de contado o con una inicial del 20% y cuotas mensuales. Pero más que el dinero, que lo cuadramos, lo que buscamos es que estés dispuesto a entrenarte con nosotros. Tu número exacto te lo doy en dos minutos."', 'Si lo pide antes (una sola vez):')
item('Con la tarjeta de becas (contado) o con el motor de cuotas (inicial del 20% y hasta 11 cuotas). El Arranque A1+A2 es $2.990.000, precio único. La cuota se dice solo con el motor abierto.', 'En el paso 7:')
item('"Hay opciones con inicial y cuotas, y con cesantías o crédito de tu banco. ¿Quieres que te muestre cómo puedes entrenarte con nosotros con menos esfuerzo económico?" Sí → escalera (Kit, sección 1). No → "Perfecto. Cuando sea tu momento, me escribes y te agendo." Fin.', '"No tengo el dinero ahora" (una sola vez):')

sec('5. LO QUE EL ASISTENTE DEL CHAT YA HIZO (lee el panel ANTES de llamar)')
t('Desde hoy el asistente hace los pasos 1 a 6 en el chat y marca estos campos. Tú arrancas la llamada confirmando, no repitiendo: "Vi que me contaste que necesitas el inglés para [motivo] y que te le mides al ritmo del video diario: ¿sigue siendo así?"')
tabla([
 ['Campo del panel', 'Qué significa para ti'],
 ['Motivo / destino', 'La palanca de la llamada: comparador → cuenta por hora; viajero → Pasaporte; papá → carácter; quemado → Pacto completo.'],
 ['Para quién / decisor', 'Si decide con alguien más, lo invitas a la llamada o propones videollamada.'],
 ['Nivel / experiencia', 'Cero → Arranque directo. Ya sabe algo → sesión de ubicación agendada antes de despedirte.'],
 ['Presencial + jornada', 'Confirmas la jornada y el "fijo o flexible".'],
 ['Pacto: sí / duda / no', 'SÍ → no lo repites completo: una frase de confirmación. DUDA → lo dices completo. NO → no debería haber cita; si la hay, lo dices completo y decides.'],
 ['Forma de pago', 'Contado / cuotas / algo más bajo. Define por dónde entras en la escalera. Nunca preguntar ingresos ni consultar centrales.'],
 ['Audio (opcional)', 'Si mandó el audio de 30 segundos, escúchalo antes de llamar.'],
 ['Puntaje y etiqueta', 'Caliente (8-10): le escribes en menos de 15 minutos. Tibio (5-7): el mismo día. Frío (0-4): no se agenda; tú decides si lo revisas. El tibio es el que más gana con tu llamada.'],
], [4.5, 12.7])
item('Lead de FORMULARIO (no pasó por el chat): lo llamas en menos de 5 minutos en horario legal; si no contesta, un mensaje. En la llamada haces las seis preguntas completas.', 'Formulario:')
item('Lead que chateó y no agendó: primero un mensaje por la misma línea, "¿te llamo ahora 5 minutos? te marco desde el 3xx…", y llamas cuando dice sí. Lead que YA agendó: no se le llama antes de la cita; se le confirma por escrito.', 'Llamada anunciada:')

sec('6. CHECKLIST DE TODA LLAMADA (para auditar, no para decorar)')
for c in ['Perfilé antes de informar (pasos 1 a 6 antes del 7).', 'Motivo y destino anotados.', 'Nivel y experiencia; marqué QUEMADO si aplica.', 'Presencial, jornada y "fijo o flexible".',
          'Dije El Pacto completo (o lo confirmé si el panel ya decía SÍ) y anoté la respuesta.', 'Forma de pago y Cajasan.', 'Programa completo con SU beca primero, Arranque segundo.',
          'Las 4 anclas y la frase del expediente (garantía + video diario se firman juntos).', 'Una frase de cierre y silencio. Salida 1, 2 o 3 con fecha o motivo.', 'Llamé desde la plataforma, en horario legal.']:
    item(c, '☐')
t('Todo lo demás del Kit sigue igual: escalera de cierre, anclas, tarjeta de becas, libros como gancho, separación digital, contador, Cajasan, lo que jamás se hace.', bold=True, size=10, color=GRIS)

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\INDUCCION_COMERCIAL\ANEXO_KIT_PACTO_Y_PERFILACION.docx'
try: doc.save(out); print('OK', out)
except PermissionError: doc.save(out.replace('.docx', '_v2.docx')); print('ABIERTO -> _v2')
