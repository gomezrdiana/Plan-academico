# -*- coding: utf-8 -*-
"""Recibo de separacion de cupo — 1 pagina, lo firma el cliente al dar el abono."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NARANJA = RGBColor(0xE5, 0x4C, 0x09)
GRIS = RGBColor(0x77, 0x77, 0x77)
CREMA = 'FFF8E7'
NARANJA_HEX = 'E54C09'

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.3); s.bottom_margin = Cm(1.1); s.left_margin = Cm(1.9); s.right_margin = Cm(1.9)
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(10)

def shd(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement('w:shd'); el.set(qn('w:val'), 'clear'); el.set(qn('w:fill'), color)
    tcPr.append(el)

def t(txt, bold=False, size=10, after=2):
    p = doc.add_paragraph(); r = p.add_run(txt); r.font.size = Pt(size); r.bold = bold
    p.paragraph_format.space_after = Pt(after)
    return p

def num(n, txt, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(f'{n}. '); r.bold = True; r.font.size = Pt(10); r.font.color.rgb = NARANJA
    r2 = p.add_run(txt); r2.font.size = Pt(10); r2.bold = bold
    p.paragraph_format.left_indent = Cm(0.4); p.paragraph_format.space_after = Pt(3)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('RECIBO Y ACUERDO DE SEPARACIÓN DE CUPO'); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = NARANJA
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Heiiu English Academy · Global Teacher S.A.S. — NIT 900.422.478-2 · Carrera 27 # 48-49, Bucaramanga'); r.font.size = Pt(9); r.font.color.rgb = GRIS
doc.add_paragraph()

# Tabla de datos
tb = doc.add_table(rows=6, cols=4); tb.style = 'Table Grid'
datos = [
    ('Recibo No.', '', 'Fecha', ''),
    ('Nombre del cliente', '', 'C.C.', ''),
    ('Teléfono', '', 'Correo', ''),
    ('Programa', '', 'Cohorte / inicio', ''),
    ('Valor del abono', '$', 'Medio de pago', ''),
    ('Precio total congelado', '$', 'Plan de pago acordado', ''),
]
for i, fila in enumerate(datos):
    for j, val in enumerate(fila):
        c = tb.rows[i].cells[j]
        pp = c.paragraphs[0]; rr = pp.add_run(val); rr.font.size = Pt(10)
        if j in (0, 2):
            rr.bold = True; shd(c, CREMA)

doc.add_paragraph()
t('CONDICIONES DEL ABONO — leídas y aceptadas por el cliente al firmar:', bold=True, after=4)

num(1, 'Este abono SEPARA el cupo del cliente en el programa y la cohorte indicados, y CONGELA el precio registrado arriba.')
num(2, 'El saldo del programa se paga según el plan de pago acordado y registrado en este recibo, a más tardar antes del inicio de clases, salvo acuerdo escrito distinto.')
num(3, 'DERECHO DE RETRACTO: dentro de los cinco (5) días hábiles siguientes al pago, el cliente puede retractarse y se le devuelve la totalidad del abono (Ley 1480 de 2011).')
num(4, 'VENCIDO EL RETRACTO: el abono NO se devuelve en dinero. Queda como SALDO A FAVOR del cliente, aplicable a cualquier programa de la academia, con vigencia de doce (12) meses desde esta fecha.', bold=True)
num(5, 'PAGO CON CRÉDITO DE TERCEROS O CESANTÍAS: cuenta como pago de contado. El desembolso o cheque debe recibirse dentro de los diez (10) días hábiles siguientes y siempre antes del inicio de clases; de lo contrario, el cliente puede pagar por otro medio a la tarifa que corresponda, o su abono queda como saldo a favor (numeral 4).')
num(6, 'La aceptación de este recibo por mensaje de datos (WhatsApp o correo, con copia de la cédula y comprobante del pago) tiene plena validez como separación del cupo (Ley 527 de 1999). La matrícula se FORMALIZA en la sede con la firma del contrato de matrícula y la sesión de ubicación realizada; este recibo no constituye matrícula.')
num(7, 'El cupo y el precio congelado quedan sujetos a completar el pago en los términos acordados; el inicio de la cohorte no se aplaza por pagos pendientes.')

doc.add_paragraph()
t('Firmas:', bold=True, after=8)
t('CLIENTE: _____________________________  C.C. _________________     RECIBE POR LA ACADEMIA: _____________________________', after=2)

doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Original para la academia · copia (foto) para el cliente · Septiembre 2026'); r.font.size = Pt(8.5); r.font.color.rgb = GRIS

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\documentos_operativos\plantillas\RECIBO_SEPARACION_CUPO.docx'
doc.save(out); print('OK', out)
