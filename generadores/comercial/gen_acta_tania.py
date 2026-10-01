# -*- coding: utf-8 -*-
"""Caso Tania (01/10/2026): acta de compromiso por excepcion (2 semanas + examen oral externo) y mensaje.
Salida: recursos/experiencia_cliente/Tania/"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import io

OUT = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\experiencia_cliente\Tania'
GRIS = RGBColor(0x77, 0x77, 0x77)

d = Document()
for s in d.sections:
    s.top_margin = Cm(1.6); s.bottom_margin = Cm(1.4); s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)
d.styles['Normal'].font.name = 'Calibri'; d.styles['Normal'].font.size = Pt(10.5)

def t(txt, bold=False, after=5, size=10.5, color=None, align=None):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    if align: p.alignment = align
    r = p.add_run(txt); r.bold = bold; r.font.size = Pt(size)
    if color: r.font.color.rgb = color

def num(n, tit, txt):
    p = d.add_paragraph(); p.paragraph_format.left_indent = Cm(0.4); p.paragraph_format.space_after = Pt(5)
    r = p.add_run(f'{n}. {tit} '); r.bold = True
    p.add_run(txt)

t('ACTA DE COMPROMISO ACADÉMICO — EXCEPCIÓN POR ÚNICA VEZ', bold=True, size=13.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
t('Heiiu English Academy — Global Teacher S.A.S. · Coordinación académica · Bucaramanga, 1 de octubre de 2026', size=9, color=GRIS, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)

t('Estudiante: TANIA ______________________________   Documento: ____________________   Nivel cursado: A2 nocturno', after=8)

t('Antecedentes', bold=True, after=3)
t('La estudiante presentó el examen final de A2 y obtuvo 38 sobre 100. Tomó el paquete de recuperación previsto en la política de la academia (tres sesiones de refuerzo y un reintento con evaluador externo) y en el reintento obtuvo 56 sobre 100. El mínimo para aprobar el nivel es 70. Conforme a la política, el reintento es uno solo.', after=8)

t('Decisión de la academia', bold=True, after=3)
t('En atención al avance demostrado en el refuerzo (de 38 a 56) y a la solicitud de la estudiante de continuar, la coordinación académica concede, POR ÚNICA VEZ Y COMO EXCEPCIÓN QUE NO CONSTITUYE PRECEDENTE, un periodo de compromiso de dos semanas y una evaluación oral final adicional, bajo las condiciones siguientes:', after=6)

num(1, 'Ingreso provisional al B1 nocturno.', 'La estudiante asiste al grupo B1 nocturno desde el 1 de octubre de 2026 en calidad de provisional. Su permanencia en el nivel depende del resultado de la evaluación del punto 5.')
num(2, 'Asistencia del 100%.', 'Entre el 1 y el 15 de octubre de 2026 asiste a todas las clases, puntual (la clase inicia a las 6:30 PM; después de las 6:45 cuenta como llegada tarde).')
num(3, 'Video diario.', 'Todos los días de clase graba y envía al grupo del cohorte un video en inglés de mínimo CUATRO minutos, hablado de corrido y sin leer. Un video leído, cortado o de menos de cuatro minutos no cuenta. No se admite audio.')
num(4, 'Tareas y trabajo en clase.', 'Entrega todas las tareas asignadas y el ticket de salida de cada clase, y participa en las simulaciones.')
num(5, 'Evaluación oral final.', 'El jueves 15 de octubre de 2026 presenta una evaluación ORAL del nivel A2, aplicada por el evaluador externo de la academia (no por su profesor). Se aprueba con 70 o más.')
num(6, 'Resultado.', 'Si aprueba, queda matriculada en el B1 con todos los efectos. Si no aprueba, o si incumple cualquiera de los puntos 2 a 4, no continúa en el B1; coordinación le indicará sus opciones según la política vigente. No habrá una tercera oportunidad de evaluación.')
num(7, 'Lo que no cambia.', 'El plan de pagos de la estudiante y las demás condiciones de su contrato siguen iguales. Los resultados de las evaluaciones los comunica únicamente coordinación.')

d.add_paragraph()
t('Firmado en Bucaramanga el ____ / ____ / 2026', after=12)
t('LA ESTUDIANTE                                                                  COORDINACIÓN ACADÉMICA', after=16)
t('_____________________________                                _____________________________', after=2)
t('Tania ____________________________                              Heiiu English Academy', after=2)
t('C.C. ______________________________', after=2)
d.save(OUT + r'\ACTA_COMPROMISO_TANIA_01OCT2026.docx'); print('OK acta')

msg = """# MENSAJE A TANIA — 01/10/2026 (lo envía Experiencia al Cliente / gerencia)

Hola Tania. Te escribo con el resultado de tu reintento y con una decisión.

El resultado: obtuviste 56 sobre 100. El mínimo para aprobar el A2 es 70, así que el reintento no quedó aprobado. Te lo digo de frente porque así te hablamos siempre.

Lo que sí vimos: pasaste de 38 a 56 en tres sesiones de refuerzo. Eso es trabajo real, y por eso no queremos que tu proceso se corte aquí.

La decisión: por única vez, y como excepción, te damos dos semanas de compromiso y una evaluación oral adicional. Funciona así:

1. Entras hoy mismo al grupo de B1 nocturno, de forma provisional.
2. Del 1 al 15 de octubre: asistencia del 100%, puntual (6:30 PM).
3. Todos los días de clase, un video en inglés de mínimo 4 minutos, hablado de corrido y sin leer, enviado al grupo. Un video leído, cortado o más corto no cuenta.
4. Todas las tareas y el ticket de salida de cada clase.
5. El jueves 15 de octubre presentas una evaluación oral del A2 con el evaluador externo. Se aprueba con 70.

Si apruebas, quedas en B1 y seguimos. Si no apruebas, o si no cumples alguno de los puntos, no continúas en B1 y te explicamos tus opciones. No habrá una tercera oportunidad, y te lo decimos ahora para que no haya sorpresas.

Antes de entrar a clase hoy firmas un acta con estas condiciones; Jacqueline la tiene en la sede.

Tania, en tres días de refuerzo subiste 18 puntos. Eso no lo hace cualquiera, y es la señal de que con ese mismo esfuerzo, sostenido dos semanas, no solo apruebas: puedes ser la mejor del grupo. Depende solo de ti. Nos vemos esta noche.

Depto. Experiencia al Cliente — Heiiu English Academy

---
(Para después, SOLO si aprueba el 15/10, en un mensaje aparte: "Tania, felicitaciones. Si quieres, nos encantaría que grabaras un video corto contando cómo fue pasar de 38 a aprobar; a otros estudiantes que están donde tú estabas les sirve mucho. Solo si quieres.")
"""
io.open(OUT + r'\MENSAJE_TANIA_01OCT2026.md', 'w', encoding='utf-8').write(msg); print('OK mensaje')
