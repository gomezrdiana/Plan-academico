# -*- coding: utf-8 -*-
"""Anexo de habilitacion comercial para docentes: solo lo firma quien recibe induccion con el kit.
Activa las comisiones de asesoria y cierre (5/2.5 leads de la casa, 8 canal propio) sobre el contrato docente v3.1."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

GRIS = RGBColor(0x77, 0x77, 0x77)
doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.6); s.bottom_margin = Cm(1.4); s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)
doc.styles['Normal'].font.name = 'Calibri'; doc.styles['Normal'].font.size = Pt(10.5)

def t(txt, bold=False):
    p = doc.add_paragraph(); r = p.add_run(txt); r.font.size = Pt(10.5); r.bold = bold; p.paragraph_format.space_after = Pt(4)

def num(n, txt):
    p = doc.add_paragraph(); r = p.add_run(f'{n}. '); r.bold = True
    p.add_run(txt); p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.space_after = Pt(3)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('ANEXO DE HABILITACIÓN COMERCIAL — DOCENTE'); r.bold = True; r.font.size = Pt(13)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Anexo al Contrato de Prestación de Servicios Profesionales — Docente Hora Cátedra v3.1 · Hace parte integral del contrato'); r.font.size = Pt(9.5); r.font.color.rgb = GRIS
doc.add_paragraph()

t('Entre GLOBAL TEACHER S.A.S., NIT 900.422.478-2, representada por DIANA MARCELA GÓMEZ RANGEL (EL CONTRATANTE), y [NOMBRE COMPLETO DEL DOCENTE], C.C. [___________] (EL CONTRATISTA), se suscribe el presente anexo, en desarrollo de la Cláusula Segunda literal (c) del contrato docente vigente entre las partes.')

t('1. HABILITACIÓN.', bold=True)
t('EL CONTRATANTE certifica que EL CONTRATISTA recibió la inducción comercial de la institución y el Kit de Venta vigente, y lo habilita para conducir asesorías y cerrar ventas de los programas de la academia, usando exclusivamente los precios, becas, planes de pago y materiales aprobados por gerencia. Esta habilitación es facultativa para EL CONTRATISTA: no lo obliga a vender ni su omisión genera incumplimiento.')

t('2. REMUNERACIÓN DE LA ACTIVIDAD COMERCIAL (siempre sobre el recaudo efectivo a caja).', bold=True)
num(1, 'REFERIDO (contacto registrado por EL CONTRATISTA, cerrado por el equipo comercial): DOS por ciento (2%), conforme al contrato.')
num(2, 'LEAD DE LA CASA cerrado por EL CONTRATISTA (contacto asignado por la institución en la planilla, asesoría y cierre a cargo del CONTRATISTA): CINCO por ciento (5%) del recaudo de cliente nuevo y 2,5% de cliente antiguo.')
num(3, 'CANAL PROPIO (contacto originado en los canales personales del CONTRATISTA, asesoría y cierre a cargo del CONTRATISTA): OCHO por ciento (8%) del recaudo de cliente nuevo.')
t('Cierra quien conduce la asesoría y llega hasta el abono o la matrícula, con el registro correspondiente en la planilla comercial. Si el contacto lo registró EL CONTRATISTA y la asesoría la hizo otra persona, aplica el numeral 1. Para reclamar cualquier comisión, el contacto debe estar registrado en la planilla desde el primer contacto; en caso de concurrencia, la comisión corresponde a quien registró primero. Sin recaudo no hay comisión; si un pago comisionado es objeto de retracto o devolución, la comisión se descuenta de la siguiente liquidación.')

t('3. CORTES Y PAGO.', bold=True)
t('Las comisiones se liquidan mensualmente con corte el último día de cada mes y se pagan contra cuenta de cobro dentro de los cinco (5) días hábiles siguientes, respecto de pagos cuyo término legal de retracto ya esté vencido.')

t('4. REGLAS DE VENTA.', bold=True)
t('EL CONTRATISTA se obliga a: ofrecer únicamente la matriz vigente aprobada por gerencia, sin descuentos ni condiciones inventadas; referirse siempre a "beca del Fondo" y jamás a "descuento" ni a "patrocinio"; no prometer empleo, visas, facilidad, rapidez ni resultados distintos del texto exacto de la garantía institucional; someter a aprobación escrita de gerencia toda pieza publicitaria que use la marca antes de publicarla; registrar cada gestión en la planilla comercial; y no recibir dinero de clientes en efectivo ni en cuentas propias (todo pago ingresa a las cuentas oficiales de la institución). El aula jamás se usa para actividad comercial. EL CONTRATISTA no tiene facultad de representación: no firma convenios ni compromete jurídicamente a GLOBAL TEACHER S.A.S.')

t('5. VIGENCIA.', bold=True)
t('Este anexo rige mientras esté vigente el contrato docente, y cualquiera de las partes puede dejarlo sin efecto por escrito en cualquier momento, sin que ello afecte el contrato docente ni las comisiones ya causadas.')

doc.add_paragraph()
t('Para constancia se firma en Bucaramanga, el día ______ de ____________ de 2026.')
doc.add_paragraph()
t('EL CONTRATANTE                                                        EL CONTRATISTA')
doc.add_paragraph()
t('_____________________________                                _____________________________')
t('DIANA MARCELA GÓMEZ RANGEL                              Nombre: ________________________')
t('Representante Legal — GLOBAL TEACHER S.A.S.                 C.C. ______________________________')

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\contratos\ANEXO_HABILITACION_COMERCIAL_DOCENTE_v1.docx'
doc.save(out); print('OK', out)
