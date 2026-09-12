# -*- coding: utf-8 -*-
"""Motor de cuotas 2026 con matriz hibrida. REGLA: se termina de pagar ANTES de terminar el programa
-> el maximo de cuotas depende de la modalidad (duracion), igual que el Excel original de Diana."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

FULL = {'A1-B2 completo': 8696000, 'A2-B2': 7100000, 'B1-B2': 5185000, 'Solo B2': 2343000}
PCT = {
 'A1-B2 completo': {'TRANSFORMACION': (45, 36), 'GENERAL': (32, 25.6), 'EJECUTIVO': (15, 12)},
 'A2-B2':          {'TRANSFORMACION': (36, 27), 'GENERAL': (25.6, 19.2), 'EJECUTIVO': (12, 9)},
 'B1-B2':          {'TRANSFORMACION': (24.8, 15.8), 'GENERAL': (17.6, 11.2), 'EJECUTIVO': (8.3, 5.3)},
 'Solo B2':        {'TRANSFORMACION': (20.3, 11.3), 'GENERAL': (14.4, 8), 'EJECUTIVO': (6.8, 3.8)},
}
# (nombre modalidad, duracion aprox, columnas de cuotas) por paquete — tomado del Excel original
MODALIDADES = {
 'A1-B2 completo': [('SUPER INTENSIVO 20h/sem - ~7 meses - MAX 5 CUOTAS', [3, 4, 5]),
                    ('INTENSIVO 10h/sem - ~14 meses - MAX 11 CUOTAS', [4, 6, 8, 10, 11])],
 'A2-B2':          [('SUPER INTENSIVO 20h/sem - ~6 meses - MAX 5 CUOTAS', [3, 4, 5]),
                    ('INTENSIVO 10h/sem - ~12 meses - MAX 11 CUOTAS', [4, 6, 8, 10, 11])],
 'B1-B2':          [('SUPER INTENSIVO 20h/sem - ~4,5 meses - MAX 3 CUOTAS', [2, 3]),
                    ('INTENSIVO 10h/sem - ~9 meses - MAX 8 CUOTAS', [4, 6, 8])],
 'Solo B2':        [('INTENSIVO 10h/sem - ~5 meses - MAX 4 CUOTAS (en super intensivo dura ~2,5 meses: practicamente contado)', [2, 3, 4])],
}
INICIALES = [0.20, 0.30, 0.40, 0.50]
r = 0.019

def cuota(saldo, n):
    return round(saldo * r / (1 - (1 + r) ** -n))

wb = openpyxl.Workbook()
wb.remove(wb.active)

NARANJA = PatternFill('solid', fgColor='E54C09')
GRISF = PatternFill('solid', fgColor='EEEEEE')
CREMA = PatternFill('solid', fgColor='FFF8E7')
BLANCO_B = Font(bold=True, color='FFFFFF', size=11)
B = Font(bold=True, size=11)
NORMAL = Font(size=11)
borde = Border(*[Side(style='thin')]*4)
MONEDA = '#,##0'

for paq, full in FULL.items():
    ws = wb.create_sheet(paq)
    ws.column_dimensions['A'].width = 24
    for col in 'BCDEFG':
        ws.column_dimensions[col].width = 14
    fila = 1
    c = ws.cell(fila, 1, 'HEIIU - MOTOR DE CUOTAS 2026 - USO INTERNO (jamas mostrar al cliente)')
    c.font = Font(bold=True, size=12, color='E54C09'); fila += 1
    ws.cell(fila, 1, paq + ' - Precio FULL: $' + format(full, ',').replace(',', '.') +
            ' - Interes credito directo: 1,9% mensual sobre saldo').font = B; fila += 1
    ws.cell(fila, 1, 'REGLA: el estudiante termina de PAGAR antes de terminar el programa - por eso cada modalidad tiene su maximo de cuotas.').font = B; fila += 1
    ws.cell(fila, 1, 'Separacion de cupo $300.000 (se cruza con la inicial) -> cuota inicial min 20% max 3 dias -> cuotas desde dia 30. Credito SIEMPRE con pagare + carta + plan de pagos firmados.').font = NORMAL
    fila += 2
    for perfil, (pc_c, pc_q) in PCT[paq].items():
        base_c = round(full * (1 - pc_c / 100))
        base_q = round(full * (1 - pc_q / 100))
        c = ws.cell(fila, 1, perfil + '  -  Beca del Fondo: contado ' + str(pc_c) + '% / cuotas ' + str(pc_q) + '%')
        c.font = BLANCO_B; c.fill = NARANJA
        for j in range(2, 8):
            ws.cell(fila, j).fill = NARANJA
        fila += 1
        c = ws.cell(fila, 1, 'PRECIO CONTADO'); c.font = B; c.fill = CREMA
        c2 = ws.cell(fila, 2, base_c); c2.font = B; c2.fill = CREMA; c2.number_format = MONEDA
        ws.cell(fila, 3, 'PRECIO BASE A CUOTAS').font = B
        ws.cell(fila, 3).fill = CREMA
        c3 = ws.cell(fila, 5, base_q); c3.font = B; c3.fill = CREMA; c3.number_format = MONEDA
        fila += 1
        for nombre_mod, ncols in MODALIDADES[paq]:
            c = ws.cell(fila, 1, nombre_mod); c.font = B; c.fill = GRISF
            for j in range(2, 8):
                ws.cell(fila, j).fill = GRISF
            fila += 1
            ws.cell(fila, 1, 'Cuota inicial').font = B
            for j, n in enumerate(ncols):
                c = ws.cell(fila, 2 + j, str(n) + ' cuotas'); c.font = B; c.border = borde
            fila += 1
            for frac in INICIALES:
                ini = round(base_q * frac)
                saldo = base_q - ini
                c = ws.cell(fila, 1, str(int(frac * 100)) + '% = $' + format(ini, ',').replace(',', '.'))
                c.font = NORMAL; c.border = borde
                for j, n in enumerate(ncols):
                    c = ws.cell(fila, 2 + j, cuota(saldo, n))
                    c.number_format = MONEDA; c.font = NORMAL; c.border = borde
                fila += 1
        fila += 1

# ---- hoja ARRANQUE (precio unico, sin becas) ----
ws = wb.create_sheet('ARRANQUE A1+A2')
ws.column_dimensions['A'].width = 24
for col in 'BCDE':
    ws.column_dimensions[col].width = 14
fila = 1
c = ws.cell(fila, 1, 'HEIIU - MOTOR DE CUOTAS 2026 - USO INTERNO (jamas mostrar al cliente)')
c.font = Font(bold=True, size=12, color='E54C09'); fila += 1
ws.cell(fila, 1, 'ARRANQUE A1+A2 - PRECIO UNICO DE LANZAMIENTO: $2.990.000 (200 horas) - SIN BECAS - Interes credito directo: 1,9% mensual').font = B; fila += 1
ws.cell(fila, 1, 'REGLA: se termina de pagar antes de terminar. Super intensivo 20h/sem ~2,5 meses: MAX 2 CUOTAS. Intensivo 10h/sem ~5 meses: MAX 4 CUOTAS.').font = B; fila += 1
ws.cell(fila, 1, 'Separacion de cupo $300.000 (se cruza con la inicial) -> cuota inicial min 20% max 3 dias -> cuotas desde dia 30. Credito SIEMPRE con pagare + carta + plan de pagos firmados.').font = NORMAL
fila += 2
base = 2990000
c = ws.cell(fila, 1, 'PRECIO (contado o a cuotas): $2.990.000'); c.font = BLANCO_B; c.fill = NARANJA
for j in range(2, 6):
    ws.cell(fila, j).fill = NARANJA
fila += 1
ws.cell(fila, 1, 'Cuota inicial').font = B
for j, n in enumerate([2, 3, 4]):
    c = ws.cell(fila, 2 + j, str(n) + ' cuotas'); c.font = B; c.border = borde
ws.cell(fila, 5, 'Nota').font = B
fila += 1
for k, frac in enumerate(INICIALES):
    ini = round(base * frac)
    saldo = base - ini
    c = ws.cell(fila, 1, str(int(frac * 100)) + '% = $' + format(ini, ',').replace(',', '.'))
    c.font = NORMAL; c.border = borde
    for j, n in enumerate([2, 3, 4]):
        c = ws.cell(fila, 2 + j, cuota(saldo, n))
        c.number_format = MONEDA; c.font = NORMAL; c.border = borde
    if k == 0:
        ws.cell(fila, 5, 'Super intensivo: solo columna 2 cuotas').font = NORMAL
    fila += 1

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\INDUCCION_COMERCIAL\MOTOR_CUOTAS_FONDO_2026.xlsx'
wb.save(out)
print('OK', out)
