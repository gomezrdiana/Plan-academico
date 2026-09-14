# -*- coding: utf-8 -*-
"""Hoja de arranque del docente B2 — checklist de todo lo que hace si o si (por rol)."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NARANJA = RGBColor(0xE5, 0x4C, 0x09)
GRIS = RGBColor(0x77, 0x77, 0x77)
CREMA = 'FFF8E7'

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.5); s.bottom_margin = Cm(1.3); s.left_margin = Cm(1.9); s.right_margin = Cm(1.9)
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)

def shd(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement('w:shd'); el.set(qn('w:val'), 'clear'); el.set(qn('w:fill'), color)
    tcPr.append(el)

def titulo(txt, sub):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(17); r.font.color.rgb = NARANJA
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(sub); r.font.size = Pt(10); r.font.color.rgb = GRIS
    doc.add_paragraph()

def sec(txt):
    p = doc.add_paragraph(); r = p.add_run(txt); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = NARANJA
    p.paragraph_format.space_before = Pt(7); p.paragraph_format.space_after = Pt(3)

def t(txt, bold=False):
    p = doc.add_paragraph(); r = p.add_run(txt); r.font.size = Pt(11); r.bold = bold
    p.paragraph_format.space_after = Pt(3)

def frase(txt):
    tab = doc.add_table(rows=1, cols=1); tab.style = 'Table Grid'
    c = tab.rows[0].cells[0]; shd(c, CREMA)
    p = c.paragraphs[0]; r = p.add_run(txt); r.italic = True; r.font.size = Pt(11)
    esp = doc.add_paragraph(); esp_r = esp.add_run(''); esp_r.font.size = Pt(4)
    esp.paragraph_format.space_after = Pt(0)

titulo('HOJA DE ARRANQUE — DOCENTE B2',
       'Todo lo que se hace SÍ O SÍ, para arrancar sin dudas · Uso interno · v1 · Se entrega con el contrato')

sec('ANTES DE LA PRIMERA CLASE (en este orden)')
t('□ 1. CONTRATO FIRMADO — sin firma no se entrega ningún material.', bold=True)
t('□ 2. TU PROPIO EF SET (50 min, gratis, efset.org): aquí todos nos medimos, profes incluidos. El certificado va a coordinación — y además nos sirve comercialmente.')
t('□ 3. LEER: las guías de las últimas 5 clases del cohorte (continuidad) + la guía de la Clase 1 (ahí está el cronograma que ya se anunció a los estudiantes). Coordinación te las entrega impresas y te explica el estado del PROYECTO del cohorte y sus fechas — nada del proyecto o de evaluaciones se anuncia a estudiantes sin instrucción de coordinación.')
t('□ 4. ACTA DE PROGRAMACIÓN firmada: nivel, horario, horas, fechas. Tus viajes del año se avisan desde ya para programarlos.')

sec('CADA CLASE — EL RITUAL (no cambia nunca)')
t('□ Llegar 10 MINUTOS ANTES: revisar la guía del día, preparar el tablero, recibir indicaciones.')
t('□ Tablero listo ANTES de empezar: FRASE DEL DÍA escrita + la virtud de la semana.')
t('□ Ejecutar LA GUÍA TAL CUAL se entrega — 4 bloques largos, con los tiempos de la hoja de ruta. No se improvisa contenido ni se cambia el método; si algo no funciona, la guía trae Plan B, y lo demás se reporta a coordinación.')
t('□ B2 es ~94% ORAL: banco de vocabulario + trabalenguas + fase, como lo indique la guía del día.')
t('□ Simulaciones FÍSICAS y realistas con el patrón de la guía: UN estudiante protagonista + observadores activos + tú como coach desde afuera. TÚ NUNCA juegas el rol.')
t('□ CERO material impreso para estudiantes: todo va al tablero o de tarea. La escritura evaluable se hace SOLO en clase, a mano.')

sec('CIERRE DE CADA CLASE (los 10 minutos que sostienen todo)')
t('□ TICKET DE SALIDA: cada estudiante escribe lo suyo; se recogen TODOS, sin evaluar ni comentar.')
t('□ ERROR PAPER doble protocolo: papel físico ANÓNIMO que se entrega a coordinación + tu libreta privada CON nombres y citas literales, solo para coordinación.')
t('□ TAREA con fecha y hora explícitas (B2: entrega 6:00 PM), completa — no se fragmenta ni se negocia. Incluye SIEMPRE el audio/video diario del portafolio: el estudiante lo envía al WhatsApp de la academia (no a tu celular) — tu único trabajo es preguntar en clase quién lo hizo y anotarlo.')
t('□ REPORTE escrito y FIRMADO + el PAQUETE DE EVIDENCIA: 4 fotos al chat interno esa misma noche — (1) reporte firmado, (2) error paper, (3) 2-3 tickets, (4) el tablero con la Frase del Día. Es lo que la academia le firma al cliente en la garantía: sin evidencia no hay producto.', bold=True)

sec('LOS JAMASES (los mismos para todos los docentes)')
t('· JAMÁS comunicar a estudiantes notas, resultados de evaluaciones ni criterios — eso es exclusivo de coordinación.')
t('· JAMÁS mencionar nombres de metodologías o sistemas internos a estudiantes o terceros.')
t('· JAMÁS sacar material de la sede ni compartirlo digital — todo es Material Protegido del contrato.')
t('· JAMÁS actividad comercial en el aula — ni cursos propios, ni de terceros, ni los de la casa: lo comercial es de la asesora y gerencia.')
t('· JAMÁS cancelar o mover una clase sin avisar a coordinación con la antelación de la política interna.')

sec('TU CANAL COMERCIAL (por aparte del aula)')
t('Tu esquema de ventas es independiente y va en tu contrato: leads de la casa al 5% · tu canal propio autofinanciado al 10%, siempre sobre plata que entra a caja. Reglas: toda pieza de pauta con la marca la aprueba gerencia ANTES de publicar · todo lead se registra en la planilla comercial al primer contacto · y la base de estudiantes de Heiiu no se toca para nada distinto de Heiiu.')

frase('La regla de oro de la casa: la guía es la clase. Tu talento está en ejecutarla con calidad humana — no en reinventarla. Cualquier duda, coordinación; cualquier idea de mejora, se propone por escrito, no se prueba en el aula.')

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\documentos_operativos\plantillas\HOJA_ARRANQUE_DOCENTE_B2.docx'
doc.save(out); print('OK', out)
