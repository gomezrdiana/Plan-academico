# -*- coding: utf-8 -*-
"""KIT DE INDUCCION bonito: mismo estilo del Kit de Venta (colores Heiiu, letra grande)."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NARANJA = RGBColor(0xE5, 0x4C, 0x09)
NEGRO = RGBColor(0x22, 0x22, 0x22)
GRIS = RGBColor(0x77, 0x77, 0x77)
CREMA = 'FFF8E7'
NARANJA_HEX = 'E54C09'

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.8); s.bottom_margin = Cm(1.6); s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(12)

def shd(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement('w:shd'); el.set(qn('w:val'), 'clear'); el.set(qn('w:fill'), color)
    tcPr.append(el)

def seccion(num, txt):
    p = doc.add_paragraph()
    r = p.add_run(num + '  ' + txt); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = NARANJA
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(6)

def sub(txt):
    p = doc.add_paragraph()
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = NEGRO
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(3)

def texto(txt, size=12, bold=False):
    p = doc.add_paragraph(); r = p.add_run(txt); r.font.size = Pt(size); r.bold = bold
    p.paragraph_format.space_after = Pt(3)
    return p

def frase(txt, size=12.5):
    tab = doc.add_table(rows=1, cols=1); tab.style = 'Table Grid'
    c = tab.rows[0].cells[0]; shd(c, CREMA)
    p = c.paragraphs[0]; r = p.add_run(txt); r.italic = True; r.font.size = Pt(size)
    esp = doc.add_paragraph(); esp.paragraph_format.space_after = Pt(0)
    esp_r = esp.add_run(''); esp_r.font.size = Pt(4)

def tabla(filas, header=True, fs=11.5):
    tab = doc.add_table(rows=len(filas), cols=len(filas[0])); tab.style = 'Table Grid'
    for i, fila in enumerate(filas):
        for j, val in enumerate(fila):
            c = tab.rows[i].cells[j]
            for k, ln in enumerate(str(val).split('\n')):
                p = c.paragraphs[0] if k == 0 else c.add_paragraph()
                r = p.add_run(ln); r.font.size = Pt(fs)
                if header and i == 0:
                    r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            if header and i == 0:
                shd(c, NARANJA_HEX)
    esp = doc.add_paragraph(); esp.paragraph_format.space_after = Pt(0)
    esp_r = esp.add_run(''); esp_r.font.size = Pt(4)

def salto():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def leer(txt):
    p = doc.add_paragraph(); r = p.add_run('Leer: ' + txt); r.font.size = Pt(10.5); r.font.color.rgb = GRIS
    p.paragraph_format.space_after = Pt(6)

# ===== PORTADA =====
for _ in range(6): doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('HEIIU ENGLISH ACADEMY · CONFIDENCIAL — SOLO USO INTERNO'); r.font.size = Pt(14); r.font.color.rgb = GRIS
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('KIT DE INDUCCIÓN'); r.bold = True; r.font.size = Pt(34); r.font.color.rgb = NARANJA
for linea in ['Asesora Comercial · v2 · Septiembre 2026',
              'Meta: que vendas como el sistema Heiiu, no como te enseñaron en otro lado.',
              'Se avanza por dominio demostrado, no por calendario.',
              'Aquí todo ya está escrito — tu trabajo es ejecutarlo con calidad humana.']:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(linea); r.font.size = Pt(13); r.font.color.rgb = NEGRO
salto()

# ===== 1 =====
seccion('1.', 'QUÉ VENDE HEIIU')
texto('Heiiu no vende "clases de inglés". Vende dos cosas que tocan la identidad:', bold=True)
texto('1. PERTENENCIA: dejar de ser "el raro que quiere más" y estar rodeado de gente con la misma ambición.')
texto('2. RESPALDO: alguien apuesta por ti (la beca del Fondo) y te abre puertas VERIFICADAS (Au Pair, Work & Travel, H2B) — en un mercado lleno de agencias que abusan, Heiiu solo abre puertas seguras.')
sub('EL BIG DOMINO (memorizarlo — se siembra en la conversación, NUNCA se recita)')
frase('Si puedo convencer a un colombiano con chispa competitiva que quiere oportunidades reales de que la mejor manera de afilar esa chispa Y abrir puertas con el inglés NO es una app gratis (el 88% la abandona el primer mes — ahí estudias en soledad: nadie te reta, nadie apuesta por ti), NI una academia que vende un curso y te suelta, NI YouTube o IA — SINO un programa serio donde está rodeado de gente con su misma ambición + alianzas internacionales verificadas + un Fondo de Becas que premia al que se compromete… entonces todas las objeciones se caen y tendrá que invertir.', 11.5)
texto('Regla de balance: se lidera con PERTENENCIA, se refuerza con RESPALDO. Si la beca se vuelve protagonista, atraes gente que viene por la plata y no por el compromiso.', bold=True)


# ===== 2 =====
seccion('2.', 'PRODUCTOS Y PRECIOS')
tabla([
 ['Nivel', 'Horas', 'Precio público'],
 ['A1', '90 h', '$1.596.000'],
 ['A2', '110 h', '$1.915.000'],
 ['B1', '175 h', '$2.842.000'],
 ['B2', '200 h', '$2.343.000'],
 ['PROGRAMA COMPLETO', '575 h', '$8.696.000 — este valor NUNCA cambia'],
])
texto('ARRANQUE A1+A2 (lanzamiento, cohorte 5 de octubre): $2.990.000, 20 cupos, precio único SIN becas, desde 17 años. Jornadas, pago, prueba de ubicación, módulos de graduación y cierre: TODO en el KIT DE VENTA — el documento hermano de este. Se estudian juntos.', bold=True)
texto('Reglas: 100% presencial (grupos máx. 16) · sábados solo A1-A2-B1, por nivel suelto y SIN beca · B2 obligatorio entre semana · no hay horarios rotativos · Personalizados: $40.000/h nacional.')
texto('CAJASAN: se pregunta y se REGISTRA la afiliación en TODA cita (el reporte sostiene la alianza). El afiliado que paga a cuotas recibe la beca de contado. JAMÁS se envía a nadie a crédito o libranza de Cajasan — tienen programa propio de inglés.', bold=True)

# ===== 3 =====
seccion('3.', 'EL FONDO DE BECAS — el corazón del cierre')
frase('La palabra "PATROCINIO" NO EXISTE. Jamás. (Muerta legalmente — Ley 1480.) El término único es: BECA DEL FONDO.')
texto('Modelo híbrido: el PERFIL define el techo — Transformación 45% (SOLO estrato 1-2/SISBÉN con soporte, cupos limitados por trimestre) · GENERAL 32% · Ejecutivo 15%. El COMPROMISO define cuánto del techo se gana: paquete más largo hacia B2 + contado = más beca. La matriz completa en pesos está en el KIT DE VENTA (sección 3) — es back-office: nunca se muestra al cliente.')
sub('Tu guion (la postura que hace todo el trabajo)')
texto('· "Heiiu tiene un Fondo de Becas institucional, administrado por nuestra contaduría. Yo no decido las becas — yo presento tu solicitud y el Fondo la evalúa."')
texto('· Ante regateo: "No puedo tocar el porcentaje — el Fondo se audita y una excepción lo tumba para todas las familias. Lo que sí puedo hacer es presentar bien tu caso."')
texto('· Al aprobar: SIEMPRE se entrega carta impresa + reglamento.')
texto('El reglamento del Fondo (2 páginas) te lo entrega gerencia impreso — es el mismo que recibe el estudiante becado con su carta.')

salto()
# ===== 4 =====
seccion('4.', 'EL FUNNEL Y TU PUESTO EN ÉL')
frase('Meta/TikTok → bot de WhatsApp (precalifica y agenda) → CITA de 45-55 min → TÚ cierras → matrícula + kit firmado.')
texto('La cita es el producto que el bot te entrega; la matrícula es el producto que tú entregas. Tu métrica de vida: citas asistidas → matrículas → CAJA REAL. El hueco histórico del negocio no es tráfico: es cierre y cobro.', bold=True)
texto('Reglas: contacto en menos de 5 minutos al lead caliente · confirmación de cita el día antes y el día de · el que no asiste se rescata máximo 2 veces · nunca agendar sábado tarde, domingos ni festivos.')

# ===== 5 =====
seccion('5.', 'LA CITA Y LA ESCALERA DE CIERRE — tu oficio')
texto('Estructura: conexión (motivo real) → diagnóstico de perfil → programa presentado desde el DOMINÓ (pertenencia primero) → números de SU caso (beca aplicada) → cierre → si no cierra: escalera.')
sub('La escalera de cierre — la regla: nadie sale sin que se le ofrezca el escalón siguiente')
texto('SIEMPRE se empieza ARRIBA. Solo se baja un escalón cuando hay un NO real YA TRABAJADO (objeción respondida y aun así no). Bajar sin ese no = regalar plata. Saltar escalones = perder ventas.', bold=True)
tabla([
 ['Escalón', 'Qué se ofrece', 'Se baja cuando…'],
 ['1', 'PROGRAMA COMPLETO con su beca del Fondo (el número exacto sale de la tabla en pesos)', 'No real, ya respondido con plan de pagos y Fondo'],
 ['2', 'ARRANQUE A1+A2 — $2.990.000 (si empieza de cero) o el paquete menor hacia B2 (si ya tiene nivel)', 'No real trabajado'],
 ['3', 'NIVEL SUELTO (ej. solo A1: $1.596.000). Aquí — y SOLO aquí — está tu única carta: hasta 20% de descuento por pago de contado, SOLO en cita presencial, JAMÁS por teléfono o WhatsApp', 'No real trabajado'],
 ['4', 'SABATINO — la entrada más económica (4h/semana, por nivel, sin beca)', 'No real trabajado'],
 ['5', 'NO es una venta, es una SIEMBRA: se anota el MOTIVO REAL del no + teléfono, a la base de rescate', '—'],
], fs=10.5)
texto('El escalón 5 es el que paga el mes siguiente: cuando hagas la tanda de rescate no llamas a ciegas — llamas con el motivo en la mano ("me contaste que en diciembre te pagan la prima…"). El "no" de hoy con motivo anotado es la venta del mes siguiente.')
texto('Las dos trampas que te vigilas: (1) jamás ARRANCAR abajo — quien abre ofreciendo sabatino nunca vende un completo; (2) jamás bajar rápido — le enseñas al cliente que esperando consigue menos precio.', bold=True)
texto('La evidencia que respalda tu palabra: garantía firmada en contrato · avance medido con examen internacional · portafolio de audios y videos del estudiante · evaluador externo en los finales ("aquí nadie pasa por pasar").')
texto('El paso a paso operativo de la cita — frases, objeciones, checklist de las 10 — vive en el KIT DE VENTA. Este módulo te da el mapa; aquel te da el libreto.', bold=True)

# ===== 6 =====
seccion('6.', 'LO QUE NUNCA SE HACE — las líneas rojas de la casa')
texto('1. Decir "patrocinio/patrocinador" — el término es beca del Fondo.')
texto('2. Prometer empleo, visa, resultados en X meses, o beneficios en construcción.')
texto('3. Tocar precios de lista o inventar descuentos — la única flexibilidad es la que ya está escrita.')
texto('4. Comunicar evaluaciones de estudiantes (eso es SOLO de coordinación) o hablar de casos de otros estudiantes con nombre.')
texto('5. Separar el paquete en devoluciones — el paquete no es separable; las tarifas individuales son el escudo. Toda solicitud de retiro/devolución va a gerencia con el contrato en mano, JAMÁS se resuelve en caliente.')
texto('6. Prometer cupos o becas sin contrato firmado y primer pago.')

salto()
# ===== 7 =====
seccion('7.', 'REPORTES Y COMISIONES')
texto('REPORTE DIARIO (6:00 PM por WhatsApp, formato exacto): "Hoy contacté a: [nombres] = N · Citas agendadas: N · Citas asistidas: N · Cierres/abonos: N · Caja: $". Los contactos van POR NOMBRE — nuevos y seguimientos cuentan, y la misma persona puede repetirse en días distintos. Día sin reporte = día sin actividad verificable.', bold=True)
texto('REPORTE SEMANAL (cada lunes antes del mediodía): citas agendadas y asistidas · matrículas con nombres y valores · caja que ENTRÓ · motivo #1 de no-cierre · lista de "lo voy a pensar" con teléfonos. Lo que no llegue queda SIN DATO en el tablero de gerencia — y sin dato se nota.', bold=True)
texto('COMISIONES: 5% del recaudo de cliente nuevo · 2,5% de cliente antiguo — sobre lo que ENTRA a caja, no sobre lo vendido. Si el cliente no paga la cuota, no hay comisión: cobrar también es vender. Liquidación mensual contra el reporte de ingresos.')
texto('KIT DE MATRÍCULA COMPLETO, SIEMPRE: contrato + pagaré con carta de instrucciones + anexo plan de pagos + checklist + (si aplica) carta de beca del Fondo + acuerdo de asistencia. Con huellas y copia para el estudiante.')

# ===== 8 =====
seccion('8.', 'CANALES ESPECIALES')
texto('· ALIANZAS INTERNACIONALES (Au Pair, Work & Travel, H2B): canal probado (16 matrículas en 2024 solo por Au Pair). Posicionamiento: Heiiu es el FILTRO SEGURO — "trabajamos solo con operadores verificados". El inglés es el requisito del camino — y el Módulo Pasaporte es el producto para esta gente.')
texto('· ALIANZAS EMPRESARIALES: las empresas premian a su gente con Becas [Empresa]; los cierres corporativos y todo lo B2B pago pasan por gerencia. Tú atiendes a los BECADOS que lleguen a matricularse — regla del mejor beneficio: si el becado califica a una beca superior del Fondo, se le aplica la mejor, nunca ambas.')
texto('· RESCATE DE LEADS: tandas de mensajes a la base de "lo voy a pensar" — siempre lenguaje Fondo de Becas, con plantillas aprobadas. Nunca improvisar oferta.')

# ===== PLAN =====
seccion('9.', 'EL PLAN PARA QUIEN LLEGA NUEVA (referencia institucional)')
texto('Si ya vienes con oficio comercial, la mayoría de estas etapas las recorres en días — o ya las tienes superadas. Lo único que aplica a TODAS, experta o no, es la evaluación de dominio de abajo: no mide tu experiencia vendiendo (esa ya la trajiste), mide el SISTEMA HEIIU — el Fondo, la escalera, las líneas rojas — que nadie trae de fábrica. Cada etapa se supera demostrando la verificación de la derecha, el mismo día si la demuestras.', bold=True)
tabla([
 ['Etapa', 'Actividad', 'Se supera cuando…'],
 ['1', 'Módulos 1-2 + tour + observar 1 cita real', 'Explica el Big Domino en sus palabras, sin leer'],
 ['2', 'Módulo 3 (Fondo) + role-play de regateo + sesión de objeciones reales con la asesora senior', 'Sostiene el "no puedo tocar el porcentaje" 3 veces seguidas'],
 ['3', 'Módulo 4 + escuchar llamadas y chats del bot', 'Dibuja el funnel de memoria con sus métricas'],
 ['4', 'Módulos 5-6-7 + role-play de cita completa (gerencia de prospecto difícil) + 2ª sesión de cierres', 'Cierra el role-play usando la escalera sin inventar descuentos'],
 ['5', 'Citas reales ACOMPAÑADA (la asesora senior presente; la nueva lidera)', '2 citas lideradas sin corrección mayor'],
 ['6', 'Citas sola + Módulo 8 + primer reporte semanal real', 'Reporte entregado completo el lunes'],
 ['7', 'EVALUACIÓN DE DOMINIO (abajo)', '8/10 mínimo — si no, refuerzo y se repite antes de soltar'],
], fs=11)

sub('Evaluación de dominio (10 preguntas — responder sin consultar)')
texto('1. ¿Qué dos cosas compra realmente el estudiante Heiiu? · 2. Diga el Big Domino en sus palabras. · 3. ¿Por qué está prohibida la palabra "patrocinio" y qué se dice en su lugar? · 4. Estrato 3 quiere A2→B2 a cuotas: ¿qué beca le corresponde? (GENERAL, 19,2%.) · 5. ¿Los 5 escalones de la escalera de cierre? · 6. "Me hacen 40% o no firmo": ¿respuesta exacta? · 7. ¿Qué lleva el kit de matrícula completo? · 8. ¿Sobre qué se paga su comisión y por qué? · 9. Un becado de empresa aliada resulta estrato 1: ¿qué se le aplica? · 10. Cliente pide devolución parcial del paquete: ¿qué hace usted?', 11)

doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Heiiu English Academy · Kit de Inducción v2 · Se actualiza con cada cambio de doctrina — responsable: gerencia.'); r.font.size = Pt(10); r.font.color.rgb = GRIS

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\INDUCCION_COMERCIAL\KIT_INDUCCION_ASESORA_COMERCIAL.docx'
doc.save(out); print('OK', out)
