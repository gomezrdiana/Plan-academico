# -*- coding: utf-8 -*-
"""Mensajes de la encuesta al grupo de Margarita Pasos (Diana + Carlos): reemplazan el envio del video de 45 min.
Salida: recursos/comercial/CFO/MENSAJES_ENCUESTA_GRUPO_MARGARITA.docx"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x16, 0x33, 0x5A); GRIS = RGBColor(0x66, 0x66, 0x66)
d = Document()
for s in d.sections:
    s.top_margin = Cm(1.6); s.bottom_margin = Cm(1.4); s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)
d.styles['Normal'].font.name = 'Calibri'; d.styles['Normal'].font.size = Pt(10.5)

def shd(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear'); el.set(qn('w:color'), 'auto'); el.set(qn('w:fill'), color); tcPr.append(el)

def h(txt, size=12.5, before=10):
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = NAVY

def t(txt, bold=False, after=4, size=10.5, color=None, italic=False):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(txt); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if color: r.font.color.rgb = color

def caja(txt):
    tb = d.add_table(rows=1, cols=1); tb.style = 'Table Grid'
    c = tb.rows[0].cells[0]; shd(c, 'F3F1EA')
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    r = p.add_run(txt); r.font.size = Pt(11)
    d.add_paragraph().paragraph_format.space_after = Pt(2)

h('MENSAJES AL GRUPO DE MARGARITA PASOS — ENCUESTA', 15, before=0)
t('Reemplazan el mensaje del video de 45 minutos. Se copian del cuadro gris y se pegan en WhatsApp, cambiando [nombre].', color=GRIS, size=9.5, after=8)

h('CÓMO SE ENVÍA')
for x in [
    '1. Uno por uno, con el nombre de la persona. Nunca como lista de difusión ni todos en el mismo minuto.',
    '2. Primero a 20 personas. Si contestan 2 o más de cada 10, se sigue con el resto. Si contestan menos, se ajusta la pregunta antes de seguir: a cada persona se le escribe una sola vez con esto.',
    '3. A los que ya recibieron el video de YouTube se les manda el mensaje 1B, 4 o 5 días después del primero.',
    '4. El test y la oferta solo se mandan a quien contesta. Nada de enlaces en el primer mensaje.',
    '5. El enlace de WhatsApp correcto de Carlos es wa.me/573102748487 (el de "wa.me/@ceremolina" no funciona). La landing corta es ceremolina.com/aplica.html.',
]:
    t(x, after=3)

h('MENSAJE 1A — para quienes NO han recibido nada')
caja('Hola [nombre], soy Diana, del grupo de Margarita Pasos. Con mi socio, que es financiero, estamos haciendo una pequeña encuesta entre empresarios, porque vemos algo que se repite mucho: la mayoría busca a un financiero solo cuando ya tiene una situación grave o urgente. Queremos entender qué pasa antes de llegar a ese punto. ¿Qué te cuesta más hoy en tu negocio: vender, cobrar o saber cuánto te queda al final del mes? Con una palabra me ayudas mucho.')

h('MENSAJE 1B — para quienes YA recibieron el video de YouTube (4 o 5 días después)')
caja('Hola [nombre], hace unos días te mandé un video largo; perdón, era mucho. Mejor te hago una sola pregunta. Con mi socio, que es financiero, estamos haciendo una pequeña encuesta entre empresarios, porque vemos algo que se repite mucho: la mayoría busca a un financiero solo cuando ya tiene una situación grave o urgente. Queremos entender qué pasa antes de llegar a ese punto. ¿Qué te cuesta más hoy en tu negocio: vender, cobrar o saber cuánto te queda al final del mes? Con una palabra me ayudas mucho.')

h('MENSAJE 2 — solo a quien conteste (según su respuesta)')
t('Si dice "SABER CUÁNTO ME QUEDA":', bold=True, after=2)
caja('Gracias, [nombre]. Es la que más se repite, y es la más peligrosa: muchas empresas pierden plata durante meses con el saldo sano. Carlos tiene un test de 7 preguntas que en un minuto te dice si es tu caso. ¿Te lo mando?')
t('Si dice "COBRAR":', bold=True, after=2)
caja('Gracias, [nombre]. Cobrar mal es plata que ya ganaste y no tienes en la mano, y casi siempre termina tapándose con crédito. Carlos tiene un test de 7 preguntas que en un minuto te muestra si se te está quedando plata por fuera. ¿Te lo mando?')
t('Si dice "VENDER":', bold=True, after=2)
caja('Gracias, [nombre]. Una cosa que Carlos repite mucho: vender más no siempre deja más plata. Él cuenta el caso de un negocio en Florida que vendía más cada mes y perdía. Son 90 segundos, ¿te lo mando?')
t('(Para "vender" se manda el CLIP CORTO de la entrevista, de 60 a 90 segundos, con el caso de Florida. Pendiente: que Carlos lo recorte. Mientras no exista, se ofrece el test igual que en las otras respuestas.)', size=9, color=GRIS, italic=True)
t('Si contesta otra cosa o varias:', bold=True, after=2)
caja('Gracias, [nombre], me sirve mucho. Carlos tiene un test de 7 preguntas que en un minuto te muestra cómo están los números de tu negocio. ¿Te lo mando?')
t('Si dice que no le interesa o no quiere contestar:', bold=True, after=2)
caja('Tranquilo, [nombre], gracias por responder. Un abrazo.')
t('(Y no se le vuelve a escribir sobre esto.)', size=9, color=GRIS, italic=True)

h('MENSAJE 3 — cuando dice "sí, mándalo"')
caja('Aquí está. Son 7 preguntas de sí o no, te toma un minuto. Cuando lo hagas, cuéntame cuántos "NO" te salieron.')
t('(Se adjunta CARLOS_TEST_7_NO.pdf, de la carpeta recursos/comercial/CFO.)', size=9, color=GRIS, italic=True)

h('MENSAJE 4 — cuando dice cuántos NO le salieron')
t('Si le salieron 3 o más NO:', bold=True, after=2)
caja('Gracias por hacerlo, [nombre]. Con [X] NO vale la pena mirar los números de cerca. Carlos hace una Radiografía Financiera: tú entregas tus documentos y en 3 páginas, sin tecnicismos, te dice dónde se escapa la plata, cuánto tienes que vender para no perder y cómo te ve un banco hoy. Acepta 5 empresas al mes. Si quieres postular la tuya: ceremolina.com/aplica.html. O escríbele directo y dile que vas de mi parte: wa.me/573102748487')
t('Si le salieron 0 a 2 NO:', bold=True, after=2)
caja('¡Muy bien, [nombre]! Estás mejor que la mayoría. Si algún día quieres una segunda mirada a tus números antes de pedir un crédito o de crecer, Carlos es la persona: wa.me/573102748487. Gracias por ayudarnos con la encuesta.')

h('DESPUÉS: LOS RESULTADOS AL GRUPO (solo si se van a contar bien)')
t('Cuando haya 30 o más respuestas, se comparte el resultado en el grupo. Es lo que se promete al llamarlo encuesta, y es el mejor contenido para Carlos porque sale de empresarios reales.', after=3)
caja('Gracias a todos los que respondieron nuestra pregunta. De [N] empresarios del grupo, [X] dijeron que lo que más les cuesta es saber cuánto les queda al final del mes, [Y] cobrar y [Z] vender. Carlos preparó un test de 7 preguntas, de un minuto, para quien quiera revisar su caso: escríbanme y se los mando.')

h('REGISTRO (para medir)')
tb = d.add_table(rows=11, cols=5); tb.style = 'Table Grid'
for j, x in enumerate(['Nombre', 'Fecha envío', '1A o 1B', 'Respuesta', 'Siguiente paso']):
    c = tb.rows[0].cells[j]; shd(c, 'F3F1EA'); r = c.paragraphs[0].add_run(x); r.bold = True; r.font.size = Pt(9.5)

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\CFO\MENSAJES_ENCUESTA_GRUPO_MARGARITA.docx'
d.save(out); print('OK', out)
