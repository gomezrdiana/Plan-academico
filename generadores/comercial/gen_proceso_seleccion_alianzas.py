# -*- coding: utf-8 -*-
"""Proceso de selección de 5 filtros — Ejecutivo(a) Senior de Alianzas Empresariales (basado en investigación 10/10/2026:
Sackett et al. 2022/2023, Schmidt & Hunter 1998, Vinchur 1998, Kuncel 2013, Taylor 2004; prácticas de Google, Amazon, Bain, P&G; normas colombianas).
Salida: recursos/comercial/INDUCCION_COMERCIAL/contratacion comercial/PROCESO_SELECCION_ALIANZAS.docx"""
import re
from docx import Document
from docx.shared import Pt, RGBColor

NAR = RGBColor(0xE5, 0x4C, 0x09); GRIS = RGBColor(0x66, 0x66, 0x66)
d = Document(); d.styles['Normal'].font.name = 'Calibri'; d.styles['Normal'].font.size = Pt(10.5)
for s in d.sections: s.left_margin = s.right_margin = Pt(54); s.top_margin = s.bottom_margin = Pt(48)
def h(t): p = d.add_paragraph(); p.paragraph_format.space_before = Pt(12); r = p.add_run(t); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = NAR
def p(t, color=None, size=10.5, after=4):
    x = d.add_paragraph(); x.paragraph_format.space_after = Pt(after)
    for part in re.split(r'(\*\*[^*]+\*\*)', t):
        if not part: continue
        r = x.add_run(part[2:-2] if part.startswith('**') else part); r.bold = part.startswith('**'); r.font.size = Pt(size)
        if color: r.font.color.rgb = color
def tabla(rows):
    t = d.add_table(rows=1, cols=len(rows[0])); t.style = 'Table Grid'
    for j, c in enumerate(rows[0]):
        r = t.rows[0].cells[j].paragraphs[0].add_run(c); r.bold = True; r.font.size = Pt(9)
    for f in rows[1:]:
        cells = t.add_row().cells
        for j, v in enumerate(f):
            r = cells[j].paragraphs[0].add_run(v); r.font.size = Pt(9)
    d.add_paragraph().paragraph_format.space_after = Pt(2)

p('PROCESO DE SELECCIÓN — EJECUTIVO(A) SENIOR DE ALIANZAS EMPRESARIALES', size=15, color=NAR)
p('5 filtros + período de prueba. Basado en la evidencia de qué predice el desempeño (metaanálisis de Sackett et al. 2022-2023 y otros) y en cómo contratan Google, Amazon, Bain y P&G, adaptado a una pyme: unas 3 horas de tu tiempo por finalista, 10 días hábiles.', GRIS, 9.5)

h('LO QUE DICE LA EVIDENCIA (en una línea cada una)')
for t in ['**La entrevista ESTRUCTURADA es el mejor método** (validez .42): mismas preguntas para todos, rúbrica escrita y puntaje. La entrevista libre "para conocerlo" vale la mitad (.19) y además confunde.',
          '**Las pruebas específicas del cargo** (simulación, tarea práctica, conocimiento del trabajo: .33-.40) predicen más que las pruebas generales de personalidad o de IQ. En ventas, el IQ casi no predice lo que se vende (.04).',
          '**Los años de experiencia casi no predicen** (.07): los 5 años solo limpian el embudo, no dicen quién vende.',
          '**Ningún filtro solo basta**: combinarlos sube la validez (~.47). Por eso son 5.',
          '**Decidir con fórmula supera al juicio "a ojo"** en más de 50% (Kuncel 2013). Los criterios se escriben ANTES de ver candidatos.',
          '**Las referencias estructuradas por teléfono** (10-15 min, preguntas fijas) sí predicen (.25-.36).',
          '**Los candidatos más encantadores y menos honestos sacan mejores entrevistas** (Frontiers 2017): las cifras se verifican con documentos y con exjefes, no con la palabra del candidato.']:
    p('• ' + t, after=3)

h('ANTES DE VER CANDIDATOS (una sola vez, 30 min)')
p('Firma esta hoja con la rúbrica, los pesos y los puntos de corte. No se cambian después de conocer a nadie.')

h('LOS 5 FILTROS')
tabla([
 ['Filtro', 'Quién', 'Tiempo', 'Qué se evalúa', 'Pasa si…'],
 ['F1 · Hoja de vida + audio de 1 min', 'Asesora con lista de chequeo; tú solo ves a los que pasan', '5 min c/u', '¿Hay al menos una venta B2B con cifra, empresa cliente y decisor gerencial? ¿El audio dice qué hizo ÉL/ELLA, con número?', 'Hay número y decisor. Sin eso, fuera.'],
 ['F2 · Llamada estructurada', 'Tú o la asesora, con guion', '20 min', 'Eliminatorias: (1) disponible presencial en Bucaramanga; (2) acepta el esquema (mínimo legal + comisión, piso $3,5M con indicadores, prueba de 2 meses, tablero diario y seguimiento diario); (3) puede entrar en ≤1 semana. Pedir: meta mensual, % de cumplimiento del último año, ticket promedio, ciclo de venta y nombre de 2 jefes directos.', 'Pasa las 3 eliminatorias. Se le envía la autorización de datos (Ley 1581) y se le pide traer soportes de cifras a F4.'],
 ['F3 · Tarea práctica (sin pago, máx. 90 min, 48 h)', 'Candidato; tú calificas', '20 min tuyos', '(a) 15 empresas reales del área metropolitana (40+ empleados) con el cargo del decisor y por qué encajan; (b) mensaje de apertura al decisor (máx. 120 palabras); (c) plan de sus 2 primeras semanas frente a los mínimos del mes 1.', 'Puntaje ≥ 3/5. (Se le dice por escrito que Heiiu no usa su lista si no lo contrata.)'],
 ['F4 · Entrevista estructurada + simulación', 'Tú + un segundo evaluador (asesora o docente líder). Califican POR SEPARADO antes de hablar', '90 min', '40 min: 6 preguntas STAR (sección "Preguntas"); 25 min: simulación "gerente de empresa de 40 empleados" con 4 objeciones fijas; 15 min: defensa de la tarea y revisión de soportes de cifras; 10 min: preguntas del candidato.', 'Promedio ≥ 3,5 y ninguna eliminatoria < 3.'],
 ['F5 · Verificación', 'Tú', '45 min', 'Soportes de cifras (desprendibles con comisiones) + 2 referencias de jefes directos (guion) + antecedentes públicos (Procuraduría, Contraloría, Policía).', 'Las cifras cuadran y hay referencias verificables. Si no, NO hay oferta.'],
])
p('**Decisión con fórmula:** Entrevista STAR 35% · Simulación 30% · Tarea 15% · Verificación 20%. Se ofrece al mayor puntaje si llega a ≥ 3,5/5 y pasó F5. **Si nadie llega, no se contrata.**')
p('**Calendario:** días 1-3 F1 y F2 · días 3-5 F3 · días 6-7 F4 (2-3 finalistas el mismo día, para comparar) · días 8-9 F5 · día 10 oferta y firma.', GRIS, 9.5)

h('RÚBRICA (1 = no cumple · 3 = a medias · 5 = excelente)')
tabla([
 ['Dimensión', '1', '3', '5'],
 ['Prospección propia (ELIMINATORIA)', 'Espera leads o le daban la base', 'Tiene método, sin cifras', 'Construyó cartera desde cero, con números de contactos → reuniones → cierres'],
 ['Venta a decisor y ciclo largo', 'Solo venta a persona natural', 'Una venta B2B contada a medias', '2+ cierres B2B con varios contactos, paso a paso y con cifra'],
 ['Honestidad de cifras (ELIMINATORIA)', 'Cifras vagas o que cambian', 'Cifras sin soporte', 'Cifras iguales en F2, F4 y F5, con soporte'],
 ['Disciplina de registro y reporte', 'Rechaza el seguimiento diario', 'Lo acepta sin historia previa', 'Muestra cómo reportaba, con ejemplo'],
 ['Ética (ELIMINATORIA)', 'Justifica prometer lo que no hay', 'Ambiguo', 'Caso real de perder una venta por no mentir'],
 ['Simulación: diagnóstico', 'Ofrece sin preguntar', '2-3 preguntas genéricas', 'Descubre el dolor de la empresa antes de ofrecer'],
 ['Simulación: propuesta', 'Habla del curso', 'Menciona las 2 ofertas', 'Conecta alianza vs. plan empresa con el dolor y usa la garantía'],
 ['Simulación: objeciones', 'Discute o regala descuento', 'Responde 2 de 4', 'Responde las 4 sin bajar precio'],
 ['Simulación: cierre', 'Termina en "lo pienso"', 'Pide otra reunión vaga', 'Cierra fecha concreta (charla, diagnóstico o firma)'],
 ['Tarea F3', 'Empresas genéricas o inventadas', 'Lista correcta sin decisor', 'Lista real con decisor, mensaje breve y plan que cubre los mínimos'],
])
p('**Las 4 objeciones fijas de la simulación:** 1) "Ya tenemos convenio con otra academia." 2) "Eso es gasto, no es prioridad este año." 3) "Mándeme la información al correo." 4) "Lo tengo que consultar con la junta."')

h('PREGUNTAS STAR (escoge 6, las mismas para todos)')
for i, (q, b) in enumerate([
 ('Cuénteme la última cartera B2B que construyó desde cero. ¿Cómo consiguió las primeras 10 empresas y en qué números cerró el primer trimestre?', 'Fuentes concretas, embudo con números. Alerta: "me pasaban la base".'),
 ('Su venta B2B más larga: ¿cuántos meses, cuántos contactos, quién decidía y qué pasó en cada etapa?', 'Varios interlocutores, manejo del comité de decisión.'),
 ('Una venta grande que perdió en la etapa final. ¿Qué hizo usted y qué cambió después?', 'Responsabilidad propia; no culpa al precio.'),
 ('Su peor mes del último año contra la meta. ¿Qué hizo en la segunda quincena?', 'Cifra exacta de meta y resultado.'),
 ('¿Cómo llegó a un gerente que no le contestaba? Deme el caso concreto.', 'Canales alternos, referidos, persistencia con criterio.'),
 ('Descríbame su registro en CRM o Excel: ¿qué campos llevaba, cada cuánto reportaba y a quién?', 'Rutina concreta. Alerta: incomodidad con el seguimiento.'),
 ('Una vez que trabajó sin jefe encima durante semanas: ¿cómo organizó su agenda y qué entregó?', 'Autonomía con evidencia.'),
 ('Un cliente le pidió una condición fuera de política. ¿Qué hizo?', 'Escaló o sostuvo la política sin perder la relación.'),
 ('Una vez que pudo cerrar prometiendo algo que no era cierto. ¿Qué hizo?', 'Un caso real y su costo. Alerta: "nunca me ha pasado".'),
 ('Una alianza que firmó y no produjo. ¿Por qué y qué hizo para activarla?', 'Entiende que firmar no basta.')], 1):
    p(f'**{i}. {q}**', after=1); p('   Buscar: ' + b, GRIS, 9.5, 3)
p('**Repreguntas fijas en todas:** "¿Qué hizo usted?" · "¿En qué número terminó?" · "¿Quién me lo confirma?"')

h('LLAMADA DE REFERENCIAS (con autorización escrita del candidato, 10 min)')
p('Apertura: "Hablo de Heiiu. [Candidato] nos autorizó por escrito a contactarle como su jefe directo en [empresa], entre [fechas]. Son 10 minutos."')
for t in ['1. ¿Cuál era su meta mensual o trimestral y en qué porcentaje la cumplía normalmente? (contrastar con lo que dijo)',
          '2. ¿Abría clientes nuevos por su cuenta o trabajaba cartera asignada?',
          '3. De 1 a 5, ¿cómo era su disciplina de reporte y seguimiento? ¿Por qué no un 5?',
          '4. ¿En qué necesitaba más acompañamiento?',
          '5. ¿Lo volvería a contratar para ventas? ¿Sí, no, con reservas? (una vacilación se anota)']:
    p(t, after=2)
p('**Prueba de cifras (se pide en F2, se trae a F4):** desprendibles de nómina con la línea de comisiones de 3 a 6 meses (o el anexo de comisiones), certificado laboral, y una ficha de 1 página con sus 3 mayores cierres (empresa, valor, fecha), que se contrasta con los referentes.')

h('SEÑALES DE ALARMA')
for t in ['Habla en "nosotros" y no aísla su aporte, o no recuerda su meta ni su cumplimiento.', 'Las cifras cambian entre la llamada, la entrevista y la verificación.',
          'No trae soportes o "los perdió todos".', 'Da colegas como referencia, no jefes directos, o los jefes "no están disponibles".',
          'Toda su experiencia es con base asignada o venta de mostrador presentada como B2B.', 'Cadena de empleos de menos de 12 meses con historias de "la empresa no cumplió".',
          'Se incomoda con el tablero diario o pide "sin mediciones".', 'La tarea trae empresas genéricas, o hecha con IA sin decisor real (se nota al defenderla).',
          'En la simulación ofrece descuento ante la primera objeción o promete lo que no existe.', 'Mucho encanto y poca evidencia.']:
    p('• ' + t, after=2)

h('PERÍODO DE PRUEBA = ÚLTIMO FILTRO (por escrito, firmado al ingresar)')
tabla([
 ['Día', 'Hito mínimo', 'Si no cumple'],
 ['15', '≥ 20 empresas con decisor en el tablero, ≥ 6 reuniones agendadas, seguimiento diario sin faltas', 'Plan de ajuste escrito de 1 página'],
 ['30', 'Mínimo del mes 1: 40 empresas con decisor, 12 reuniones, 2 firmas', 'Sin ingreso mínimo de $3,5M ese mes + alerta escrita'],
 ['50', 'Camino al mes 2 (≥ 3 firmas acumuladas, 1 activación en marcha) y tablero al día', 'DECISIÓN el día 50: continúa o termina en período de prueba, antes del día 60'],
])

h('QUÉ NO HACER (legal y sesgos)')
for t in ['**No consultar Datacrédito ni centrales de riesgo** para decidir una contratación: la ley lo prohíbe (Ley 1266 de 2008, modificada por Ley 2157 de 2021).',
          '**No aplicar tests psicológicos tú misma**: el material psicotécnico es exclusivo de psicólogos (Ley 1090 de 2006, art. 45). La rúbrica y la simulación no son tests psicológicos.',
          '**No preguntar** por embarazo, hijos, estado civil, religión, política, salud ni orientación sexual (discriminación, Ley 1482 de 2011; datos sensibles, Ley 1581).',
          '**Pedir autorización escrita** para tratar sus datos y llamar referencias (Ley 1581 de 2012).',
          '**El período de prueba va por escrito** y es de máximo 2 meses (CST arts. 76-80). La forma de terminarlo, confírmala con el abogado.',
          '**No cobrar nada** al candidato.',
          '**No hacer entrevista libre**, no cambiar la rúbrica después de ver al candidato, no hablar con el otro evaluador antes de calificar, no dar puntos por años de experiencia ni por desenvoltura, y no contratar "porque no hay más".']:
    p('• ' + t, after=3)
p('Fuentes: Sackett, Zhang, Berry y Lievens (2022) JAP doi:10.1037/apl0000994 y (2023) IOP doi:10.1017/iop.2023.24 · Schmidt y Hunter (1998) · Vinchur et al. (1998) · Verbeke et al. (2011) · Van Iddekinge et al. (2019) · Kuncel et al. (2013) · Taylor et al. (2004) · Dana et al. (2013) · Google re:Work, Amazon Jobs, Bain, P&G · Leyes 1581/2012, 1266/2008, 1090/2006, 1482/2011, CST. Sin verificar: etapas de selección de Bancolombia, Sura, Nutresa; jurisprudencia sobre terminación en período de prueba.', GRIS, 8.5)

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\INDUCCION_COMERCIAL\contratacion comercial\PROCESO_SELECCION_ALIANZAS.docx'
try: d.save(out); print('OK', out)
except PermissionError: d.save(out.replace('.docx', '_v2.docx')); print('ABIERTO -> _v2')
