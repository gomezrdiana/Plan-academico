# -*- coding: utf-8 -*-
"""Entrevista de NÚMEROS que hace el socio financiero (Carlos) a Reinaldo (candidato de alianzas), 30 min virtual, 2 págs.
Misma estructura que la de John para que Carlos compare a los dos con la misma vara. Salida en Candidatos/ (fuera de git):
  - ENTREVISTA_CARLOS_REINALDO_NUMEROS.docx (para Carlos)
  - LLAMADA_REFERENCIA_BBVA_DIANA.docx (para Diana: 6 preguntas al jefe directo en BBVA + 2 al de Bancolombia)"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAR = RGBColor(0xE5, 0x4C, 0x09); GRIS = RGBColor(0x66, 0x66, 0x66)
BASE = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\INDUCCION_COMERCIAL\contratacion comercial\LINKEDIN DIANA GOMEZ\Candidatos'

def nuevo_doc():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.2); s.bottom_margin = Cm(1.0); s.left_margin = Cm(1.8); s.right_margin = Cm(1.8)
    doc.styles['Normal'].font.name = 'Calibri'; doc.styles['Normal'].font.size = Pt(10)
    return doc
def shd(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear'); el.set(qn('w:color'), 'auto'); el.set(qn('w:fill'), color); tcPr.append(el)
def titulo(doc, t1, t2):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
    r = p.add_run(t1); r.bold = True; r.font.size = Pt(13.5); r.font.color.rgb = NAR
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(6)
    r = p.add_run(t2); r.font.size = Pt(9); r.font.color.rgb = GRIS
def sec(doc, txt):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(7); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(txt); r.bold = True; r.font.size = Pt(11.5); r.font.color.rgb = NAR
def t(doc, txt, bold=False, size=10, color=None, after=3):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(txt); r.font.size = Pt(size); r.bold = bold
    if color: r.font.color.rgb = color
def preg(doc, txt):
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(1)
    r = p.add_run(txt); r.font.size = Pt(10)
def busca(doc, txt):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4); p.paragraph_format.left_indent = Cm(0.6)
    r = p.add_run('Qué buscas: '); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = GRIS
    r = p.add_run(txt); r.font.size = Pt(9.5); r.font.color.rgb = GRIS
def caja(doc, lines):
    tb = doc.add_table(rows=1, cols=1); tb.style = 'Table Grid'; c = tb.cell(0, 0); shd(c, 'FFF8E7')
    c.paragraphs[0].paragraph_format.space_after = Pt(0)
    r = c.paragraphs[0].add_run('\n'.join(lines)); r.font.size = Pt(9.5); r.italic = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

# ---------- 1. ENTREVISTA PARA CARLOS ----------
doc = nuevo_doc()
titulo(doc, 'ENTREVISTA DE NÚMEROS — REINALDO JOYA', 'Cargo: Ejecutivo Senior de Alianzas Empresariales · La hace Carlos, socio financiero · Virtual · 30 minutos · Uso interno, no se entrega al candidato')
t(doc, 'Misma hoja que la de John, con sus cifras: así comparas a los dos con la misma vara. Lo que Reinaldo ya contó (BBVA, Bancolombia, por qué salió en mayo y qué hizo en los dos huecos de su hoja de vida) lo verifica Diana con la ficha y con su jefe directo de BBVA. Esta entrevista no repite esa historia: lo pone a construir números en vivo. Un comercial de verdad se sabe su embudo de memoria; el que no, da cifras redondas y vagas. Tú anotas cada número exacto que diga.', bold=True)
t(doc, 'Lo que ya tienes para leer antes: la ficha de cifras que llenó en la reunión con Diana, su tarea práctica (5 empresas con decisor, mensaje de apertura, plan de 2 semanas) y la propuesta del cargo. Su HV dice: canal de concesionarios de BBVA en Santander creado desde cero (2021), más de 10 concesionarios con convenio, $12.000 millones colocados al año; antes, 10 años en Bancolombia llevando una zona de $500 a $7.000 millones. Viene de banca: vendía con la marca de un banco detrás.', size=9.5, color=GRIS)

sec(doc, 'ARRANQUE (1 min)')
caja(doc, ['"Soy Carlos, socio financiero de Heiiu. Reviso los números de todo cargo comercial. Diana ya habló contigo de la causa; vamos a hablar solo de cifras."'])

sec(doc, '1. SU EMBUDO EN BBVA, DE MEMORIA (8 min)')
preg(doc, '"En una semana normal en BBVA: ¿cuántos concesionarios o vendedores visitabas, cuántos te recibían, cuántos firmaban convenio y cuánto tardaba el primer cliente referido en llegar?"')
preg(doc, '"De los más de 10 concesionarios con convenio: ¿cuántos firmaste tú solo, cuánto tardaste en cada uno, y cuántos te referían clientes de verdad el último mes?"')
preg(doc, '"De los $12.000 millones del año: ¿qué porcentaje venía de los 3 concesionarios más grandes? ¿Cuántas operaciones eran y de qué ticket promedio?"')
preg(doc, '"¿Cuánto le costaba a BBVA cada concesionario activo, contando tu salario y la comisión del convenio? ¿Lo calculaste alguna vez?"')
busca(doc, 'números específicos y rápidos (ej. "visitaba 8, me recibían 6, firmaba 1 cada mes y medio, el primer referido llegaba a las 2 semanas"). Si dice "muchos", "bastante", "depende", o se demora, no manejaba sus indicadores: se los manejaban. Ojo: en banca los referidos llegan por la marca; que diga cuánto era suyo y cuánto era del banco.')

sec(doc, '2. CÓMO REPORTABA (4 min)')
preg(doc, '"¿Qué le reportabas a tu coordinador nacional, cada cuánto y en qué formato? Descríbeme el tablero."')
preg(doc, '"¿Cuál era el indicador que te medían primero? ¿Y cuál era el que tú mirabas primero?"')
preg(doc, '"Cuéntame un mes en que ibas por debajo de meta: qué viste en los números, en qué semana, y qué cambiaste."')
busca(doc, 'que mencione indicadores de actividad (visitas, contactos), de conversión y de plata, y que sepa cuál fue el que le avisó. Si solo habla del resultado final, no leía el proceso. Su HV nombra Power BI y CRM: que describa lo que él miraba, no la herramienta.')

sec(doc, '3. QUE ARME EL EMBUDO DE HEIIU EN VIVO (10 min)')
preg(doc, '"Los mínimos del cargo son 40 empresas con decisor y 12 reuniones al mes, y 2, 4 y 6 firmas acumuladas en los meses 1 a 3. Constrúyeme el embudo: de 40 contactos, cuántas reuniones, cuántas propuestas, cuántas firmas, cuántas matrículas por convenio, cuánto recaudo y en qué mes entra."')
preg(doc, '"¿Qué cuenta para ti como decisor contactado?"')
busca(doc, 'la respuesta correcta es "hablé con la persona que decide y tengo nombre, cargo y celular". Si cuenta correos enviados o mensajes de LinkedIn, sus 40 son 8.')
preg(doc, '"En BBVA el concesionario te recibía porque eras BBVA. Aquí nadie conoce la academia. ¿Cómo abres la puerta de una empresa de 40 empleados la primera semana, con nombres: a quién llamas primero?"')
busca(doc, 'que tenga una red propia en Bucaramanga (concesionarios, vendedores de maquinaria, alcaldías, pymes de la zona) y la ponga en números: cuántas puertas abre en el mes 1 sin marca. Si la respuesta es "con la propuesta" o "con la causa", no tiene red: tiene discurso.')
preg(doc, '"Compara lo que escribiste en la ficha (a cuántos gerentes puedes sentarte en 2 semanas y a quiénes nombraste) con las 5 empresas de tu tarea. ¿Cuál es tu mes 1 de verdad: la gente que ya conoces o las 5 de la tarea? ¿Y cuántas matrículas da cada una?"')
busca(doc, 'lee la ficha y la tarea antes y anota aquí si coinciden o no: ______________________. La ficha es lo que tiene; la tarea es lo que cree que queríamos oír. Que lo reconozca y que ponga números a su red real. Si defiende empresas grandes y lentas como su mes 1, no está leyendo su propio embudo.')
preg(doc, '"Dime los 5 indicadores que me vas a mostrar cada martes y, para cada uno, el número al que te sacas tú solo."')
busca(doc, 'que las cifras del embudo cuadren entre sí (si dice 40 contactos y 20 reuniones, no sabe lo que es un decisor). Y que el último punto lo diga sin pensar: un ejecutivo senior sabe cuándo está fallando antes de que se lo digan.')

sec(doc, '4. LA PLATA Y LAS CONDICIONES (5 min)')
preg(doc, '"Explícame con tus palabras cómo se te paga: el básico, el piso de los 3 meses y la comisión. ¿Sobre qué se paga la comisión?"')
busca(doc, 'sobre lo cobrado, no sobre lo firmado. Si no lo sabe, se le explica una vez y se anota.')
preg(doc, '"Si tu plan trae una meta en pesos: ¿en qué mes entra esa plata a la cuenta?"')
busca(doc, 'que distinga firmado, inicial pagada y recaudo real.')
preg(doc, '"Vienes de un salario de banca. El básico aquí es el mínimo legal y lo que protege los meses 1 a 3 es el piso. Si en el mes 4 la comisión todavía no reemplaza lo que ganabas, ¿qué haces?"')
busca(doc, 'una respuesta concreta sobre su caja personal ("tengo para X meses", "mi esposa trabaja", "ya lo calculé") en vez de "no hay problema". Lleva 5 meses sin trabajar: si no lo ha calculado, se va en el mes 5.')
preg(doc, '"Martes y jueves en la sede, con tu planilla abierta, con Diana y conmigo. ¿Te sirve?"')
busca(doc, 'un sí sin condiciones. Es un hecho del cargo, no se negocia.')

sec(doc, 'CIERRE (2 min)')
caja(doc, ['"¿Hay algo de todo esto que te haría decir que no?"  → y te callas. Lo que diga es su objeción real; se anota, no se responde.',
           '"Gracias. Diana te escribe esta semana."  → sin oferta, sin fecha, sin comparar con el otro candidato.'])

sec(doc, 'DESPUÉS: 3 LÍNEAS A DIANA, EL MISMO DÍA')
preg(doc, '¿Se sabe sus números o se los sabían por él?')
preg(doc, '¿El embudo de Heiiu que armó cuadra o es discurso? ¿Tiene red propia sin la marca del banco?')
preg(doc, 'Veredicto: contratar, no contratar, o una condición para contratar. Y si entrevistaste a los dos: cuál y por qué, en una línea.')
t(doc, 'Lo que no haces: venderle Heiiu, ofrecer plata o fecha, contarle del otro candidato. Tú verificas; Diana ofrece después. No preguntar por embarazo, hijos, estado civil, religión, política ni salud.', size=9.5, color=GRIS)

out1 = BASE + r'\ENTREVISTA_CARLOS_REINALDO_NUMEROS.docx'
try: doc.save(out1); print('OK', out1)
except PermissionError: doc.save(out1.replace('.docx', '_v2.docx')); print('ABIERTO -> _v2')

# ---------- 2. LLAMADA A LA REFERENCIA (DIANA) ----------
doc = nuevo_doc()
titulo(doc, 'LLAMADA A LA REFERENCIA DE BBVA — LA HACE DIANA', 'Javier Nicolás Garcés, coordinador nacional Vendor, BBVA · 312 483 1818 · 10 minutos · Antes de la entrevista con Carlos')
t(doc, 'No es una confrontación: es escuchar y anotar hechos. No le menciones ninguna cifra de Reinaldo; que las diga él. Si una respuesta no cuadra con la ficha, no lo discutes con Javier: lo anotas y se lo pasas a Carlos para la entrevista.', bold=True)
sec(doc, 'Apertura (tal cual)')
caja(doc, ['"Buenas, Javier, le habla Diana Gómez, gerente de Heiiu English Academy en Bucaramanga. Reinaldo Joya lo dio como referencia y autorizó por escrito que lo llamáramos. ¿Tiene 10 minutos?"'])
sec(doc, 'Las 6 preguntas, en este orden')
for q, nota in [
 ('"¿Qué cargo tenía Reinaldo con usted y de qué fechas a qué fechas?"', 'Su HV: ejecutivo de cuenta Leasing, junio 2021 a mayo 2025.'),
 ('"¿Cuántos concesionarios con convenio tenía la zona cuando él llegó y cuántos cuando salió?"', 'Él escribió: canal creado desde cero, más de 10 convenios. No se lo digas.'),
 ('"¿Esos convenios los firmó él solo o con un equipo? ¿Cuántos le referían clientes de verdad al final?"', ''),
 ('"¿Cómo fue su cumplimiento de presupuesto, año a año?"', 'Él dice $12.000 millones al año, presupuesto cumplido.'),
 ('"¿Por qué salió en mayo?"', 'Lo que diga Javier es la respuesta. Compárala después con lo que Reinaldo escribió en la ficha.'),
 ('"¿Lo volvería a contratar para abrir un canal desde cero, sin la marca del banco detrás?"', 'Y te callas. El silencio o el "pues..." ya responde.'),
]:
    p = doc.add_paragraph(style='List Number'); p.paragraph_format.space_after = Pt(1)
    r = p.add_run(q); r.font.size = Pt(10)
    if nota:
        r = p.add_run('  ' + nota); r.font.size = Pt(9.5); r.font.color.rgb = GRIS
    t(doc, 'Respuesta: ________________________________________________________________________', size=9.5, after=4)
sec(doc, 'Cierre')
caja(doc, ['"Le agradezco mucho su tiempo. Que esté bien."  → no cuentas si lo vas a contratar ni qué te pareció.'])
sec(doc, 'Segunda llamada, solo si la primera deja dudas: Bancolombia')
t(doc, 'Nelson Gómez, coordinador regional · 320 672 8560 · 5 minutos · Misma apertura. Dos preguntas: "¿De qué fechas a qué fechas estuvo Reinaldo con usted y por qué salió en 2018?" y "¿Lo volvería a contratar?". Lo de 2018 importa porque entre Bancolombia y BBVA hay 3 años sin nada escrito en la HV.', size=9.5)
t(doc, 'Después: una línea a Carlos con lo que no cuadró (si nada, "BBVA confirma la ficha").', size=9.5, color=GRIS)

out2 = BASE + r'\LLAMADA_REFERENCIA_BBVA_DIANA.docx'
try: doc.save(out2); print('OK', out2)
except PermissionError: doc.save(out2.replace('.docx', '_v2.docx')); print('ABIERTO -> _v2')
