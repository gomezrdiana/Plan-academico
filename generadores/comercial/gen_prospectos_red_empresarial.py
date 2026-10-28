# -*- coding: utf-8 -*-
"""Excel de prospectos B2B a partir de la red empresarial de Bucaramanga (lista de gerencia, 27/10/2026),
verificada por agente (fuente: recursos/comercial/alianzas/prospectos_red_empresarial.csv).
Por rol, sin nombres de personas. Institutos de idiomas no van en el archivo (regla: competidores nunca en docs que circulan).
Salida: recursos/comercial/alianzas/PROSPECTOS_RED_EMPRESARIAL.xlsx + copia en la carpeta del gerente de alianzas.
Uso: python -X utf8 generadores/comercial/gen_prospectos_red_empresarial.py"""
import os, csv, shutil
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CSV = os.path.join(ROOT, 'recursos', 'comercial', 'alianzas', 'prospectos_red_empresarial.csv')
OUT = os.path.join(ROOT, 'recursos', 'comercial', 'alianzas', 'PROSPECTOS_RED_EMPRESARIAL.xlsx')
JOHN = os.path.join(ROOT, 'recursos', 'comercial', 'INDUCCION_COMERCIAL', 'contratacion comercial', 'LINKEDIN DIANA GOMEZ', 'Candidatos', 'John Hinestroza')
NARANJA = 'E54C09'; CREMA = 'FFF8E7'

rows = list(csv.DictReader(open(CSV, encoding='utf-8'), delimiter=';'))
# fuera: institutos de idiomas (no se nombran); PCG = primera empresa aliada (caso de referencia)
rows = [r for r in rows if 'praxis' not in r['nombre'].lower()]
for r in rows:
    if r['nombre'].lower().startswith('pcg'):
        r['prioridad'] = '1'; r['angulo'] = 'aliada: primera firma, caso de referencia'
        r['necesidad_ingles'] = 'baja (nomina pequena); sirve como primer convenio firmado para mostrar a las demas'
        r['nota'] = 'acceso directo de gerencia'
        r['sector'] = r['sector'] or 'construccion'

def clave(r):
    p = r['prioridad'].strip()
    return (0 if p == '1' else 1 if p == '2' else 2 if p == '3' else 3, r['nombre'].lower())
rows.sort(key=clave)

wb = Workbook(); ws = wb.active; ws.title = 'Prospectos'
cab = ['#', 'Empresa / entidad', 'Sector', 'Sede en Bucaramanga', 'Tamano (empleados)', 'Tipo', 'Necesidad de ingles',
       'Angulo Heiiu', 'Prioridad', 'Quien presenta', 'Estado', 'Fecha contacto', 'Siguiente paso', 'LinkedIn', 'Nota / fuente']
ws.append(cab)
for i, r in enumerate(rows, 1):
    ws.append([i, r['nombre'], r['sector'], r['presencia_bga'], r['tamano'], r['tipo'], r['necesidad_ingles'],
               r['angulo'], r['prioridad'], '', 'sin contactar', '', '', r['linkedin'], r['nota']])
anchos = [4, 38, 26, 16, 22, 16, 44, 30, 9, 16, 14, 13, 26, 28, 60]
for c, w in enumerate(anchos, 1):
    ws.column_dimensions[get_column_letter(c)].width = w
fino = Side(style='thin', color='DDDDDD')
for c in range(1, len(cab) + 1):
    cell = ws.cell(row=1, column=c); cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor=NARANJA)
    cell.alignment = Alignment(wrap_text=True, vertical='center')
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    p = str(row[8].value).strip()
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical='top'); cell.border = Border(top=fino, bottom=fino, left=fino, right=fino)
    if p == '1':
        for cell in row: cell.fill = PatternFill('solid', fgColor=CREMA)
        row[1].font = Font(bold=True)
    if p == '-':
        for cell in row: cell.font = Font(color='999999')
ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions
ws.row_dimensions[1].height = 30

g = wb.create_sheet('Como usar')
for l in ['PROSPECTOS DE LA RED EMPRESARIAL DE BUCARAMANGA · gerente de alianzas',
          '',
          'Prioridad 1 = nomina grande en Bucaramanga y necesidad clara de ingles. Prioridad 2 = candidata, verificar tamano. Prioridad 3 = pequena o decide fuera de Santander. "-" = excluida (universidades, cajas y cooperativas van por otra via).',
          'Quien presenta: gerencia anota en esa columna las empresas donde hay acceso directo; el gerente de alianzas NO llega en frio a esas: pide la presentacion primero.',
          'Nunca se escribe en el grupo de la red: cada empresa se aborda una a una, en privado.',
          'Oferta por angulo: "aliada" = convenio empresa aliada (sin costo para la empresa; empleados e hijos con beca del Fondo). "plan empresa" = la empresa paga cupos (tarifa por volumen, orden firmada por gerencia). "individual" = programa de 100 horas uno a uno para directivos, medicos y abogados.',
          'Pago directo de becas por la empresa: argumento tributario pendiente de validacion del contador; no se usa hasta que gerencia lo autorice por escrito.',
          'Entidades publicas con sindicato (acueducto, electrificadora, alcantarillado): la capacitacion entra por acuerdos sindicales y planes de bienestar; la ventana es enero (planes de capacitacion del ano).',
          'Empresa con riesgo reputacional (mineria en zona de paramo): no se contacta sin decision de gerencia.',
          'Tamanos y fuentes: directorios de terceros (EMIS, FinalScout y similares), prensa o informes de la empresa; "est." = estimacion; "sin dato" = no se encontro, verificar en RUES o Camara de Comercio antes de la cita.',
          'Estado: sin contactar / presentado / cita / propuesta / firmado / no. Esta hoja alimenta el tablero semanal de alianzas; quien registra primero una empresa es su dueno.']:
    g.append([l])
g.column_dimensions['A'].width = 150
for row in g.iter_rows():
    for cell in row: cell.alignment = Alignment(wrap_text=True, vertical='top')
g['A1'].font = Font(bold=True, color=NARANJA, size=12)

try:
    wb.save(OUT); print('OK', OUT)
except OSError:
    OUT2 = OUT.replace('.xlsx', '_v2.xlsx'); wb.save(OUT2); print('ABIERTO ->', OUT2); OUT = OUT2
if os.path.isdir(JOHN):
    dst = '\\\\?\\' + os.path.join(JOHN, 'PROSPECTOS_RED_EMPRESARIAL.xlsx')
    try:
        shutil.copyfile(OUT, dst); print('copia gerente OK')
    except OSError as e:
        print('copia gerente NO (archivo abierto?)', e)
print('filas:', len(rows), '| P1:', sum(1 for r in rows if r['prioridad'].strip() == '1'), '| P2:', sum(1 for r in rows if r['prioridad'].strip() == '2'))
