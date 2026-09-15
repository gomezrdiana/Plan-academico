# -*- coding: utf-8 -*-
"""Analizador de productos para reventa en Mercado Libre Colombia.
Excel con la matematica de margen: se llena mirando ML en el navegador + cotizaciones reales.
Correr desde la raiz del repo: python "recursos/comercial/CFO/MERCADO_LIBRE/gen_analizador_productos.py"
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

NARANJA = 'E54C09'
CREMA = 'FFF8E7'
wb = openpyxl.Workbook()

def estilo_encabezado(ws, ncols, fila=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=fila, column=c)
        cell.font = Font(bold=True, color='FFFFFF', size=10)
        cell.fill = PatternFill('solid', fgColor=NARANJA)
        cell.alignment = Alignment(wrap_text=True, vertical='center')

# ============ HOJA 1: CANDIDATOS ============
ws = wb.active
ws.title = 'CANDIDATOS'
cols = [
    ('Producto', 28), ('Link ML del lider', 30), ('Precio venta ML (lider)', 14),
    ('Vendidos visibles (+N)', 13), ('# vendedores fuertes', 12), ('Costo cotizado (proveedor)', 15),
    ('% comision ML', 11), ('Costo envio a cargo vendedor', 14), ('Otros costos/unidad (empaque, 4x1000)', 14),
    ('Comision ML $', 13), ('MARGEN NETO $', 14), ('MARGEN %', 10), ('VEREDICTO', 13),
    ('Proveedor / contacto', 20), ('Notas', 24),
]
for i, (nombre, ancho) in enumerate(cols, 1):
    ws.cell(row=1, column=i, value=nombre)
    ws.column_dimensions[get_column_letter(i)].width = ancho
estilo_encabezado(ws, len(cols))

ws.cell(row=2, column=1, value='>> Ejemplo (borrar): cepillo secador 5 en 1')
ws.cell(row=2, column=3, value=95000)
ws.cell(row=2, column=4, value=500)
ws.cell(row=2, column=5, value=3)
ws.cell(row=2, column=6, value=48000)
ws.cell(row=2, column=7, value=0.16)
ws.cell(row=2, column=8, value=6000)
ws.cell(row=2, column=9, value=1500)
for f in range(2, 42):
    ws.cell(row=f, column=10, value=f'=IF(C{f}="","",C{f}*G{f})')
    ws.cell(row=f, column=11, value=f'=IF(C{f}="","",C{f}-F{f}-J{f}-H{f}-I{f})')
    ws.cell(row=f, column=12, value=f'=IF(C{f}="","",K{f}/C{f})')
    ws.cell(row=f, column=13, value=f'=IF(C{f}="","",IF(AND(L{f}>=0.2,D{f}>=100),"GO",IF(L{f}>=0.2,"GO (validar demanda)",IF(L{f}>=0.12,"AJUSTAR COSTO","NO GO"))))')
    ws.cell(row=f, column=12).number_format = '0.0%'
    ws.cell(row=f, column=7).number_format = '0.0%'
    for c in (3, 6, 8, 9, 10, 11):
        ws.cell(row=f, column=c).number_format = '#,##0'
ws.freeze_panes = 'A2'

dv = DataValidation(type='list', formula1='"0.13,0.16,0.18,0.20"', allow_blank=True, showDropDown=False)
ws.add_data_validation(dv)
dv.add('G2:G41')

# ============ HOJA 2: REGLAS Y COSTOS ML ============
ws2 = wb.create_sheet('REGLAS ML')
datos = [
    ('CONCEPTO', 'VALOR DE REFERENCIA', 'NOTA'),
    ('Comision ML publicacion clasica', '~13%', 'Menos exposicion, sin cuotas'),
    ('Comision ML publicacion premium', '~16-18% segun categoria', 'Mas exposicion + pago a cuotas'),
    ('VERIFICAR comisiones exactas', 'mercadolibre.com.co -> Costos de vender', 'Cambian por categoria — confirmar antes de publicar'),
    ('Envio gratis obligatorio', 'Productos desde ~$60.000: el vendedor subsidia parte', 'Ver tabla de Mercado Envios vigente'),
    ('Reputacion nueva', 'Primeras ~10 ventas sin termometro', 'Precio agresivo al arrancar; despachar YA; cero reclamos'),
    ('4x1000', '0,4% de cada retiro bancario', 'Va en "otros costos"'),
    ('Devoluciones', 'Presupuestar 3-5% de las ventas', 'ML favorece al comprador'),
    ('REGLA GO', 'Margen neto >= 20% Y demanda visible (+100 vendidos del lider)', 'Debajo de 12% no es negocio: es trasteo'),
    ('REGLA DE ORO', 'Primero medir en ML que se vende -> DESPUES buscar la distribucion de eso', 'Nunca al reves'),
    ('La ventaja real', 'No es publicar en ML (eso lo hace cualquiera): es el costo de distribuidor que otros no tienen', ''),
]
for f, fila in enumerate(datos, 1):
    for c, val in enumerate(fila, 1):
        ws2.cell(row=f, column=c, value=val)
estilo_encabezado(ws2, 3)
ws2.column_dimensions['A'].width = 32
ws2.column_dimensions['B'].width = 48
ws2.column_dimensions['C'].width = 45

# ============ HOJA 3: COMO SE USA ============
ws3 = wb.create_sheet('COMO SE USA')
pasos = [
    'COMO SE USA ESTE ANALIZADOR (10 minutos por producto)',
    '',
    '1. En el navegador: buscar el producto en mercadolibre.com.co. Ordenar por "Mas vendidos" o mirar los primeros resultados.',
    '2. Del lider de la busqueda copiar: precio, "+N vendidos" (aparece bajo el titulo del producto), y contar cuantos vendedores con reputacion verde venden lo mismo.',
    '3. Conseguir COTIZACION REAL del producto (distribuidor, mayorista, fabricante). Sin cotizacion real no se llena la fila — el video de Instagram no es una cotizacion.',
    '4. Llenar la fila en CANDIDATOS. Las columnas J-M se calculan solas.',
    '5. Veredicto GO = margen neto >= 20% con demanda probada. Solo entonces: pedido de PRUEBA chico (10-30 unidades), venderlo COMPLETO, medir rotacion real.',
    '6. Si la prueba rota en menos de 45 dias con el margen previsto: escalar. Si no: siguiente producto. La plata se decide con este cuadro, jamas con un reel.',
    '',
    'AUTOMATIZACION: el script ml_dashboard_api.py de esta carpeta llena candidatos automaticamente con los mas vendidos por categoria de la API oficial de ML (requiere token gratuito — ver instrucciones en el script).',
]
for f, txt in enumerate(pasos, 1):
    ws3.cell(row=f, column=1, value=txt)
ws3.cell(row=1, column=1).font = Font(bold=True, size=12, color=NARANJA)
ws3.column_dimensions['A'].width = 140

out = r'C:\Users\pedro\Downloads\diana gt\heiiu\estrategia global Heiiu\recursos\comercial\CFO\MERCADO_LIBRE\ANALIZADOR_PRODUCTOS_ML.xlsx'
wb.save(out)
print('OK', out)
